"""Readable, exact-snapshot financial workbooks; no formula or cross-tenant data."""
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
import json
import tempfile

import openpyxl
import pytest

from models import db, MarketplaceFinanceFact, MarketplaceFinanceSync
from services import marketplace_finance_export as export
from tests.test_ozon_finance_workspace import workspace, get, listing, child


def download(f, suffix='', *, snapshot=None, anchor=None):
    data=get(f).get_json()['data']
    snapshot = snapshot if snapshot is not None else data['snapshot_sync']['id']
    anchor = anchor if anchor is not None else data['coverage']['requested_end']
    auth,login=f._auth()
    with auth,login:
        return f.client.get(f'/marketplaces/api/finance/export.xlsx?account_id={f.own_account_id}&snapshot_id={snapshot}&as_of={anchor}'+suffix)


def workbook(response):
    assert response.status_code==200,response.get_json(silent=True)
    assert response.headers['Content-Type']==export.MIME
    assert response.headers['Cache-Control']=='private, no-store'
    return openpyxl.load_workbook(BytesIO(response.data),data_only=False)


def test_complete_workbook_not_preview_and_no_double_count(workspace):
    f=workspace
    with f.app.app_context():
        own=listing(f)
        for i in range(121):child(f,i,own.id)
        original=db.session.get(MarketplaceFinanceFact,f.fact_id)
        for i in range(30):
            db.session.add(MarketplaceFinanceFact(sync_id=original.sync_id,seller_id=original.seller_id,
                marketplace_id=original.marketplace_id,account_id=original.account_id,accrual_id=f'new-{i}',
                fact_date=date.today(),accrued_category='ITEM',total_amount=Decimal('0.0001'),currency='USD',
                amount_sign='positive',source_fingerprint='c'*64,observed_at=datetime.utcnow()))
        db.session.commit()
    response=download(f);wb=workbook(response)
    assert wb.sheetnames==['О выгрузке','Итоги','Начисления','Товары','Расшифровка']
    assert wb['Начисления'].max_row==32
    assert wb['Товары'].max_row==wb['Расшифровка'].max_row==122
    rows=list(wb['Итоги'].values)[1:]
    assert rows[0][0:2]==('RUB',1) and rows[0][6:8]==(-10,'-10.0000')
    assert rows[1][0:2]==('USD',30) and Decimal(str(rows[1][6]))==Decimal('0.0030')
    assert wb['Начисления'].freeze_panes=='A2' and wb['Начисления'].auto_filter.ref=='A1:J32'
    assert wb['Товары']['C2'].value=='offer'
    assert response.headers['X-Finance-Account-Id']==str(f.own_account_id)
    assert '.xlsx' in response.headers['Content-Disposition']
    filtered=workbook(download(f,'&type_id=7&sign=negative'))
    assert filtered['Начисления'].max_row==2 and filtered['Расшифровка'].max_row==122
    assert filtered['Итоги']['G2'].value==-10
    empty=workbook(download(f,'&search=%25'))
    assert empty['Начисления'].max_row==empty['Итоги'].max_row==1


def test_foreign_children_and_listing_never_appear(workspace):
    f=workspace
    with f.app.app_context():
        secret=listing(f,foreign=True);child(f,1,secret.id);child(f,2,secret.id,foreign=True);db.session.commit()
    wb=workbook(download(f));rows=list(wb['Товары'].values)
    assert len(rows)==2 and rows[1][2:4]==(None,None) and rows[1][4]=='Связанная карточка недоступна'
    assert wb['Расшифровка'].max_row==2
    assert 'Secret other store' not in repr([list(ws.values) for ws in wb])
    assert workbook(download(f,'&type_id=9'))['Начисления'].max_row==1


def test_exact_text_and_formula_free_cells(workspace):
    f=workspace
    with f.app.app_context():
        own=listing(f);own.title='=HYPERLINK("https://example.invalid","text")';own.offer_id='0000123456789012345678'
        fact=db.session.get(MarketplaceFinanceFact,f.fact_id);fact.accrual_id='=1+2';fact.unit_number='@SUM(1,2)'
        child(f,1,own.id);db.session.commit()
    wb=workbook(download(f))
    assert wb['Товары']['A2'].value=='=1+2' and wb['Товары']['A2'].data_type=='s'
    assert wb['Товары']['C2'].value=='0000123456789012345678'
    assert wb['Товары']['D2'].data_type=='s'
    assert wb['Начисления']['D2'].value=='@SUM(1,2)' and wb['Начисления']['D2'].data_type=='s'
    assert all(cell.data_type!='f' and cell.hyperlink is None for sheet in wb for row in sheet for cell in row)
    assert export._money_cells(Decimal('9999999999999999.1234'))==('9999999999999999.1234','9999999999999999.1234')
    assert export._money_cells(Decimal('-0.0001'))==(Decimal('-0.0001'),'-0.0001')


def test_exact_snapshot_scope_and_strict_export_query(workspace):
    f=workspace
    data=get(f).get_json()['data'];sid=data['snapshot_sync']['id']
    for suffix in ['&account_id=1','&page=1','&format=csv','&seller_id=2','&category=bad','&type_id=٠']:
        assert download(f,suffix).status_code==400,suffix
    assert download(f,snapshot=99999).status_code==404
    assert download(f,anchor='9999-01-01').status_code==400
    auth,login=f._auth()
    with auth,login:
        for query in [f'account_id={f.own_account_id}',f'account_id={f.own_account_id}&snapshot_id={sid}']:
            assert f.client.get('/marketplaces/api/finance/export.xlsx?'+query).status_code==400
    with f.app.app_context():
        db.session.get(MarketplaceFinanceSync,sid).status='running';db.session.commit()
    assert download(f,snapshot=sid).status_code==404
    with f.app.app_context():
        snap=db.session.get(MarketplaceFinanceSync,sid);snap.status='completed';snap.account_id=f.foreign_account_id;db.session.commit()
    assert download(f,snapshot=sid).status_code==404


def test_partial_and_empty_export_never_claim_full_zero_balance(workspace):
    f=workspace
    with f.app.app_context():
        fact=db.session.get(MarketplaceFinanceFact,f.fact_id);fact.sync.period_start=date.today()-timedelta(days=3);db.session.commit()
    response=download(f);wb=workbook(response)
    metadata=dict(list(wb['О выгрузке'].values)[1:])
    assert 'НЕПОЛНОЕ' in metadata['Покрытие периода']
    assert '-partial.xlsx' in response.headers['Content-Disposition']
    assert 'не означает нулевой' in metadata['Пустая выборка']


def test_budgets_fail_without_partial_file_and_temp_files_are_removed(workspace,monkeypatch,tmp_path):
    f=workspace
    with f.app.app_context():child(f,1);db.session.commit()
    with monkeypatch.context() as m:
        m.setattr(export,'MAX_FACTS',0);r=download(f);assert r.status_code==422 and r.is_json
    with monkeypatch.context() as m:
        m.setattr(export,'MAX_CHILDREN',0);r=download(f);assert r.status_code==422 and r.is_json
    with monkeypatch.context() as m:
        m.setattr(tempfile,'tempdir',str(tmp_path))
        original=export._check_budget
        def stop_after_writer_started(deadline):
            if list(tmp_path.glob('openpyxl.*')):raise export.FinanceExportLimit('Budget test')
            original(deadline)
        m.setattr(export,'_check_budget',stop_after_writer_started)
        r=download(f);assert r.status_code==422 and r.is_json
        assert not list(tmp_path.glob('openpyxl.*'))
    with f.app.app_context():
        db.session.get(MarketplaceFinanceFact,f.fact_id).unit_number='bad\x00field';db.session.commit()
    assert download(f).status_code==400


def test_read_transaction_survives_concurrent_snapshot_retention(workspace,tmp_path):
    """Retention commits between reading the parent and its facts on real WAL SQLite."""
    import sqlite3
    from sqlalchemy import create_engine,event
    f=workspace
    with f.app.app_context():
        child(f,1);db.session.commit();sid=db.session.get(MarketplaceFinanceFact,f.fact_id).sync_id
        db.session.remove();original_engine=db.engines[None]
        path=tmp_path/'finance-read-race.db'
        with original_engine.connect() as source:
            target=sqlite3.connect(path);source.connection.driver_connection.backup(target)
            target.execute('PRAGMA journal_mode=WAL');target.close()
        engine=create_engine('sqlite:///'+str(path));db.engines[None]=engine
        deleted=[]
        def prune(connection,cursor,statement,parameters,context,executemany):
            if (not deleted and statement.startswith('SELECT marketplace_finance_syncs.id \n')
                    and connection.connection.driver_connection.in_transaction):
                with sqlite3.connect(path) as writer:
                    for table in ['marketplace_finance_components','marketplace_finance_fact_items']:
                        writer.execute(f'DELETE FROM {table} WHERE fact_id=?',(f.fact_id,))
                    writer.execute('DELETE FROM marketplace_finance_facts WHERE id=?',(f.fact_id,))
                    writer.execute('DELETE FROM marketplace_finance_syncs WHERE id=?',(sid,))
                deleted.append(True)
        event.listen(engine,'after_cursor_execute',prune)
        try:
            result=download(f,snapshot=sid);wb=workbook(result)
            assert deleted and wb['Начисления'].max_row==2 and wb['Товары'].max_row==2
            assert wb['Расшифровка'].max_row==2 and wb['Итоги']['G2'].value==-10
            with sqlite3.connect(path) as reader:
                assert reader.execute('SELECT COUNT(*) FROM marketplace_finance_facts').fetchone()[0]==0
        finally:
            event.remove(engine,'after_cursor_execute',prune);db.session.remove();db.engines[None]=original_engine;engine.dispose()


def test_new_completed_run_does_not_replace_pinned_export(workspace):
    f=workspace
    with f.app.app_context():
        source=db.session.get(MarketplaceFinanceFact,f.fact_id).sync;sid=source.id
        newer=MarketplaceFinanceSync(seller_id=source.seller_id,marketplace_id=source.marketplace_id,
            account_id=source.account_id,period_code=source.period_code,period_start=source.period_start,
            period_end=source.period_end,current_date=source.current_date,status='completed',phase='completed',
            contract_version=source.contract_version,request_fingerprint=source.request_fingerprint,
            completed_at=datetime.utcnow()+timedelta(seconds=1))
        db.session.add(newer);db.session.commit()
    wb=workbook(download(f,snapshot=sid))
    assert wb['Начисления'].max_row==2 and wb['Итоги']['G2'].value==-10
