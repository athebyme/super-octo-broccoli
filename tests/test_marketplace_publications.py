# -*- coding: utf-8 -*-
"""Ozon product import is strict, durable, idempotent and tenant-scoped."""

from datetime import datetime, timedelta
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch

from flask import Flask
from openpyxl import load_workbook

from models import (
    BackgroundJob,
    ImportedProduct,
    Marketplace,
    MarketplaceAttributeDefinition,
    MarketplaceAttributeValue,
    MarketplaceCategoryMapping,
    MarketplaceListing,
    MarketplaceOperation,
    MarketplaceProductDraft,
    MarketplaceProductType,
    MarketplaceTaxonomyCategory,
    Seller,
    SellerMarketplaceAccount,
    User,
    db,
)
from services.marketplace_adapters import MarketplaceCredentials
from services.marketplace_drafts import (
    MarketplaceDraftConflict,
    MarketplaceDraftService,
)
from services.marketplace_fact_pack import MarketplaceFactPackBuilder
from services.image_lab_service import ImageLabError
from services.marketplace_publications import (
    MarketplacePublicationNotFound,
    MarketplacePublicationService,
)
from services.ozon_api_client import (
    OzonAPIError,
    OzonAmbiguousWriteError,
)
from services.ozon_product_import import (
    OzonProductImportContract,
    OzonProductImportPayloadError,
    OzonProductImportProtocolError,
)
from services.ozon_product_state import OzonProductStateContract
from services.ozon_reference_service import OzonReferenceService
from services.ozon_bulk_upload import (
    OzonBulkUploadError,
    OzonBulkUploadService,
    OzonBulkUploadValidationError,
)
from services.ozon_bulk_repair import OzonBulkRepairService
from routes.ozon_bulk_uploads import register_ozon_bulk_upload_routes


SYNTHETIC_CREDENTIALS = MarketplaceCredentials(
    external_account_id="synthetic-client",
    api_key="synthetic-key",
)


class SyntheticPublicationAdapter:
    capabilities = {"catalog_read", "catalog_write"}

    def __init__(
        self,
        *,
        offer_exists=False,
        ambiguous=False,
        preflight_error=False,
        status="imported",
        cursor_on_exact=False,
    ):
        self.offer_exists = offer_exists
        self.ambiguous = ambiguous
        self.preflight_error = preflight_error
        self.status = status
        self.cursor_on_exact = cursor_on_exact
        self.archived = False
        self.list_calls = []
        self.submitted_payloads = []
        self.status_calls = []
        self.live_payload = None
        self.full_read_calls = []

    def require_capability(self, capability):
        if capability not in self.capabilities:
            raise AssertionError(f"missing capability {capability}")

    def list_products(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.list_calls.append(payload)
        if self.preflight_error:
            raise OzonAPIError(
                "synthetic read outage",
                code="synthetic_read_outage",
                request_id="synthetic-request",
            )
        visibility = payload["filter"]["visibility"]
        items = []
        if self.offer_exists and visibility == "ALL":
            items = [{
                "product_id": 987654,
                "offer_id": payload["filter"]["offer_id"][0],
                "archived": False,
                "has_fbo_stocks": False,
                "has_fbs_stocks": False,
            }]
        return {
            "result": {
                "items": items,
                "total": len(items),
                "last_id": (
                    "opaque-current-ozon-cursor"
                    if items and self.cursor_on_exact else ""
                ),
            }
        }

    def get_operation_limits(self, credentials):
        assert credentials == SYNTHETIC_CREDENTIALS
        return {
            "operation_limits": [{
                "operation": "product_create",
                "limit": 100,
                "usage": 4,
                "remaining": 96,
                "reset_at": "2026-07-16T00:00:00Z",
            }]
        }

    def submit_products(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.submitted_payloads.append(payload)
        if self.ambiguous:
            self.offer_exists = True
            self.live_payload = deepcopy(payload)
            raise OzonAmbiguousWriteError(
                "synthetic ambiguous write",
                code="synthetic_ambiguous_write",
                request_id="synthetic-write-request",
            )
        return {"result": {"task_id": 456}}

    def get_submission(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.status_calls.append(payload)
        errors = []
        product_id = 987654 if self.status == "imported" else 0
        if self.status == "imported" and self.submitted_payloads:
            self.live_payload = deepcopy(self.submitted_payloads[-1])
            self.offer_exists = True
        if self.status in {"failed", "skipped"}:
            errors = [{
                "code": "SYNTHETIC_REJECT",
                "message": "Synthetic item rejection",
            }]
        return {
            "result": {
                "items": [{
                    "offer_id": "safe-offer",
                    "product_id": product_id,
                    "status": self.status,
                    "errors": errors,
                }],
                "total": 1,
            }
        }

    def _item(self):
        if self.live_payload is None:
            raise AssertionError("synthetic live payload is unavailable")
        return self.live_payload["items"][0]

    @staticmethod
    def _media(item):
        return SyntheticFullStateAdapter._media(item)

    @staticmethod
    def _description(item):
        return SyntheticFullStateAdapter._description(item)

    def get_products(self, credentials, payload):
        return SyntheticFullStateAdapter.get_products(self, credentials, payload)

    def get_product_attributes(self, credentials, payload):
        return SyntheticFullStateAdapter.get_product_attributes(
            self, credentials, payload,
        )

    def read_prices(self, credentials, payload):
        return SyntheticFullStateAdapter.read_prices(self, credentials, payload)

    def get_product_pictures(self, credentials, payload):
        return SyntheticFullStateAdapter.get_product_pictures(
            self, credentials, payload,
        )


class SyntheticFullStateAdapter(SyntheticPublicationAdapter):
    """Synthetic exact-state adapter for update/archive state machines."""

    def __init__(
        self,
        live_payload=None,
        *,
        create_mode=False,
        cursor_on_exact=False,
    ):
        super().__init__(
            offer_exists=not create_mode,
            cursor_on_exact=cursor_on_exact,
        )
        self.live_payload = deepcopy(live_payload)
        self.archived = False
        self.pending_payload = None
        self.archive_calls = []
        self.full_read_calls = []

    def _item(self):
        if not self.live_payload:
            raise AssertionError("synthetic live payload is unavailable")
        return self.live_payload["items"][0]

    def list_products(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.list_calls.append(payload)
        visibility = payload["filter"]["visibility"]
        visible = self.live_payload is not None and (
            (visibility == "ALL" and not self.archived)
            or (visibility == "ARCHIVED" and self.archived)
        )
        items = []
        if visible:
            items = [{
                "product_id": 987654,
                "offer_id": self._item()["offer_id"],
                "archived": self.archived,
                "has_fbo_stocks": False,
                "has_fbs_stocks": False,
            }]
        return {
            "result": {"items": items, "total": len(items), "last_id": ""}
        }

    def get_operation_limits(self, credentials):
        assert credentials == SYNTHETIC_CREDENTIALS
        return {
            "operation_limits": [{
                "operation": "product_update" if self.live_payload else "product_create",
                "limit": 100,
                "usage": 1,
                "remaining": 99,
            }]
        }

    def submit_products(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.submitted_payloads.append(deepcopy(payload))
        self.pending_payload = deepcopy(payload)
        return {"result": {"task_id": 456}}

    def get_submission(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.status_calls.append(payload)
        if self.pending_payload is not None:
            self.live_payload = self.pending_payload
            self.pending_payload = None
            self.offer_exists = True
            self.archived = False
        return {
            "result": {
                "items": [{
                    "offer_id": self._item()["offer_id"],
                    "product_id": 987654,
                    "status": "imported",
                    "errors": [],
                }],
                "total": 1,
            }
        }

    def get_products(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.full_read_calls.append(("info", deepcopy(payload)))
        item = self._item()
        media = self._media(item)
        return {"items": [{
            "id": 987654,
            "offer_id": item["offer_id"],
            "name": item["name"],
            "description": self._description(item),
            "description_category_id": item["description_category_id"],
            "type_id": item["type_id"],
            "is_archived": self.archived,
            "barcodes": [item["barcode"]] if item.get("barcode") else [],
            "primary_image": [media["primary_image"]],
            "images": [media["primary_image"], *media["images"]],
            "statuses": {},
            "visibility_details": {},
            "errors": [],
        }]}

    def get_product_attributes(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.full_read_calls.append(("attributes", deepcopy(payload)))
        item = self._item()
        return {
            "result": [{
                "id": 987654,
                "offer_id": item["offer_id"],
                "name": item["name"],
                "description_category_id": item["description_category_id"],
                "type_id": item["type_id"],
                "attributes": deepcopy(item["attributes"]),
                "complex_attributes": deepcopy(item["complex_attributes"]),
                "width": item["width"],
                "height": item["height"],
                "depth": item["depth"],
                "dimension_unit": item["dimension_unit"],
                "weight": item["weight"],
                "weight_unit": item["weight_unit"],
                "barcodes": [item["barcode"]] if item.get("barcode") else [],
                "images": [],
                "sku": 7654321,
            }],
            "total": 1,
            "last_id": "",
        }

    def read_prices(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.full_read_calls.append(("prices", deepcopy(payload)))
        item = self._item()
        price = {
            "price": item["price"],
            "currency_code": item["currency_code"],
            "vat": item["vat"],
        }
        if item.get("old_price"):
            price["old_price"] = item["old_price"]
        return {
            "items": [{
                "product_id": 987654,
                "offer_id": item["offer_id"],
                "price": price,
            }],
            "total": 1,
            "cursor": (
                "opaque-current-ozon-price-cursor"
                if self.cursor_on_exact else ""
            ),
        }

    def get_product_pictures(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.full_read_calls.append(("pictures", deepcopy(payload)))
        media = self._media(self._item())
        return {"items": [{
            "product_id": 987654,
            "primary_photo": [media["primary_image"]],
            "photo": media["images"],
            "color_photo": [media["color_image"]] if media.get("color_image") else [],
            "errors": [],
        }]}

    def archive_products(self, credentials, payload):
        assert credentials == SYNTHETIC_CREDENTIALS
        self.archive_calls.append(deepcopy(payload))
        self.archived = True
        return {"result": True}

    @staticmethod
    def _media(item):
        images = list(item.get("images", []))
        primary = item.get("primary_image")
        if not primary:
            primary = images.pop(0)
        result = {"primary_image": primary, "images": images}
        if item.get("color_image"):
            result["color_image"] = item["color_image"]
        return result

    @staticmethod
    def _description(item):
        for attribute in item.get("attributes", []):
            if attribute.get("id") == 4191 and attribute.get("values"):
                return attribute["values"][0].get("value")
        return None


class SyntheticTypeOmittingAdapter(SyntheticFullStateAdapter):
    """Mirror Ozon's import-only 8229 omission in attributes read-back."""

    def get_product_attributes(self, credentials, payload):
        response = super().get_product_attributes(credentials, payload)
        item = response["result"][0]
        item["attributes"] = [
            attribute
            for attribute in item["attributes"]
            if str(attribute.get("id")) != "8229"
        ]
        return response


class OzonPublicationFixture:
    def setUp(self):
        self.app = Flask(__name__)
        self.app.config.update(
            TESTING=True,
            SQLALCHEMY_DATABASE_URI="sqlite://",
            SQLALCHEMY_TRACK_MODIFICATIONS=False,
        )
        db.init_app(self.app)
        self.context = self.app.app_context()
        self.context.push()
        db.create_all()
        now = datetime.utcnow()
        self.user = User(
            username="ozon-publication",
            email="ozon-publication@test.local",
            is_active=True,
        )
        self.user.set_password("synthetic-password")
        self.seller = Seller(user=self.user, company_name="Publication Seller")
        self.foreign_user = User(
            username="ozon-publication-foreign",
            email="ozon-publication-foreign@test.local",
            is_active=True,
        )
        self.foreign_user.set_password("synthetic-password")
        self.foreign_seller = Seller(
            user=self.foreign_user,
            company_name="Foreign Seller",
        )
        self.marketplace = Marketplace(
            name="Ozon",
            code="ozon",
            adapter_code="ozon",
            is_active=True,
            categories_synced_at=now,
            categories_snapshot_hash="tree-hash",
        )
        db.session.add_all([
            self.seller,
            self.foreign_seller,
            self.marketplace,
        ])
        db.session.flush()
        self.account = SellerMarketplaceAccount(
            seller_id=self.seller.id,
            marketplace_id=self.marketplace.id,
            external_account_id="synthetic-client",
            label="Synthetic Ozon",
            is_active=True,
            connection_status="connected",
        )
        category = MarketplaceTaxonomyCategory(
            marketplace_id=self.marketplace.id,
            external_category_id="10",
            name="Категория",
            full_path="Категория",
            is_available=True,
            last_seen_at=now,
        )
        db.session.add_all([self.account, category])
        db.session.flush()
        self.product_type = MarketplaceProductType(
            marketplace_id=self.marketplace.id,
            category_id=category.id,
            external_type_id="777",
            name="Тип",
            is_available=True,
            is_enabled=True,
            attributes_synced_at=now,
            attributes_sync_status="success",
            attributes_schema_hash="schema-hash",
            attributes_version=3,
            attributes_count=1,
            required_attributes_count=0,
        )
        self.source = ImportedProduct(
            seller_id=self.seller.id,
            external_id="source-1",
            external_vendor_code="safe-offer",
            source_type="synthetic",
            title="Безопасный товар",
            description="Наблюдаемое описание",
            category="Категория",
        )
        db.session.add_all([self.product_type, self.source])
        db.session.flush()
        description = MarketplaceAttributeDefinition(
            marketplace_id=self.marketplace.id,
            product_type_id=self.product_type.id,
            external_attribute_id="4191",
            name="Аннотация",
            data_type="String",
            is_required=False,
            max_value_count=1,
            is_available=True,
            is_enabled=True,
            last_seen_at=now,
        )
        self.draft = MarketplaceProductDraft(
            seller_id=self.seller.id,
            marketplace_id=self.marketplace.id,
            account_id=self.account.id,
            imported_product_id=self.source.id,
            product_type_id=self.product_type.id,
            offer_id="safe-offer",
            external_category_id="10",
            external_type_id="777",
            status="ready",
            source_fact_hash="a" * 64,
            source_facts_json="{}",
            provenance_json="{}",
            content_json=json.dumps({
                "name": "Безопасный товар",
                "description": "Наблюдаемое описание",
            }, ensure_ascii=False),
            attributes_json="[]",
            complex_attributes_json="[]",
            media_json=json.dumps({
                "images": ["https://img.test/product.jpg"],
            }),
            dimensions_json=json.dumps({
                "width": "200",
                "height": "30",
                "depth": "300",
                "dimension_unit": "MILLIMETERS",
                "weight": "250",
                "weight_unit": "GRAMS",
            }),
            barcodes_json=json.dumps(["4600000000001"]),
            commercial_json=json.dumps({
                "price": "1000",
                "old_price": "1200",
                "vat": "0.22",
                "currency_code": "RUB",
            }),
            schema_version=3,
            schema_hash="schema-hash",
            validation_status="valid",
            validation_result_json=json.dumps({
                "publishable": True,
                "errors": [],
                "warnings": [],
            }),
            validated_at=now,
        )
        db.session.add_all([description, self.draft])
        db.session.commit()
        self.expected_version = self.draft.version

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def validation_result(self):
        result = {
            "publishable": True,
            "errors": [],
            "warnings": [],
            "schema": {
                "hash": "schema-hash",
                "version": 3,
            },
        }
        if self.draft.published_listing_id is not None:
            _, baseline = MarketplaceDraftService.publication_documents(
                self.draft
            )
            result["update_baseline"] = baseline
        return result

    def start(self, adapter, *, key="publication-key-0001"):
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            return MarketplacePublicationService.start_publication(
                seller_id=self.seller.id,
                draft_id=self.draft.id,
                expected_version=self.expected_version,
                idempotency_key=key,
                created_by_user_id=self.user.id,
                adapter=adapter,
                credentials=SYNTHETIC_CREDENTIALS,
            )

    def desired_payload(self):
        return OzonProductImportContract.build_payload(self.draft)

    def prior_payload(self):
        payload = deepcopy(self.desired_payload())
        item = payload["items"][0]
        item["name"] = "Предыдущее название"
        item["price"] = "900"
        item["old_price"] = "1100"
        item["images"] = ["https://img.test/prior.jpg"]
        for attribute in item["attributes"]:
            if attribute["id"] == 4191:
                attribute["values"] = [{"value": "Предыдущее описание"}]
        return payload

    def add_required_dictionary_attribute(
        self,
        *,
        external_attribute_id,
        name,
        external_value_id,
        value,
    ):
        now = datetime.utcnow()
        definition = MarketplaceAttributeDefinition(
            marketplace_id=self.marketplace.id,
            product_type_id=self.product_type.id,
            external_attribute_id=str(external_attribute_id),
            name=name,
            data_type="String",
            is_required=True,
            dictionary_id=f"dictionary-{external_attribute_id}",
            max_value_count=1,
            is_collection=False,
            is_available=True,
            is_enabled=True,
            last_seen_at=now,
            values_synced_at=now,
            values_sync_status="success",
            values_snapshot_hash=f"values-{external_attribute_id}",
            values_version=1,
            values_count=1,
        )
        db.session.add(definition)
        db.session.flush()
        official = MarketplaceAttributeValue(
            marketplace_id=self.marketplace.id,
            product_type_id=self.product_type.id,
            attribute_id=definition.id,
            external_value_id=str(external_value_id),
            value=value,
            value_normalized=value.casefold(),
            is_available=True,
            last_seen_at=now,
        )
        attributes = json.loads(self.draft.attributes_json)
        attributes.append({
            "attribute_id": str(external_attribute_id),
            "complex_id": "0",
            "values": [{
                "dictionary_value_id": str(external_value_id),
                "value": value,
            }],
        })
        self.draft.attributes_json = json.dumps(
            attributes,
            ensure_ascii=False,
        )
        self.product_type.attributes_count += 1
        self.product_type.required_attributes_count += 1
        db.session.add(official)
        db.session.commit()
        self.expected_version = self.draft.version
        return definition

    def enable_import_only_type_attribute(self):
        return self.add_required_dictionary_attribute(
            external_attribute_id="8229",
            name="Тип",
            external_value_id="93477",
            value=self.product_type.name,
        )

    def attach_listing(self, payload=None):
        payload = payload or self.prior_payload()
        item = payload["items"][0]
        now = datetime.utcnow()

        def listing_attribute(raw):
            return {
                "id": str(raw["id"]),
                "complex_id": (
                    str(raw.get("complex_id"))
                    if raw.get("complex_id") not in (None, 0, "0")
                    else None
                ),
                "values": [
                    {
                        key: (
                            str(value)
                            if key == "dictionary_value_id"
                            else value
                        )
                        for key, value in raw_value.items()
                    }
                    for raw_value in raw["values"]
                ],
            }

        attributes = [
            listing_attribute(raw)
            for raw in item["attributes"]
        ]
        complex_attributes = [{
            "attributes": [
                listing_attribute(raw)
                for raw in group["attributes"]
            ]
        } for group in item["complex_attributes"]]
        media = {"images": list(item["images"])}
        if item.get("primary_image"):
            media["primary_image"] = item["primary_image"]
        if item.get("color_image"):
            media["color_image"] = item["color_image"]
        price_values = {
            "price": item["price"],
            "vat": item["vat"],
            "currency_code": item["currency_code"],
        }
        if item.get("old_price"):
            price_values["old_price"] = item["old_price"]
        listing = MarketplaceListing(
            seller_id=self.seller.id,
            marketplace_id=self.marketplace.id,
            account_id=self.account.id,
            imported_product_id=self.source.id,
            product_type_id=self.product_type.id,
            offer_id=item["offer_id"],
            external_product_id="987654",
            external_category_id=str(item["description_category_id"]),
            external_type_id=str(item["type_id"]),
            title=item["name"],
            normalized_status="active",
            is_available=True,
            is_archived=False,
            attributes_json=json.dumps(attributes),
            complex_attributes_json=json.dumps(complex_attributes),
            media_json=json.dumps(media),
            dimensions_json=json.dumps({
                "width": item["width"],
                "height": item["height"],
                "depth": item["depth"],
                "dimension_unit": item["dimension_unit"],
                "weight": item["weight"],
                "weight_unit": item["weight_unit"],
            }),
            barcodes_json=json.dumps(
                [item["barcode"]] if item.get("barcode") else []
            ),
            price_summary_json=json.dumps({
                "available": True,
                "currency": item["currency_code"],
                "values": price_values,
            }),
            stock_summary_json=json.dumps({"preserve": True}),
            list_synced_at=now,
            info_synced_at=now,
            attributes_synced_at=now,
            prices_synced_at=now,
            last_seen_at=now,
            sync_fingerprint="b" * 64,
        )
        db.session.add(listing)
        db.session.flush()
        self.draft.published_listing_id = listing.id
        db.session.commit()
        self.expected_version = self.draft.version
        return listing

    def start_update(self, adapter, *, key="product-update-key-0001"):
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            return MarketplacePublicationService.start_update(
                seller_id=self.seller.id,
                draft_id=self.draft.id,
                expected_version=self.expected_version,
                idempotency_key=key,
                created_by_user_id=self.user.id,
                adapter=adapter,
                credentials=SYNTHETIC_CREDENTIALS,
            )


class OzonProductImportContractTest(OzonPublicationFixture, unittest.TestCase):
    def test_builder_is_whitelist_only_and_maps_description_attribute(self):
        payload = OzonProductImportContract.build_payload(self.draft)
        self.assertEqual(set(payload), {"items"})
        item = payload["items"][0]
        self.assertEqual(item["offer_id"], "safe-offer")
        self.assertEqual(item["description_category_id"], 10)
        self.assertEqual(item["type_id"], 777)
        self.assertEqual(item["dimension_unit"], "mm")
        self.assertEqual(item["weight_unit"], "g")
        self.assertEqual(item["barcode"], "4600000000001")
        self.assertNotIn("description", item)
        self.assertNotIn("images360", item)
        description = next(
            value for value in item["attributes"]
            if value["id"] == 4191
        )
        self.assertEqual(description, {
            "id": 4191,
            "complex_id": 0,
            "values": [{"value": "Наблюдаемое описание"}],
        })

    def test_builder_rejects_long_offer_fractional_physical_and_extra_barcode(self):
        self.draft.offer_id = "x" * 51
        with self.assertRaises(OzonProductImportPayloadError):
            OzonProductImportContract.build_payload(self.draft)
        self.draft.offer_id = "safe-offer"
        dimensions = json.loads(self.draft.dimensions_json)
        dimensions["weight"] = "250.5"
        self.draft.dimensions_json = json.dumps(dimensions)
        with self.assertRaises(OzonProductImportPayloadError):
            OzonProductImportContract.build_payload(self.draft)
        dimensions["weight"] = "250"
        self.draft.dimensions_json = json.dumps(dimensions)
        self.draft.barcodes_json = json.dumps(["1", "2"])
        with self.assertRaises(OzonProductImportPayloadError):
            OzonProductImportContract.build_payload(self.draft)

    def test_provider_roundtrip_omission_never_hides_other_drift(self):
        self.enable_import_only_type_attribute()
        submitted = self.desired_payload()
        live = deepcopy(submitted)
        live["items"][0]["attributes"] = [
            attribute
            for attribute in live["items"][0]["attributes"]
            if attribute["id"] != 8229
        ]

        self.assertEqual(
            OzonProductStateContract.provider_roundtrip_omissions(
                submitted,
                live,
                omittable_simple_attribute_ids=["8229"],
            ),
            ("8229",),
        )

        changed = deepcopy(live)
        changed["items"][0]["price"] = "1001"
        self.assertIsNone(
            OzonProductStateContract.provider_roundtrip_omissions(
                submitted,
                changed,
                omittable_simple_attribute_ids=["8229"],
            )
        )

        returned_wrong_value = deepcopy(submitted)
        type_attribute = next(
            attribute
            for attribute in returned_wrong_value["items"][0]["attributes"]
            if attribute["id"] == 8229
        )
        type_attribute["values"][0]["value"] = "Другой тип"
        self.assertIsNone(
            OzonProductStateContract.provider_roundtrip_omissions(
                submitted,
                returned_wrong_value,
                omittable_simple_attribute_ids=["8229"],
            )
        )

    def test_status_exact_set_and_skipped_fail_closed(self):
        skipped = OzonProductImportContract.normalize_status({
            "result": {
                "items": [{
                    "offer_id": "safe-offer",
                    "product_id": 0,
                    "status": "skipped",
                    "errors": [{"code": "SKIPPED"}],
                }],
                "total": 1,
            }
        }, expected_offer_ids=["safe-offer"])
        self.assertEqual(skipped["aggregate_status"], "failed")
        with self.assertRaises(OzonProductImportProtocolError):
            OzonProductImportContract.normalize_status({
                "result": {
                    "items": [{
                        "offer_id": "foreign-offer",
                        "product_id": 1,
                        "status": "imported",
                        "errors": [],
                    }],
                    "total": 1,
                }
            }, expected_offer_ids=["safe-offer"])

    def test_quota_prefers_new_contract_and_malformed_new_shape_does_not_fallback(self):
        normalized = OzonProductImportContract.normalize_quota({
            "operation_limits": [{
                "operation": "product_create",
                "limit": 100,
                "usage": 10,
                "remaining": 90,
            }],
            "daily_create": {"limit": 999, "usage": 0},
        })
        self.assertEqual(normalized["source"], "operation_limits")
        self.assertEqual(normalized["remaining"], 90)
        split = OzonProductImportContract.normalize_quota({
            "operation_limits": [
                {
                    "operation": "product_create",
                    "limit": 100,
                    "usage": 10,
                    "remaining": 90,
                },
                {
                    "operation": "product_update",
                    "limit": 100,
                    "usage": 100,
                    "remaining": 0,
                },
            ],
        }, mode="create")
        self.assertEqual(split["remaining"], 90)
        self.assertEqual(
            [item["name"] for item in split["entries"]],
            ["product_create"],
        )
        with self.assertRaises(OzonProductImportProtocolError):
            OzonProductImportContract.normalize_quota({
                "operation_limits": {"unexpected": True},
                "daily_create": {"limit": 999, "usage": 0},
            })

    def test_quota_accepts_current_operation_caps_with_daily_counters(self):
        normalized = OzonProductImportContract.normalize_quota({
            "daily_create": {
                "limit": 1500,
                "usage": 12,
                "reset_at": "2026-07-25T00:00:00Z",
            },
            "daily_update": {
                "limit": 20000,
                "usage": 100,
                "reset_at": "2026-07-25T00:00:00Z",
            },
            "operation_limits": [{
                "limit": 30000,
                "limit_type": "PRODUCT_IMPORT",
            }],
            "total": {
                "limit": 13219,
                "usage": 8719,
            },
        }, mode="create")

        self.assertEqual(
            normalized["source"],
            "operation_caps_with_daily_counters",
        )
        self.assertEqual(normalized["remaining"], 1488)
        self.assertEqual(
            [entry["name"] for entry in normalized["entries"]],
            ["daily_create", "total"],
        )
        self.assertEqual(normalized["operation_caps"], [{
            "limit_type": "PRODUCT_IMPORT",
            "limit": 30000,
        }])

    def test_quota_does_not_infer_remaining_from_operation_cap_only(self):
        with self.assertRaises(OzonProductImportProtocolError):
            OzonProductImportContract.normalize_quota({
                "operation_limits": [{
                    "limit": 30000,
                    "limit_type": "PRODUCT_IMPORT",
                }],
            })



class BulkEnqueueTest(OzonPublicationFixture, unittest.TestCase):
    """Bulk create/update enqueue durable operations without provider writes."""

    def _second_draft(self, *, offer_id="bulk-offer-2", published_listing_id=None):
        import json as _json
        source = ImportedProduct(
            seller_id=self.seller.id,
            external_id=f"source-{offer_id}",
            external_vendor_code=offer_id,
            source_type="synthetic",
            title="Безопасный товар",
            description="Наблюдаемое описание",
            category="Категория",
        )
        db.session.add(source)
        db.session.flush()
        draft = MarketplaceProductDraft(
            seller_id=self.seller.id,
            marketplace_id=self.marketplace.id,
            account_id=self.account.id,
            imported_product_id=source.id,
            product_type_id=self.product_type.id,
            offer_id=offer_id,
            external_category_id="10",
            external_type_id="777",
            status="ready",
            source_fact_hash="a" * 64,
            source_facts_json="{}",
            provenance_json="{}",
            content_json=self.draft.content_json,
            attributes_json="[]",
            complex_attributes_json="[]",
            media_json=self.draft.media_json,
            dimensions_json=self.draft.dimensions_json,
            barcodes_json=_json.dumps(["4600000000002"]),
            commercial_json=self.draft.commercial_json,
            schema_version=3,
            schema_hash="schema-hash",
            validation_status="valid",
            validation_result_json=self.draft.validation_result_json,
            validated_at=self.draft.validated_at,
            published_listing_id=published_listing_id,
        )
        db.session.add(draft)
        db.session.commit()
        return draft

    def _enqueue(self, draft_ids):
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            return MarketplacePublicationService.enqueue_bulk_publications(
                seller_id=self.seller.id,
                account_id=self.account.id,
                draft_ids=draft_ids,
                created_by_user_id=self.user.id,
            )

    def _enqueue_updates(self, draft_ids):
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            return MarketplacePublicationService.enqueue_bulk_updates(
                seller_id=self.seller.id,
                account_id=self.account.id,
                draft_ids=draft_ids,
                created_by_user_id=self.user.id,
            )

    def test_mixed_set_enqueues_valid_and_skips_rest_with_reasons(self):
        linked = self._second_draft(
            offer_id="bulk-linked", published_listing_id=1,
        )
        result = self._enqueue([self.draft.id, 999999, linked.id])

        self.assertEqual(len(result["queued"]), 1)
        self.assertEqual(result["queued"][0]["draft_id"], self.draft.id)
        skipped_ids = {row["draft_id"] for row in result["skipped"]}
        self.assertEqual(skipped_ids, {999999, linked.id})

        operation = db.session.get(
            MarketplaceOperation,
            result["queued"][0]["operation_id"]
        )
        # Durable queued: провайдер не вызывался, submit оставлен scheduler-у
        self.assertEqual(operation.status, "queued")
        self.assertEqual(operation.attempt_count, 0)
        self.assertIsNotNone(operation.next_poll_at)

        # Повторный bulk не создаёт дублей: активная операция уже есть
        repeat = self._enqueue([self.draft.id])
        self.assertEqual(repeat["queued"], [])
        self.assertEqual(len(repeat["skipped"]), 1)
        self.assertEqual(
            MarketplaceOperation.query.filter_by(
                draft_id=self.draft.id,
            ).count(),
            1,
        )

    def test_id_validation_rejects_whole_request(self):
        from services.marketplace_publications import (
            MarketplacePublicationValidationError,
        )
        with self.assertRaises(MarketplacePublicationValidationError):
            self._enqueue([self.draft.id, self.draft.id])
        with self.assertRaises(MarketplacePublicationValidationError):
            self._enqueue([True])
        with self.assertRaises(MarketplacePublicationValidationError):
            self._enqueue(list(range(1, 52)))

    def test_non_publishable_draft_is_skipped_not_fatal(self):
        broken = self._second_draft(offer_id="bulk-broken")
        invalid = {"publishable": False, "errors": [
            {"code": "missing_required"},
        ], "warnings": []}
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            side_effect=[self.validation_result(), invalid],
        ):
            result = MarketplacePublicationService.enqueue_bulk_publications(
                seller_id=self.seller.id,
                account_id=self.account.id,
                draft_ids=[self.draft.id, broken.id],
                created_by_user_id=self.user.id,
            )
        self.assertEqual(len(result["queued"]), 1)
        self.assertEqual(len(result["skipped"]), 1)
        self.assertEqual(result["skipped"][0]["draft_id"], broken.id)

    def test_bulk_update_queues_exact_full_state_preflight_without_provider_io(self):
        listing = self.attach_listing()

        result = self._enqueue_updates([self.draft.id])

        self.assertEqual(result["skipped"], [])
        self.assertEqual(len(result["queued"]), 1)
        operation = db.session.get(
            MarketplaceOperation,
            result["queued"][0]["operation_id"],
        )
        self.assertEqual(operation.operation_kind, "product_update")
        self.assertEqual(operation.status, "queued")
        self.assertEqual(operation.attempt_count, 0)
        self.assertEqual(operation.listing_id, listing.id)
        self.assertEqual(
            json.loads(operation.snapshot.before_state_json)["state"],
            "pending_live_preflight",
        )

        repeated = self._enqueue_updates([self.draft.id])
        self.assertEqual(repeated["queued"], [])
        self.assertEqual(len(repeated["skipped"]), 1)
        self.assertEqual(
            MarketplaceOperation.query.filter_by(
                draft_id=self.draft.id,
                operation_kind="product_update",
            ).count(),
            1,
        )

    def test_bulk_update_skips_unlinked_draft_and_rejects_invalid_ids(self):
        result = self._enqueue_updates([self.draft.id])
        self.assertEqual(result["queued"], [])
        self.assertEqual(len(result["skipped"]), 1)
        self.assertIn("listing", result["skipped"][0]["reason"])

        from services.marketplace_publications import (
            MarketplacePublicationValidationError,
        )
        with self.assertRaises(MarketplacePublicationValidationError):
            self._enqueue_updates([self.draft.id, self.draft.id])
        with self.assertRaises(MarketplacePublicationValidationError):
            self._enqueue_updates(list(range(1, 52)))


class OzonBulkUploadServiceTest(OzonPublicationFixture, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self.app.config.update(
            MARKETPLACE_OZON_ENABLED=True,
            MARKETPLACE_OZON_PUBLICATION_ENABLED=True,
        )
        # The bulk HTTP path only checks that credentials were saved. It never
        # decrypts them and never calls the provider.
        self.account._credentials_encrypted = "synthetic-encrypted-value"
        self.account.settings_json = json.dumps({"default_vat": "0.22"})
        db.session.commit()

    def _create_run(self):
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            side_effect=lambda draft: self.validation_result(),
        ):
            return OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
                created_by_user_id=self.user.id,
            )

    def test_history_filters_owned_account_before_limit_and_ignores_bad_json(self):
        other = SellerMarketplaceAccount(seller_id=self.seller.id, marketplace_id=self.marketplace.id,
                                        external_account_id='second-owned', label='Second', is_active=True)
        foreign = SellerMarketplaceAccount(seller_id=self.foreign_seller.id, marketplace_id=self.marketplace.id,
                                          external_account_id='foreign-owned', label='Foreign', is_active=True)
        db.session.add_all([other, foreign])
        db.session.flush()
        for uid, seller_id, progress in [
            ('history-own', self.seller.id, json.dumps({'account_id': self.account.id, 'items': []})),
            ('history-second', self.seller.id, json.dumps({'account_id': other.id, 'items': []})),
            ('history-invalid', self.seller.id, '{invalid json'),
            ('history-foreign', self.foreign_seller.id, json.dumps({'account_id': self.account.id, 'items': []})),
        ]:
            db.session.add(BackgroundJob(job_uid=uid, seller_id=seller_id,
                                         job_type=OzonBulkUploadService.JOB_TYPE, progress_data=progress))
        db.session.commit()
        rows = OzonBulkUploadService.list_runs(seller_id=self.seller.id, account_id=self.account.id, limit=1)
        self.assertEqual([row.job_uid for row in rows], ['history-own'])
        with self.assertRaises(OzonBulkUploadError) as error:
            OzonBulkUploadService.list_runs(seller_id=self.seller.id, account_id=foreign.id)
        self.assertEqual(error.exception.status_code, 404)

    def test_create_run_prepares_and_queues_without_provider_io(self):
        with patch(
            "services.supplier_service."
            "hydrate_missing_imported_observed_snapshot",
            return_value=False,
        ) as hydrate:
            job = self._create_run()
        document = OzonBulkUploadService.public_document(job, detail=True)

        self.assertEqual(job.job_type, "ozon_bulk_upload")
        self.assertEqual(job.status, "running")
        self.assertEqual(document["summary"]["active"], 1)
        self.assertEqual(document["items"][0]["status"], "queued")
        self.assertEqual(
            document["category_mapping_preflight"]["code"],
            "observed_category_mapping_reconciled",
        )
        self.assertEqual(
            document["category_mapping_preflight"]["created"],
            0,
        )
        operation = db.session.get(
            MarketplaceOperation,
            document["items"][0]["operation_id"]
        )
        self.assertEqual(operation.status, "queued")
        self.assertEqual(operation.attempt_count, 0)
        hydrate.assert_called_once()
        self.assertEqual(
            hydrate.call_args.args[0].id,
            self.source.id,
        )

    def test_run_distinguishes_media_preparation_from_provider_queue(self):
        job = self._create_run()
        document = OzonBulkUploadService.public_document(job, detail=True)
        operation = db.session.get(
            MarketplaceOperation,
            document["items"][0]["operation_id"],
        )
        summary = json.loads(operation.request_summary_json)
        summary.update({
            "media_asset_state": "preparing",
            "media_asset_prepared": 2,
            "media_asset_source_total": 5,
        })
        operation.request_summary_json = json.dumps(summary)
        operation.error_code = "media_preparation_pending"
        operation.error_message = "Подготавливаем фото для Ozon: 2 из 5"
        db.session.commit()

        reconciled = OzonBulkUploadService.reconcile_run(
            seller_id=self.seller.id,
            job_uid=job.job_uid,
        )
        result = OzonBulkUploadService.public_document(
            reconciled,
            detail=True,
        )

        self.assertEqual(result["items"][0]["status"], "preparing_media")
        self.assertEqual(result["summary"]["preparing_media"], 1)
        self.assertEqual(result["summary"]["active"], 1)
        self.assertEqual(
            result["items"][0]["media_delivery"],
            {"state": "preparing", "prepared": 2, "total": 5},
        )

    def test_existing_source_suffix_is_queued_as_full_update_not_create(self):
        self.source.source_type = "sexoptovik"
        self.source.external_id = "7725"
        self.source.external_vendor_code = "id-7725-1366"
        self.draft.offer_id = "id-7725-1366"
        prior = self.prior_payload()
        prior["items"][0]["offer_id"] = "id-7725-1364"
        prior["items"][0]["name"] = "Существующая карточка Ozon 7725"
        listing = self.attach_listing(prior)
        listing.external_product_id = "7725001"
        listing.imported_product_id = None
        listing.link_status = "unlinked"
        listing.link_source = None
        listing.sync_fingerprint = "d" * 64
        self.draft.published_listing_id = None
        db.session.commit()

        job = self._create_run()
        document = OzonBulkUploadService.public_document(job, detail=True)
        item = document["items"][0]
        operation = db.session.get(
            MarketplaceOperation,
            item["operation_id"],
        )

        self.assertEqual(item["action"], "update")
        self.assertEqual(item["offer_id"], "id-7725-1364")
        self.assertEqual(item["listing_id"], listing.id)
        self.assertEqual(operation.operation_kind, "product_update")
        self.assertEqual(listing.imported_product_id, self.source.id)
        self.assertEqual(self.draft.published_listing_id, listing.id)

    def test_forbidden_brand_is_needs_input_without_provider_operation(self):
        self.source.brand = "SVAKOM"
        fact_pack = MarketplaceFactPackBuilder.build(self.source)
        self.draft.source_fact_hash = fact_pack["fact_hash"]
        self.draft.source_facts_json = json.dumps(
            fact_pack,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
        self.draft.provenance_json = json.dumps(
            fact_pack["provenance"],
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
        db.session.commit()

        job = OzonBulkUploadService.create_run(
            seller_id=self.seller.id,
            account_id=self.account.id,
            imported_product_ids=[self.source.id],
            created_by_user_id=self.user.id,
        )
        document = OzonBulkUploadService.public_document(
            job,
            detail=True,
        )
        self.assertEqual(job.status, "completed")
        self.assertEqual(document["items"][0]["status"], "needs_input")
        self.assertEqual(
            document["items"][0]["code"],
            "ozon_brand_forbidden",
        )
        self.assertFalse(
            document["items"][0]["completeness"]["brand_allowed"],
        )
        self.assertEqual(
            MarketplaceOperation.query.filter_by(
                draft_id=self.draft.id,
            ).count(),
            0,
        )

    def test_bulk_run_reaches_confirmed_provider_success_end_to_end(self):
        job = self._create_run()
        document = OzonBulkUploadService.public_document(job, detail=True)
        operation_id = document["items"][0]["operation_id"]
        adapter = SyntheticPublicationAdapter()
        submitted_at = datetime.utcnow()

        submitted = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation_id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
            now=submitted_at,
        )
        self.assertEqual(submitted.status, "submitted")
        self.assertEqual(submitted.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)

        confirmed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation_id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
            now=submitted_at + timedelta(seconds=16),
        )
        self.assertEqual(confirmed.status, "succeeded")
        self.assertEqual(len(adapter.submitted_payloads), 1)

        reconciled = OzonBulkUploadService.reconcile_run(
            seller_id=self.seller.id,
            job_uid=job.job_uid,
        )
        result = OzonBulkUploadService.public_document(
            reconciled,
            detail=True,
        )
        self.assertEqual(reconciled.status, "completed")
        self.assertEqual(result["summary"]["outcome"], "success")
        self.assertEqual(result["summary"]["succeeded"], 1)
        self.assertEqual(result["summary"]["created"], 1)
        self.assertEqual(result["summary"]["updated"], 0)
        self.assertEqual(result["items"][0]["status"], "succeeded")
        self.assertEqual(result["items"][0]["action"], "create")

    def test_published_card_queues_full_update_and_no_change_is_success(self):
        desired = self.desired_payload()
        listing = self.attach_listing(desired)
        job = self._create_run()
        document = OzonBulkUploadService.public_document(job, detail=True)
        item = document["items"][0]
        operation = db.session.get(
            MarketplaceOperation,
            item["operation_id"],
        )
        self.assertEqual(item["action"], "update")
        self.assertEqual(operation.operation_kind, "product_update")
        self.assertEqual(operation.listing_id, listing.id)
        self.assertEqual(operation.attempt_count, 0)

        adapter = SyntheticFullStateAdapter(desired)
        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(completed.attempt_count, 0)
        self.assertEqual(adapter.submitted_payloads, [])
        self.assertEqual(
            json.loads(completed.item_results_json)[0]["status"],
            "already_current",
        )

        reconciled = OzonBulkUploadService.reconcile_run(
            seller_id=self.seller.id,
            job_uid=job.job_uid,
        )
        result = OzonBulkUploadService.public_document(
            reconciled,
            detail=True,
        )
        self.assertEqual(result["summary"]["outcome"], "success")
        self.assertEqual(result["summary"]["already_current"], 1)
        self.assertEqual(result["summary"]["created"], 0)
        self.assertEqual(result["summary"]["updated"], 0)
        self.assertEqual(result["items"][0]["status"], "already_current")
        db.session.refresh(self.draft)
        self.assertEqual(self.draft.status, "published")

    def test_published_card_full_update_writes_once_and_reports_updated(self):
        prior = self.prior_payload()
        self.attach_listing(prior)
        job = self._create_run()
        item = OzonBulkUploadService.public_document(
            job,
            detail=True,
        )["items"][0]
        operation_id = item["operation_id"]
        adapter = SyntheticFullStateAdapter(prior)
        started = datetime.utcnow()

        submitted = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation_id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
            now=started,
        )
        self.assertEqual(submitted.status, "submitted")
        self.assertEqual(submitted.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)

        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation_id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
            now=started + timedelta(seconds=16),
        )
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(len(adapter.submitted_payloads), 1)
        result = OzonBulkUploadService.public_document(
            OzonBulkUploadService.reconcile_run(
                seller_id=self.seller.id,
                job_uid=job.job_uid,
            ),
            detail=True,
        )
        self.assertEqual(result["summary"]["updated"], 1)
        self.assertEqual(result["summary"]["created"], 0)
        self.assertEqual(result["items"][0]["action"], "update")

    def test_full_two_hundred_product_run_uses_four_bounded_chunks(self):
        extra_products = [
            ImportedProduct(
                seller_id=self.seller.id,
                external_id=f"bulk-source-{index}",
                source_type="synthetic",
                title=f"Массовый товар {index}",
            )
            for index in range(1, OzonBulkUploadService.MAX_ITEMS)
        ]
        db.session.add_all(extra_products)
        db.session.commit()
        product_ids = [self.source.id, *[item.id for item in extra_products]]
        drafts = {
            product_id: SimpleNamespace(
                id=100_000 + index,
                offer_id=f"bulk-offer-{index}",
                published_listing_id=None,
                status="ready",
                validation_result_json='{"errors":[]}',
                version=1,
            )
            for index, product_id in enumerate(product_ids, start=1)
        }

        def create_draft(**kwargs):
            return drafts[kwargs["imported_product_id"]]

        def return_draft(**kwargs):
            return next(
                draft
                for draft in drafts.values()
                if draft.id == kwargs["draft_id"]
            )

        def enqueue(**kwargs):
            return {
                "queued": [
                    {
                        "draft_id": draft_id,
                        "operation_id": 200_000 + draft_id,
                    }
                    for draft_id in kwargs["draft_ids"]
                ],
                "skipped": [],
            }

        with patch.object(
            MarketplaceDraftService,
            "create_draft",
            side_effect=create_draft,
        ), patch.object(
            MarketplaceDraftService,
            "rebase_source_defaults",
            side_effect=return_draft,
        ), patch.object(
            MarketplaceDraftService,
            "apply_reference_defaults",
            side_effect=return_draft,
        ), patch.object(
            MarketplaceDraftService,
            "apply_account_defaults",
            side_effect=return_draft,
        ), patch.object(
            MarketplaceDraftService,
            "validate_draft",
            side_effect=return_draft,
        ), patch.object(
            MarketplaceDraftService,
            "completeness_summary",
            return_value={},
        ), patch.object(
            MarketplacePublicationService,
            "enqueue_bulk_publications",
            side_effect=enqueue,
        ) as enqueue_bulk, patch.object(
            OzonBulkUploadService,
            "reconcile_run",
            side_effect=lambda **kwargs: BackgroundJob.query.filter_by(
                job_uid=kwargs["job_uid"],
            ).one(),
        ):
            job = OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=product_ids,
                created_by_user_id=self.user.id,
            )

        progress = job.get_progress()
        self.assertEqual(len(progress["items"]), 200)
        self.assertEqual(
            {item["status"] for item in progress["items"]},
            {"queued"},
        )
        self.assertEqual(enqueue_bulk.call_count, 4)
        self.assertEqual(
            [
                len(call.kwargs["draft_ids"])
                for call in enqueue_bulk.call_args_list
            ],
            [50, 50, 50, 50],
        )
        self.assertLessEqual(
            len(job.progress_data.encode("utf-8")),
            OzonBulkUploadService.MAX_PROGRESS_BYTES,
        )

    def test_mixed_create_and_update_batch_is_split_into_bounded_queues(self):
        ready_items = {
            draft_id: {
                "action": "create" if draft_id <= 60 else "update",
                "status": "queued",
            }
            for draft_id in range(1, 121)
        }

        def enqueue(**kwargs):
            return {
                "queued": [{
                    "draft_id": draft_id,
                    "operation_id": 500_000 + draft_id,
                } for draft_id in kwargs["draft_ids"]],
                "skipped": [],
            }

        with patch.object(
            MarketplacePublicationService,
            "enqueue_bulk_publications",
            side_effect=enqueue,
        ) as create_queue, patch.object(
            MarketplacePublicationService,
            "enqueue_bulk_updates",
            side_effect=enqueue,
        ) as update_queue:
            OzonBulkUploadService._enqueue_ready_items(
                seller_id=self.seller.id,
                account_id=self.account.id,
                created_by_user_id=self.user.id,
                ready_items=ready_items,
            )

        self.assertEqual(
            [len(call.kwargs["draft_ids"]) for call in create_queue.call_args_list],
            [50, 10],
        )
        self.assertEqual(
            [len(call.kwargs["draft_ids"]) for call in update_queue.call_args_list],
            [50, 10],
        )
        self.assertEqual(
            {item["status"] for item in ready_items.values()},
            {"queued"},
        )
        self.assertTrue(all(
            isinstance(item.get("operation_id"), int)
            for item in ready_items.values()
        ))

    def test_maximum_run_with_multibyte_reasons_fits_progress_budget(self):
        safe = OzonBulkUploadService._safe_text
        validation_errors = [{
            "code": safe("к" * 500, 80),
            "field": safe("поле" * 500, 120),
            "message": safe("причина" * 500, 300),
        } for _ in range(OzonBulkUploadService.MAX_VALIDATION_ERRORS)]
        document = {
            "version": OzonBulkUploadService.DOCUMENT_VERSION,
            "source": "products",
            "account_id": self.account.id,
            "account_label": safe("кабинет" * 100, 120),
            "created_at": datetime.utcnow().isoformat(),
            "updated_at": datetime.utcnow().isoformat(),
            "items": [{
                "imported_product_id": index + 1,
                "title": safe("товар" * 500, 300),
                "offer_id": safe("артикул" * 500, 200),
                "status": "needs_input",
                "code": validation_errors[0]["code"],
                "message": validation_errors[0]["message"],
                "validation_errors": validation_errors,
                "updated_at": datetime.utcnow().isoformat(),
            } for index in range(OzonBulkUploadService.MAX_ITEMS)],
        }
        job = BackgroundJob(
            job_uid="ozon-upload-" + ("b" * 32),
            seller_id=self.seller.id,
            job_type=OzonBulkUploadService.JOB_TYPE,
        )

        OzonBulkUploadService._store_progress(job, document)

        self.assertLessEqual(
            len(job.progress_data.encode("utf-8")),
            OzonBulkUploadService.MAX_PROGRESS_BYTES,
        )

    def test_legacy_missing_operation_link_requires_manual_review(self):
        job = self._create_run()
        document = job.get_progress()
        operation_id = document["items"][0].pop("operation_id")
        document["items"][0]["updated_at"] = (
            datetime.utcnow() - timedelta(
                seconds=OzonBulkUploadService.PREPARATION_STALE_SECONDS + 1,
            )
        ).isoformat()
        job.set_progress(document)
        db.session.commit()

        reconciled = OzonBulkUploadService.reconcile_run(
            seller_id=self.seller.id,
            job_uid=job.job_uid,
        )
        item = OzonBulkUploadService.public_document(
            reconciled,
            detail=True,
        )["items"][0]

        self.assertIsNone(item.get("operation_id"))
        self.assertEqual(item["status"], "needs_manual_reconciliation")
        self.assertNotEqual(item.get("operation_id"), operation_id)

    def test_reconcile_without_state_change_rotates_run_without_rewriting_progress(self):
        job = self._create_run()
        result_before = job.result_data
        progress_before = job.progress_data
        updated_before = job.updated_at

        reconciled = OzonBulkUploadService.reconcile_run(
            seller_id=self.seller.id,
            job_uid=job.job_uid,
            rotate=True,
        )

        self.assertEqual(reconciled.result_data, result_before)
        self.assertEqual(reconciled.progress_data, progress_before)
        self.assertGreaterEqual(reconciled.updated_at, updated_before)

    def test_reconcile_closes_stale_interrupted_preparation_for_manual_review(self):
        stale_at = datetime.utcnow() - timedelta(
            seconds=OzonBulkUploadService.PREPARATION_STALE_SECONDS + 1,
        )
        document = {
            "version": OzonBulkUploadService.DOCUMENT_VERSION,
            "source": "products",
            "account_id": self.account.id,
            "account_label": self.account.label,
            "created_at": stale_at.isoformat(),
            "updated_at": stale_at.isoformat(),
            "items": [{
                "imported_product_id": self.source.id,
                "title": self.source.title,
                "status": "preparing",
                "updated_at": stale_at.isoformat(),
            }],
        }
        job = BackgroundJob(
            job_uid="ozon-upload-" + ("a" * 32),
            seller_id=self.seller.id,
            job_type=OzonBulkUploadService.JOB_TYPE,
            status="running",
            total=1,
            created_at=stale_at,
            updated_at=stale_at,
        )
        OzonBulkUploadService._store_progress(job, document)
        job.set_result(OzonBulkUploadService._summary(document["items"]))
        db.session.add(job)
        db.session.commit()

        reconciled = OzonBulkUploadService.reconcile_run(
            seller_id=self.seller.id,
            job_uid=job.job_uid,
        )
        item = OzonBulkUploadService.public_document(
            reconciled,
            detail=True,
        )["items"][0]

        self.assertEqual(reconciled.status, "completed")
        self.assertEqual(item["status"], "needs_manual_reconciliation")
        self.assertEqual(item["code"], "legacy_upload_review_required")

    def test_reconcile_closes_run_after_success(self):
        job = self._create_run()
        item = OzonBulkUploadService.public_document(
            job,
            detail=True,
        )["items"][0]
        operation = db.session.get(
            MarketplaceOperation,
            item["operation_id"],
        )
        operation.status = "succeeded"
        operation.completed_at = datetime.utcnow()
        db.session.commit()

        reconciled = OzonBulkUploadService.reconcile_run(
            seller_id=self.seller.id,
            job_uid=job.job_uid,
        )
        document = OzonBulkUploadService.public_document(
            reconciled,
            detail=True,
        )

        self.assertEqual(reconciled.status, "completed")
        self.assertEqual(document["summary"]["outcome"], "success")
        self.assertEqual(document["summary"]["succeeded"], 1)
        self.assertEqual(document["items"][0]["status"], "succeeded")

    def test_manually_stopped_uncertain_operation_closes_run_without_retry(self):
        job = self._create_run()
        item = OzonBulkUploadService.public_document(
            job,
            detail=True,
        )["items"][0]
        operation = db.session.get(
            MarketplaceOperation,
            item["operation_id"],
        )
        operation.status = "uncertain"
        operation.error_code = "manual_uncertain_resolution"
        operation.error_message = (
            "Автоматическая сверка остановлена; итог Ozon неизвестен"
        )
        operation.next_poll_at = None
        db.session.commit()

        reconciled = OzonBulkUploadService.reconcile_run(
            seller_id=self.seller.id,
            job_uid=job.job_uid,
        )
        document = OzonBulkUploadService.public_document(
            reconciled,
            detail=True,
        )

        self.assertEqual(reconciled.status, "completed")
        self.assertEqual(document["summary"]["active"], 0)
        self.assertEqual(document["summary"]["uncertain"], 1)
        self.assertEqual(document["summary"]["uncertain_active"], 0)
        self.assertEqual(document["summary"]["uncertain_stopped"], 1)
        self.assertTrue(
            document["items"][0]["reconciliation_stopped"],
        )
        with self.assertRaises(OzonBulkUploadError) as retry_error:
            OzonBulkUploadService.retry_run(
                seller_id=self.seller.id,
                job_uid=job.job_uid,
                request_key="t" * 24,
                created_by_user_id=self.user.id,
            )
        self.assertEqual(retry_error.exception.code, "draft_review_required")

    def test_reconcile_exposes_normalized_provider_item_reason(self):
        job = self._create_run()
        item = OzonBulkUploadService.public_document(
            job,
            detail=True,
        )["items"][0]
        operation = db.session.get(
            MarketplaceOperation,
            item["operation_id"],
        )
        operation.status = "failed"
        operation.error_code = "product_rejected"
        operation.error_message = "Ozon rejected the card"
        operation.item_results_json = json.dumps([{
            "offer_id": self.draft.offer_id,
            "errors": [{
                "code": "INVALID_ATTRIBUTE",
                "attribute_name": "Бренд",
                "message": "значение не найдено в справочнике",
            }],
        }], ensure_ascii=False)
        operation.completed_at = datetime.utcnow()
        db.session.commit()

        reconciled = OzonBulkUploadService.reconcile_run(
            seller_id=self.seller.id,
            job_uid=job.job_uid,
        )
        document = OzonBulkUploadService.public_document(
            reconciled,
            detail=True,
        )
        failed = document["items"][0]

        self.assertEqual(document["summary"]["failed"], 1)
        self.assertEqual(failed["code"], "INVALID_ATTRIBUTE")
        self.assertIn("Бренд", failed["message"])
        self.assertIn("справочнике", failed["message"])

    def test_non_publishable_card_finishes_with_actionable_reason(self):
        invalid = {
            "publishable": False,
            "errors": [{
                "code": "missing_required_attribute",
                "field": "attributes.123",
                "message": "Заполните обязательную характеристику «Материал»",
            }],
            "warnings": [],
            "schema": {"hash": "schema-hash", "version": 3},
        }
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=invalid,
        ):
            job = OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
                created_by_user_id=self.user.id,
            )
        document = OzonBulkUploadService.public_document(job, detail=True)

        self.assertEqual(job.status, "completed")
        self.assertEqual(document["summary"]["needs_input"], 1)
        self.assertEqual(document["items"][0]["status"], "needs_input")
        self.assertIn("Материал", document["items"][0]["message"])
        self.assertIsNone(document["items"][0].get("operation_id"))

    def test_legacy_stale_reference_wait_requires_new_review_without_enqueue(self):
        stale = {
            "publishable": False,
            "errors": [{
                "code": "schema_stale",
                "field": "product_type_id",
                "message": "Схема Ozon ещё не загружена",
            }],
            "warnings": [],
            "schema": {
                "hash": None,
                "version": 0,
            },
        }
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=stale,
        ):
            job = OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
                created_by_user_id=self.user.id,
            )
        waiting = OzonBulkUploadService.public_document(
            job,
            detail=True,
        )
        self.assertEqual(job.status, "completed")
        self.assertEqual(waiting["summary"]["waiting_reference"], 0)
        self.assertEqual(
            waiting["items"][0]["status"],
            "needs_manual_reconciliation",
        )
        self.assertIsNone(waiting["items"][0].get("operation_id"))

        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            result = OzonBulkUploadService.reconcile_active_runs(limit=20)

        self.assertEqual(result["reconciled"], 0)
        resumed = OzonBulkUploadService.public_document(
            OzonBulkUploadService.get_run(
                seller_id=self.seller.id,
                job_uid=job.job_uid,
                reconcile=False,
            ),
            detail=True,
        )
        self.assertEqual(resumed["summary"]["waiting_reference"], 0)
        self.assertEqual(
            resumed["items"][0]["status"], "needs_manual_reconciliation",
        )
        self.assertIsNone(resumed["items"][0].get("operation_id"))
        self.assertEqual(MarketplaceOperation.query.count(), 0)

    def test_legacy_update_reference_wait_requires_new_review(self):
        self.attach_listing(self.prior_payload())
        stale = {
            "publishable": False,
            "errors": [{
                "code": "schema_stale",
                "field": "product_type_id",
                "message": "Схема Ozon ещё не загружена",
            }],
            "warnings": [],
            "schema": {"hash": None, "version": 0},
        }
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=stale,
        ):
            job = OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
                created_by_user_id=self.user.id,
            )
        waiting = OzonBulkUploadService.public_document(job, detail=True)
        self.assertEqual(waiting["items"][0]["status"], "needs_manual_reconciliation")
        self.assertEqual(waiting["items"][0]["action"], "update")

        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            OzonBulkUploadService.reconcile_active_runs(limit=20)
        resumed = OzonBulkUploadService.public_document(
            OzonBulkUploadService.get_run(
                seller_id=self.seller.id,
                job_uid=job.job_uid,
                reconcile=False,
            ),
            detail=True,
        )
        self.assertEqual(
            resumed["items"][0]["status"], "needs_manual_reconciliation",
        )
        self.assertEqual(resumed["items"][0]["action"], "update")
        self.assertIsNone(resumed["items"][0].get("operation_id"))
        self.assertEqual(MarketplaceOperation.query.count(), 0)

    def test_account_error_is_translated_and_foreign_product_is_rejected(self):
        with self.assertRaises(OzonBulkUploadError) as account_error:
            OzonBulkUploadService.create_run(
                seller_id=self.foreign_seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
            )
        self.assertEqual(
            account_error.exception.code,
            "marketplace_account_not_found",
        )

        foreign_product = ImportedProduct(
            seller_id=self.foreign_seller.id,
            external_id="foreign-source",
            source_type="synthetic",
            title="Чужой товар",
        )
        db.session.add(foreign_product)
        db.session.commit()
        with self.assertRaises(OzonBulkUploadValidationError):
            OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[foreign_product.id],
            )
        self.assertEqual(
            BackgroundJob.query.filter_by(
                seller_id=self.seller.id,
                job_type="ozon_bulk_upload",
            ).count(),
            0,
        )

    def test_expired_account_fails_before_creating_a_run(self):
        self.account.credential_expires_at = datetime.utcnow() - timedelta(
            seconds=1,
        )
        db.session.commit()

        with self.assertRaises(OzonBulkUploadError) as error:
            OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
            )

        self.assertEqual(
            error.exception.code,
            "ozon_account_credentials_expired",
        )
        self.assertEqual(
            BackgroundJob.query.filter_by(
                job_type=OzonBulkUploadService.JOB_TYPE,
            ).count(),
            0,
        )

    def test_selected_drafts_use_the_same_run_and_require_exact_account(self):
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            job = OzonBulkUploadService.create_run_from_drafts(
                seller_id=self.seller.id,
                account_id=self.account.id,
                draft_ids=[self.draft.id],
                created_by_user_id=self.user.id,
            )
        document = OzonBulkUploadService.public_document(job, detail=True)
        self.assertEqual(
            document["items"][0]["imported_product_id"],
            self.source.id,
        )

        with self.assertRaises(OzonBulkUploadValidationError):
            OzonBulkUploadService.create_run_from_drafts(
                seller_id=self.seller.id,
                account_id=self.account.id + 1000,
                draft_ids=[self.draft.id],
            )

    def test_reviewed_drafts_keep_exact_content_and_version_without_defaults(self):
        version = self.draft.version
        content = self.draft.content_json
        with patch.object(MarketplaceDraftService, "_build_validation_result",
                          return_value=self.validation_result()), \
                patch.object(MarketplaceDraftService, "create_draft",
                             side_effect=AssertionError("Must not prepare reviewed content")), \
                patch.object(MarketplaceDraftService, "apply_account_defaults",
                             side_effect=AssertionError("Must not change reviewed defaults")):
            job = OzonBulkUploadService.create_run_from_drafts(
                seller_id=self.seller.id, account_id=self.account.id,
                draft_ids=[self.draft.id], created_by_user_id=self.user.id,
                expected_versions={str(self.draft.id): version},
            )
        operation = MarketplaceOperation.query.one()
        self.assertEqual(operation.draft_version, version)
        self.assertEqual(operation.attempt_count, 0)
        self.assertEqual(self.draft.version, version)
        self.assertEqual(self.draft.content_json, content)
        document = OzonBulkUploadService.public_document(job, detail=True)
        self.assertEqual(document["items"][0]["status"], "queued")
        self.assertEqual(document["summary"]["source"], "drafts")

    def test_reviewed_drafts_reject_stale_or_incomplete_versions_before_job(self):
        for versions in [{}, {str(self.draft.id): True},
                         {str(self.draft.id): str(self.draft.version)},
                         {str(self.draft.id): self.draft.version + 1},
                         {str(self.draft.id): self.draft.version, "9999": 1}]:
            with self.subTest(versions=versions), self.assertRaises(OzonBulkUploadError):
                OzonBulkUploadService.create_run_from_drafts(
                    seller_id=self.seller.id, account_id=self.account.id,
                    draft_ids=[self.draft.id], created_by_user_id=self.user.id,
                    expected_versions=versions,
                )
        self.assertEqual(BackgroundJob.query.count(), 0)
        self.assertEqual(MarketplaceOperation.query.count(), 0)

    def test_reviewed_draft_changed_between_review_and_enqueue_is_not_submitted(self):
        enqueue = OzonBulkUploadService._enqueue_ready_items
        def changed(**kwargs):
            self.draft.content_json = json.dumps({"name": "Changed after review"})
            db.session.commit()
            return enqueue(**kwargs)
        with patch.object(OzonBulkUploadService, "_enqueue_ready_items", side_effect=changed):
            job = OzonBulkUploadService.create_run_from_drafts(
                seller_id=self.seller.id, account_id=self.account.id,
                draft_ids=[self.draft.id], created_by_user_id=self.user.id,
                expected_versions={str(self.draft.id): self.draft.version},
            )
        self.assertEqual(MarketplaceOperation.query.count(), 0)
        item = OzonBulkUploadService.public_document(job, detail=True)["items"][0]
        self.assertEqual(item["status"], "needs_input")
        self.assertIn("изменился", item["message"])

    def test_legacy_retry_requires_new_review_even_for_terminal_problem_items(self):
        invalid = {
            "publishable": False,
            "errors": [{
                "code": "missing_required_attribute",
                "field": "attributes.123",
                "message": "Заполните материал",
            }],
            "warnings": [],
            "schema": {"hash": "schema-hash", "version": 3},
        }
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=invalid,
        ):
            first = OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
                created_by_user_id=self.user.id,
            )

        with self.assertRaises(OzonBulkUploadError) as retry_error:
            OzonBulkUploadService.retry_run(
                seller_id=self.seller.id,
                job_uid=first.job_uid,
                request_key="u" * 24,
                created_by_user_id=self.user.id,
            )
        self.assertEqual(retry_error.exception.code, "draft_review_required")
        self.assertEqual(MarketplaceOperation.query.count(), 0)

    def test_legacy_category_review_does_not_authorize_implicit_retry_write(self):
        def product(index):
            original = {
                "external_id": f"category-retry-{index}",
                "vendor_code": f"category-retry-offer-{index}",
                "title": f"Товар для повтора {index}",
                "description": "Наблюдаемое описание товара",
                "category": "Категория",
                "barcodes": [f"46000000010{index:02d}"],
                "photo_urls": [f"https://img.test/category-retry-{index}.jpg"],
                "dimensions": {
                    "package_width_cm": 20,
                    "package_height_cm": 3,
                    "package_length_cm": 30,
                    "package_weight_g": 250,
                },
            }
            return ImportedProduct(
                seller_id=self.seller.id,
                external_id=original["external_id"],
                external_vendor_code=original["vendor_code"],
                source_type="synthetic",
                title=original["title"],
                description=original["description"],
                category=original["category"],
                original_data=json.dumps(original, ensure_ascii=False),
                photo_urls=json.dumps(original["photo_urls"]),
                barcodes=json.dumps(original["barcodes"]),
                calculated_price=1000,
                calculated_price_before_discount=1200,
            )

        representative = product(1)
        sibling = product(2)
        db.session.add_all([representative, sibling])
        db.session.commit()

        first = OzonBulkUploadService.create_run(
            seller_id=self.seller.id,
            account_id=self.account.id,
            imported_product_ids=[representative.id, sibling.id],
            created_by_user_id=self.user.id,
        )
        first_document = OzonBulkUploadService.public_document(
            first,
            detail=True,
        )
        self.assertEqual(
            [item["status"] for item in first_document["items"]],
            ["needs_input", "needs_input"],
        )

        representative_draft = MarketplaceDraftService.get_draft(
            seller_id=self.seller.id,
            draft_id=first_document["items"][0]["draft_id"],
        )
        self.app.secret_key = 'synthetic-category-review-test'
        review = MarketplaceDraftService.category_impact(
            seller_id=self.seller.id,
            draft_id=representative_draft.id,
            expected_version=representative_draft.version,
            target_product_type_id=self.product_type.id,
            save_mapping=True,
            actor_user_id=self.user.id,
        )
        confirmed = MarketplaceDraftService.update_draft(
            seller_id=self.seller.id,
            draft_id=representative_draft.id,
            expected_version=representative_draft.version,
            patch={
                "product_type_id": self.product_type.id,
                "save_mapping": True,
            },
            corrected_by_user_id=self.user.id,
            category_review_token=review['review_token'],
        )
        self.assertIsNotNone(confirmed.category_mapping_id)
        self.assertEqual(
            MarketplaceCategoryMapping.query.filter_by(
                seller_id=self.seller.id,
                mapping_status="active",
            ).count(),
            1,
        )

        with self.assertRaises(OzonBulkUploadError) as retry_error:
            OzonBulkUploadService.retry_run(
                seller_id=self.seller.id,
                job_uid=first.job_uid,
                request_key="v" * 24,
                created_by_user_id=self.user.id,
            )
        self.assertEqual(retry_error.exception.code, "draft_review_required")
        self.assertEqual(MarketplaceOperation.query.count(), 0)


class OzonBulkRepairServiceTest(OzonPublicationFixture, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self.app.config.update(
            MARKETPLACE_OZON_ENABLED=True,
            MARKETPLACE_OZON_PUBLICATION_ENABLED=True,
        )
        self.account._credentials_encrypted = "synthetic-encrypted-value"
        self.account.settings_json = json.dumps({"default_vat": "0.22"})
        now = datetime.utcnow()
        self.tnved = MarketplaceAttributeDefinition(
            marketplace_id=self.marketplace.id,
            product_type_id=self.product_type.id,
            external_attribute_id="22232",
            name="Код ТН ВЭД ЕАЭС",
            data_type="String",
            is_required=True,
            dictionary_id="tnved-dictionary",
            max_value_count=1,
            is_collection=False,
            is_available=True,
            is_enabled=True,
            last_seen_at=now,
            values_synced_at=now,
            values_sync_status="success",
            values_snapshot_hash="tnved-values-hash",
            values_version=1,
            values_count=1,
        )
        db.session.add(self.tnved)
        db.session.flush()
        self.tnved_value = MarketplaceAttributeValue(
            marketplace_id=self.marketplace.id,
            product_type_id=self.product_type.id,
            attribute_id=self.tnved.id,
            external_value_id="1001",
            value="3304990000",
            value_normalized=OzonReferenceService.normalize_value(
                "3304990000"
            ),
            is_available=True,
            last_seen_at=now,
        )
        self.product_type.attributes_count = 2
        self.product_type.required_attributes_count = 1
        self.draft.status = "blocked"
        self.draft.validation_status = "invalid"
        self.draft.validation_result_json = json.dumps({
            "publishable": False,
            "errors": [{
                "code": "required_attribute_missing",
                "field": "attributes.22232",
                "message": "Обязательный атрибут «Код ТН ВЭД ЕАЭС» не заполнен",
            }],
            "warnings": [],
        }, ensure_ascii=False)
        db.session.add(self.tnved_value)
        db.session.commit()
        self.job = self._repair_job()

    def _repair_job(self):
        now = datetime.utcnow().isoformat()
        document = {
            "version": OzonBulkUploadService.DOCUMENT_VERSION,
            "source": "products",
            "account_id": self.account.id,
            "account_label": self.account.label,
            "created_at": now,
            "updated_at": now,
            "items": [{
                "imported_product_id": self.source.id,
                "draft_id": self.draft.id,
                "offer_id": self.draft.offer_id,
                "title": self.source.title,
                "action": "create",
                "status": "needs_input",
                "code": "required_attribute_missing",
                "message": "Заполните ТН ВЭД",
                "validation_errors": [{
                    "code": "required_attribute_missing",
                    "field": "attributes.22232",
                    "message": "Заполните ТН ВЭД",
                }],
                "updated_at": now,
            }],
        }
        job = BackgroundJob(
            job_uid="ozon-upload-" + "a" * 32,
            seller_id=self.seller.id,
            job_type=OzonBulkUploadService.JOB_TYPE,
            status="completed",
            total=1,
        )
        OzonBulkUploadService._store_progress(job, document)
        db.session.add(job)
        db.session.flush()
        OzonBulkUploadService._persist(job, document)
        return job

    @staticmethod
    def _workbook_cell(workbook, key):
        sheet = workbook[OzonBulkRepairService.CARD_SHEET]
        keys = [cell.value for cell in sheet[1]]
        return sheet.cell(row=3, column=keys.index(key) + 1)

    @staticmethod
    def _payload(workbook):
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _ready_validation(self, draft):
        attributes = json.loads(draft.attributes_json or "[]")
        present = any(
            item.get("attribute_id") == "22232"
            and item.get("values") == [{
                "dictionary_value_id": "1001",
                "value": "3304990000",
            }]
            for item in attributes
            if isinstance(item, dict)
        )
        return {
            "publishable": present,
            "errors": [] if present else [{
                "code": "required_attribute_missing",
                "field": "attributes.22232",
                "message": "Заполните ТН ВЭД",
            }],
            "warnings": [],
            "schema": {"hash": "schema-hash", "version": 3},
        }

    def test_xlsx_round_trip_resolves_official_value_without_provider_write(self):
        filename, payload = OzonBulkRepairService.export_workbook(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
        )
        self.assertTrue(filename.endswith(".xlsx"))
        workbook = load_workbook(BytesIO(payload))
        self._workbook_cell(
            workbook,
            "attribute:22232",
        ).value = "3304990000"

        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            side_effect=self._ready_validation,
        ):
            report = OzonBulkRepairService.import_workbook(
                seller_id=self.seller.id,
                job_uid=self.job.job_uid,
                payload=self._payload(workbook),
                corrected_by_user_id=self.user.id,
            )

        self.assertEqual(report["updated"], 1)
        self.assertEqual(report["ready_to_retry"], 1)
        self.assertEqual(report["failed"], 0)
        self.assertEqual(MarketplaceOperation.query.count(), 0)
        draft = db.session.get(MarketplaceProductDraft, self.draft.id)
        self.assertEqual(json.loads(draft.attributes_json), [{
            "attribute_id": "22232",
            "complex_id": "0",
            "values": [{
                "dictionary_value_id": "1001",
                "value": "3304990000",
            }],
        }])
        result = OzonBulkUploadService.public_document(
            OzonBulkUploadService.get_run(
                seller_id=self.seller.id,
                job_uid=self.job.job_uid,
                reconcile=False,
            ),
            detail=True,
        )
        self.assertEqual(result["items"][0]["status"], "ready_to_retry")
        self.assertEqual(result["summary"]["ready_to_retry"], 1)

    def test_platform_editor_applies_exact_dictionary_value_locally(self):
        editor = OzonBulkRepairService.editor_document(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
        )
        row = editor["groups"][0]["rows"][0]
        self.assertEqual(
            row["attributes"][0]["external_id"],
            "22232",
        )
        self.assertTrue(row["attributes"][0]["dictionary"])

        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            side_effect=self._ready_validation,
        ):
            report = OzonBulkRepairService.apply_editor_rows(
                seller_id=self.seller.id,
                job_uid=self.job.job_uid,
                rows=[{
                    "draft_id": row["draft_id"],
                    "expected_version": row["draft_version"],
                    "imported_product_id": row["imported_product_id"],
                    "action": OzonBulkRepairService.ACTION_REPAIR,
                    "product_type_id": row["product_type_id"],
                    "save_mapping": False,
                    "values": {
                        "attribute:22232": "3304990000",
                        "attribute_value_id:22232": "1001",
                    },
                }],
                corrected_by_user_id=self.user.id,
            )

        self.assertEqual(report["ready_to_retry"], 1)
        self.assertEqual(report["failed"], 0)
        self.assertEqual(MarketplaceOperation.query.count(), 0)
        draft = db.session.get(MarketplaceProductDraft, self.draft.id)
        self.assertEqual(json.loads(draft.attributes_json), [{
            "attribute_id": "22232",
            "complex_id": "0",
            "values": [{
                "dictionary_value_id": "1001",
                "value": "3304990000",
            }],
        }])

    def test_platform_editor_exposes_tnved_hint_without_selecting_it(self):
        self.product_type.name = "Насадки, удлинители эротические"
        self.product_type.category.name = "Товары для взрослых"
        self.product_type.category.full_path = "Товары для взрослых"
        self.draft.source_facts_json = json.dumps({
            "version": 2,
            "facts": {
                "identity": {
                    "source_title": "Насадка на член с вибрацией",
                    "source_category": (
                        "Насадки и кольца > "
                        "Удлиняющие и расширяющие насадки"
                    ),
                    "source_categories": [
                        "Насадки и кольца > С вибрацией",
                    ],
                },
                "attributes": {
                    "materials": ["Эластичный TPR"],
                },
            },
        }, ensure_ascii=False)
        suggestion = MarketplaceAttributeValue(
            marketplace_id=self.marketplace.id,
            product_type_id=self.product_type.id,
            attribute_id=self.tnved.id,
            external_value_id="vibration-tnved",
            value=(
                "9019101000 - "
                "Аппараты электрические вибромассажные"
            ),
            value_normalized=OzonReferenceService.normalize_value(
                "9019101000 - Аппараты электрические вибромассажные"
            ),
            is_available=True,
            last_seen_at=datetime.utcnow(),
        )
        self.tnved.values_count = 2
        db.session.add(suggestion)
        db.session.commit()

        editor = OzonBulkRepairService.editor_document(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
        )

        field = next(
            item
            for item in editor["groups"][0]["rows"][0]["attributes"]
            if item["external_id"] == "22232"
        )
        self.assertEqual(field["value"], "")
        self.assertEqual(field["dictionary_value_id"], "")
        self.assertEqual(field["suggestions"][0]["code"], "9019101000")
        self.assertEqual(
            field["suggestions"][0]["dictionary_value_id"],
            "vibration-tnved",
        )
        self.assertEqual(MarketplaceOperation.query.count(), 0)

    def test_platform_editor_requires_explicit_legacy_attribute_cleanup(self):
        prior = self.prior_payload()
        prior["items"][0]["attributes"].append({
            "id": 99999,
            "complex_id": 0,
            "values": [{"value": "Устаревшее значение"}],
        })
        self.attach_listing(prior)
        self.draft.status = "blocked"
        self.draft.validation_status = "invalid"
        self.draft.validation_result_json = json.dumps({
            "publishable": False,
            "errors": [
                {
                    "code": "unknown_attribute",
                    "field": "attributes[0]",
                    "message": (
                        "Атрибут отсутствует в текущей Ozon schema"
                    ),
                },
                {
                    "code": "required_attribute_missing",
                    "field": "attributes.22232",
                    "message": "Заполните ТН ВЭД",
                },
            ],
            "warnings": [],
        }, ensure_ascii=False)
        document = OzonBulkUploadService._load_progress(self.job)
        document["items"][0]["action"] = "update"
        OzonBulkUploadService._persist(self.job, document)

        editor = OzonBulkRepairService.editor_document(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
        )
        row = editor["groups"][0]["rows"][0]
        self.assertEqual(
            [
                candidate["attribute_id"]
                for candidate in row["cleanup_candidates"]
            ],
            ["99999"],
        )
        self.assertFalse(row["schema_cleanup"])
        self.assertGreaterEqual(row["validation_error_count"], 2)
        self.assertIn(
            "unknown_attribute",
            {
                item["code"]
                for item in row["validation_errors"]
            },
        )
        self.assertIn(
            "required_attribute_missing",
            {
                item["code"]
                for item in row["validation_errors"]
            },
        )

        current_validator = (
            MarketplaceDraftService._build_validation_result
        )

        def validation_after_edit(current_draft):
            if any(
                item.get("attribute_id") == "22232"
                for item in json.loads(
                    current_draft.attributes_json or "[]"
                )
                if isinstance(item, dict)
            ):
                return self._ready_validation(current_draft)
            return current_validator(current_draft)

        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            side_effect=validation_after_edit,
        ):
            report = OzonBulkRepairService.apply_editor_rows(
                seller_id=self.seller.id,
                job_uid=self.job.job_uid,
                rows=[{
                    "draft_id": row["draft_id"],
                    "expected_version": row["draft_version"],
                    "imported_product_id": row["imported_product_id"],
                    "action": OzonBulkRepairService.ACTION_REPAIR,
                    "product_type_id": row["product_type_id"],
                    "save_mapping": False,
                    "schema_cleanup": True,
                    "values": {
                        "attribute:22232": "3304990000",
                        "attribute_value_id:22232": "1001",
                    },
                }],
                corrected_by_user_id=self.user.id,
            )

        self.assertEqual(report["ready_to_retry"], 1)
        self.assertEqual(report["failed"], 0)
        self.assertEqual(MarketplaceOperation.query.count(), 0)
        draft = db.session.get(MarketplaceProductDraft, self.draft.id)
        self.assertEqual(json.loads(draft.attribute_removals_json), [{
            "attribute_id": "99999",
            "complex_id": "0",
        }])
        effective, baseline = MarketplaceDraftService.publication_documents(
            draft
        )
        self.assertIsNotNone(baseline)
        self.assertNotIn(
            "99999",
            {
                item.get("attribute_id")
                for item in effective["attributes"]
            },
        )

    def test_platform_editor_does_not_use_stale_cleanup_error_paths(self):
        prior = self.prior_payload()
        prior["items"][0]["attributes"].append({
            "id": 99999,
            "complex_id": 0,
            "values": [{"value": "Устаревшее значение"}],
        })
        self.attach_listing(prior)
        self.draft.status = "blocked"
        self.draft.validation_status = "invalid"
        self.draft.validation_result_json = json.dumps({
            "publishable": False,
            "errors": [{
                "code": "unknown_attribute",
                "field": "attributes[0]",
                "message": "Старая позиционная ошибка",
            }],
            "warnings": [],
        }, ensure_ascii=False)
        document = OzonBulkUploadService._load_progress(self.job)
        document["items"][0]["action"] = "update"
        OzonBulkUploadService._persist(self.job, document)
        db.session.commit()

        current_validation = {
            "publishable": False,
            "errors": [{
                "code": "price_required",
                "field": "commercial.price",
                "message": "Текущая локальная причина",
            }],
            "warnings": [],
        }
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=current_validation,
        ):
            editor = OzonBulkRepairService.editor_document(
                seller_id=self.seller.id,
                job_uid=self.job.job_uid,
            )

        row = editor["groups"][0]["rows"][0]
        self.assertEqual(row["cleanup_candidates"], [])
        self.assertEqual(row["validation_error_count"], 1)
        self.assertEqual(
            row["validation_errors"][0]["code"],
            "price_required",
        )

    def test_platform_editor_can_bind_one_type_and_save_category_mapping(self):
        self.draft.product_type_id = None
        self.draft.external_category_id = None
        self.draft.external_type_id = None
        self.draft.attributes_json = "[]"
        self.draft.status = "needs_category"
        self.draft.validation_status = "invalid"
        self.draft.validation_result_json = json.dumps({
            "publishable": False,
            "errors": [{
                "code": "product_type_required",
                "field": "product_type_id",
                "message": "Выберите тип товара Ozon",
            }],
            "warnings": [],
        }, ensure_ascii=False)
        db.session.commit()
        editor = OzonBulkRepairService.editor_document(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
        )
        row = editor["groups"][0]["rows"][0]

        report = OzonBulkRepairService.apply_editor_rows(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
            rows=[{
                "draft_id": row["draft_id"],
                "expected_version": row["draft_version"],
                "imported_product_id": row["imported_product_id"],
                "action": OzonBulkRepairService.ACTION_REPAIR,
                "product_type_id": self.product_type.id,
                "save_mapping": True,
                "values": {},
            }],
            corrected_by_user_id=self.user.id,
        )

        self.assertEqual(report["updated"], 1)
        self.assertEqual(report["failed"], 0)
        draft = db.session.get(MarketplaceProductDraft, self.draft.id)
        self.assertEqual(draft.product_type_id, self.product_type.id)
        mapping = db.session.get(
            MarketplaceCategoryMapping,
            draft.category_mapping_id,
        )
        self.assertEqual(mapping.mapping_source, "manual")
        self.assertEqual(mapping.corrected_by_user_id, self.user.id)
        self.assertEqual(MarketplaceOperation.query.count(), 0)

    def test_platform_dictionary_search_is_exact_scoped_and_escapes_wildcards(self):
        result = OzonBulkRepairService.search_dictionary_values(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
            draft_id=self.draft.id,
            external_attribute_id="22232",
            query="3304",
        )
        self.assertEqual(result["items"], [{
            "id": "1001",
            "value": "3304990000",
        }])
        escaped = OzonBulkRepairService.search_dictionary_values(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
            draft_id=self.draft.id,
            external_attribute_id="22232",
            query="%",
        )
        self.assertEqual(escaped["items"], [])
        with self.assertRaises(OzonBulkUploadError):
            OzonBulkRepairService.search_dictionary_values(
                seller_id=self.foreign_seller.id,
                job_uid=self.job.job_uid,
                draft_id=self.draft.id,
                external_attribute_id="22232",
                query="3304",
            )

    def test_forbidden_brand_defaults_to_recoverable_run_exclusion(self):
        self.draft.validation_result_json = json.dumps({
            "publishable": False,
            "errors": [{
                "code": "ozon_brand_forbidden",
                "field": "brand",
                "message": "Бренд запрещён",
            }],
            "warnings": [],
        }, ensure_ascii=False)
        db.session.commit()
        _filename, payload = OzonBulkRepairService.export_workbook(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
        )
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(
            self._workbook_cell(workbook, "action").value,
            OzonBulkRepairService.ACTION_EXCLUDE,
        )
        report = OzonBulkRepairService.import_workbook(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
            payload=self._payload(workbook),
            corrected_by_user_id=self.user.id,
        )
        self.assertEqual(report["excluded"], 1)
        self.assertIsNotNone(
            db.session.get(MarketplaceProductDraft, self.draft.id)
        )
        document = OzonBulkUploadService.public_document(
            OzonBulkUploadService.get_run(
                seller_id=self.seller.id,
                job_uid=self.job.job_uid,
                reconcile=False,
            ),
            detail=True,
        )
        self.assertEqual(document["items"][0]["status"], "excluded")
        self.assertEqual(document["summary"]["excluded"], 1)
        with self.assertRaises(OzonBulkUploadError) as retry_error:
            OzonBulkUploadService.retry_run(
                seller_id=self.seller.id,
                job_uid=self.job.job_uid,
                request_key="forbidden-brand-retry-0001",
                created_by_user_id=self.user.id,
            )
        self.assertEqual(retry_error.exception.code, "draft_review_required")

    def test_stale_row_does_not_overwrite_newer_draft(self):
        _filename, payload = OzonBulkRepairService.export_workbook(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
        )
        workbook = load_workbook(BytesIO(payload))
        self._workbook_cell(
            workbook,
            "attribute:22232",
        ).value = "3304990000"
        self.draft.content_json = json.dumps({
            "name": "Новая версия",
            "description": "Наблюдаемое описание",
        }, ensure_ascii=False)
        db.session.commit()

        report = OzonBulkRepairService.import_workbook(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
            payload=self._payload(workbook),
            corrected_by_user_id=self.user.id,
        )
        self.assertEqual(report["failed"], 1)
        self.assertEqual(report["updated"], 0)
        self.assertEqual(
            json.loads(
                db.session.get(
                    MarketplaceProductDraft,
                    self.draft.id,
                ).attributes_json
            ),
            [],
        )

    def test_formula_and_cross_tenant_workbooks_fail_closed(self):
        _filename, payload = OzonBulkRepairService.export_workbook(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
        )
        workbook = load_workbook(BytesIO(payload))
        self._workbook_cell(workbook, "price_rub").value = "=1+1"
        with self.assertRaises(OzonBulkUploadValidationError):
            OzonBulkRepairService.import_workbook(
                seller_id=self.seller.id,
                job_uid=self.job.job_uid,
                payload=self._payload(workbook),
            )
        with self.assertRaises(OzonBulkUploadError):
            OzonBulkRepairService.export_workbook(
                seller_id=self.foreign_seller.id,
                job_uid=self.job.job_uid,
            )
        self.assertEqual(MarketplaceOperation.query.count(), 0)

    def test_formula_like_observed_text_round_trips_without_mutation(self):
        observed_description = "=это текст карточки, а не формула"
        content = json.loads(self.draft.content_json or "{}")
        content["description"] = observed_description
        self.draft.content_json = json.dumps(
            content,
            ensure_ascii=False,
        )
        db.session.commit()
        _filename, payload = OzonBulkRepairService.export_workbook(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
        )
        workbook = load_workbook(BytesIO(payload), data_only=False)
        description_cell = self._workbook_cell(
            workbook,
            "description",
        )
        self.assertEqual(description_cell.value, observed_description)
        self.assertEqual(description_cell.data_type, "s")

        report = OzonBulkRepairService.import_workbook(
            seller_id=self.seller.id,
            job_uid=self.job.job_uid,
            payload=self._payload(workbook),
            corrected_by_user_id=self.user.id,
        )

        self.assertEqual(report["failed"], 0)
        refreshed = db.session.get(
            MarketplaceProductDraft,
            self.draft.id,
        )
        self.assertEqual(
            json.loads(refreshed.content_json)["description"],
            observed_description,
        )
        self.assertEqual(MarketplaceOperation.query.count(), 0)


class OzonBulkUploadRoutesTest(OzonPublicationFixture, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self.app.config.update(
            SECRET_KEY="ozon-bulk-upload-routes",
            MARKETPLACE_OZON_ENABLED=True,
            MARKETPLACE_OZON_PUBLICATION_ENABLED=True,
        )
        self.account._credentials_encrypted = "synthetic-encrypted-value"
        self.account.settings_json = json.dumps({"default_vat": "0.22"})
        db.session.commit()
        register_ozon_bulk_upload_routes(self.app)
        self.client = self.app.test_client()

    @staticmethod
    def _user(seller, user_id):
        return SimpleNamespace(
            id=user_id,
            seller=SimpleNamespace(id=seller.id),
            is_authenticated=True,
            is_active=True,
        )

    def _auth_patches(self, seller, user_id):
        user = self._user(seller, user_id)
        return (
            patch("routes.ozon_bulk_uploads.current_user", user),
            patch("flask_login.utils._get_user", return_value=user),
        )

    def test_history_html_renders_without_live_account_lookup(self):
        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch, patch(
            "routes.ozon_bulk_uploads.render_template",
            return_value="history",
        ) as render:
            response = self.client.get("/marketplaces/ozon/uploads/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_data(as_text=True), "history")
        render.assert_called_once()
        self.assertEqual(render.call_args.args[0], "ozon_bulk_uploads.html")

    def test_detail_template_uses_dict_key_for_items_not_dict_method(self):
        source = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "ozon_bulk_upload_detail.html"
        ).read_text(encoding="utf-8")

        self.assertIn("{% for item in run['items'] %}", source)
        self.assertNotIn("{% for item in run.items %}", source)

    def test_json_create_and_detail_share_one_durable_run(self):
        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch, patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            created = self.client.post(
                "/marketplaces/ozon/uploads/",
                json={
                    "account_id": self.account.id,
                    "imported_product_ids": [self.source.id],
                    "confirm_prepare": True,
                    "request_key": "route-prepare-exact-0123456789abcdef",
                },
            )
        self.assertEqual(created.status_code, 202)
        body = created.get_json()
        self.assertTrue(body["success"])
        job_uid = body["run"]["job_uid"]
        self.assertEqual(body["run"]["mode"], "source_prepare")
        self.assertEqual(body["run"]["items"][0]["status"], "pending")
        self.assertEqual(MarketplaceOperation.query.count(), 0)

        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch:
            detail = self.client.get(
                f"/marketplaces/ozon/uploads/api/{job_uid}",
            )
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.get_json()["run"]["job_uid"], job_uid)

    def test_run_detail_is_tenant_scoped(self):
        job = self._create_for_route_test()
        user_patch, login_patch = self._auth_patches(
            self.foreign_seller,
            self.foreign_user.id,
        )
        with user_patch, login_patch:
            response = self.client.get(
                f"/marketplaces/ozon/uploads/api/{job.job_uid}",
            )
        self.assertEqual(response.status_code, 404)
        self.assertEqual(
            response.get_json()["code"],
            "ozon_bulk_upload_not_found",
        )

    def test_ready_drafts_enter_the_same_upload_result(self):
        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch, patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            response = self.client.post(
                "/marketplaces/ozon/uploads/from-drafts",
                json={
                    "account_id": self.account.id,
                    "draft_ids": [self.draft.id],
                    "confirm_write": True,
                    "expected_versions": {str(self.draft.id): self.draft.version},
                    "request_key": "route-review-exact-0123456789abcdef",
                },
            )
        self.assertEqual(response.status_code, 202)
        run = response.get_json()["run"]
        self.assertEqual(
            run["items"][0]["imported_product_id"],
            self.source.id,
        )
        self.assertEqual(run["mode"], "reviewed_drafts")
        self.assertEqual(run["items"][0]["status"], "reviewed")
        self.assertEqual(MarketplaceOperation.query.count(), 0)

    def test_reviewed_bulk_route_rejects_changed_version_and_null_contract(self):
        user_patch, login_patch = self._auth_patches(self.seller, self.user.id)
        with user_patch, login_patch:
            for versions, expected_status in [
                (None, 409), ({}, 409),
                ({str(self.draft.id): self.draft.version + 1}, 409),
            ]:
                with self.subTest(versions=versions):
                    response = self.client.post(
                        "/marketplaces/ozon/uploads/from-drafts",
                        json={"account_id": self.account.id,
                              "draft_ids": [self.draft.id], "confirm_write": True,
                              "expected_versions": versions,
                              "request_key": "route-review-conflict-0123456789abcdef"},
                    )
                    self.assertEqual(response.status_code, expected_status)
        self.assertEqual(BackgroundJob.query.count(), 0)
        self.assertEqual(MarketplaceOperation.query.count(), 0)

    def test_repair_xlsx_route_is_local_only_and_needs_no_write_confirmation(self):
        invalid = {
            "publishable": False,
            "errors": [{
                "code": "ozon_brand_forbidden",
                "field": "brand",
                "message": "Бренд запрещён для Ozon",
            }],
            "warnings": [],
            "schema": {"hash": "schema-hash", "version": 3},
        }
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=invalid,
        ):
            job = OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
                created_by_user_id=self.user.id,
            )
        self.assertEqual(MarketplaceOperation.query.count(), 0)

        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch:
            exported = self.client.get(
                f"/marketplaces/ozon/uploads/{job.job_uid}/repair.xlsx",
            )
        self.assertEqual(exported.status_code, 200)
        self.assertIn(
            "spreadsheetml.sheet",
            exported.headers["Content-Type"],
        )
        self.assertEqual(
            exported.headers["Cache-Control"],
            "private, no-store, max-age=0",
        )

        workbook = load_workbook(BytesIO(exported.data))
        sheet = workbook[OzonBulkRepairService.CARD_SHEET]
        keys = [cell.value for cell in sheet[1]]
        sheet.cell(
            row=3,
            column=keys.index("action") + 1,
        ).value = OzonBulkRepairService.ACTION_EXCLUDE
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)

        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch:
            imported = self.client.post(
                f"/marketplaces/ozon/uploads/{job.job_uid}/repair",
                data={
                    "repair_file": (
                        payload,
                        "repair.xlsx",
                    ),
                },
                content_type="multipart/form-data",
            )
        self.assertEqual(imported.status_code, 302)
        self.assertEqual(MarketplaceOperation.query.count(), 0)
        document = OzonBulkUploadService.public_document(
            OzonBulkUploadService.get_run(
                seller_id=self.seller.id,
                job_uid=job.job_uid,
                reconcile=False,
            ),
            detail=True,
        )
        self.assertEqual(document["items"][0]["status"], "excluded")

    def test_platform_mass_editor_renders_and_applies_without_provider_write(self):
        invalid = {
            "publishable": False,
            "errors": [{
                "code": "ozon_brand_forbidden",
                "field": "brand",
                "message": "Бренд запрещён для Ozon",
            }],
            "warnings": [],
            "schema": {"hash": "schema-hash", "version": 3},
        }
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=invalid,
        ):
            job = OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
                created_by_user_id=self.user.id,
            )
        self.assertEqual(job.status, "completed")
        self.assertEqual(MarketplaceOperation.query.count(), 0)

        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch, patch(
            "routes.ozon_bulk_uploads.render_template",
            return_value="platform-editor",
        ) as render:
            opened = self.client.get(
                f"/marketplaces/ozon/uploads/{job.job_uid}/repair",
            )
        self.assertEqual(opened.status_code, 200)
        self.assertEqual(opened.get_data(as_text=True), "platform-editor")
        self.assertEqual(
            render.call_args.args[0],
            "ozon_bulk_repair.html",
        )
        editor = render.call_args.kwargs["editor"]
        row = editor["groups"][0]["rows"][0]

        prefix = f"row_{row['draft_id']}_"
        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch:
            applied = self.client.post(
                (
                    f"/marketplaces/ozon/uploads/{job.job_uid}"
                    "/repair/apply"
                ),
                data={
                    "selected_draft_id": str(row["draft_id"]),
                    prefix + "draft_version": str(
                        row["draft_version"]
                    ),
                    prefix + "imported_product_id": str(
                        row["imported_product_id"]
                    ),
                    prefix + "action": (
                        OzonBulkRepairService.ACTION_EXCLUDE
                    ),
                    prefix + "product_type_id": str(
                        row["product_type_id"]
                    ),
                    prefix + "save_mapping": "0",
                    prefix + "price_rub": row["price_rub"],
                    prefix + "package_width_mm": (
                        row["package_width_mm"]
                    ),
                    prefix + "package_height_mm": (
                        row["package_height_mm"]
                    ),
                    prefix + "package_depth_mm": (
                        row["package_depth_mm"]
                    ),
                    prefix + "package_weight_g": (
                        row["package_weight_g"]
                    ),
                    prefix + "description": row["description"],
                },
            )
        self.assertEqual(applied.status_code, 302)
        self.assertIn(
            f"/marketplaces/ozon/uploads/{job.job_uid}/repair",
            applied.headers["Location"],
        )
        self.assertEqual(MarketplaceOperation.query.count(), 0)
        document = OzonBulkUploadService.public_document(
            OzonBulkUploadService.get_run(
                seller_id=self.seller.id,
                job_uid=job.job_uid,
                reconcile=False,
            ),
            detail=True,
        )
        self.assertEqual(document["items"][0]["status"], "excluded")

    def test_platform_editor_is_tenant_scoped(self):
        invalid = {
            "publishable": False,
            "errors": [{
                "code": "price_required",
                "field": "commercial.price",
                "message": "Укажите цену",
            }],
            "warnings": [],
            "schema": {"hash": "schema-hash", "version": 3},
        }
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=invalid,
        ):
            job = OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
                created_by_user_id=self.user.id,
            )
        user_patch, login_patch = self._auth_patches(
            self.foreign_seller,
            self.foreign_user.id,
        )
        with user_patch, login_patch:
            response = self.client.get(
                f"/marketplaces/ozon/uploads/{job.job_uid}/repair",
                headers={"Accept": "application/json"},
            )
        self.assertEqual(response.status_code, 404)
        self.assertEqual(
            response.get_json()["code"],
            "ozon_bulk_upload_not_found",
        )

    def test_http_rejects_oversized_product_and_draft_lists_before_service(self):
        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch, patch.object(
            OzonBulkUploadService,
            "accept_source_prepare",
        ) as create_run:
            response = self.client.post(
                "/marketplaces/ozon/uploads/",
                json={
                    "account_id": self.account.id,
                    "confirm_prepare": True,
                    "request_key": "route-oversized-0123456789abcdef",
                    "imported_product_ids": list(range(
                        1,
                        OzonBulkUploadService.MAX_ITEMS + 2,
                    )),
                },
            )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            response.get_json()["code"],
            "ozon_bulk_upload_invalid",
        )
        create_run.assert_not_called()

        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch, patch.object(
            OzonBulkUploadService,
            "accept_reviewed_publish",
        ) as create_from_drafts:
            response = self.client.post(
                "/marketplaces/ozon/uploads/from-drafts",
                json={
                    "account_id": self.account.id,
                    "confirm_write": True,
                    "draft_ids": list(range(
                        1,
                        OzonBulkUploadService.MAX_ITEMS + 2,
                    )),
                },
            )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            response.get_json()["code"],
            "ozon_bulk_upload_invalid",
        )
        create_from_drafts.assert_not_called()

    def test_disabled_publication_fails_before_creating_a_job(self):
        self.app.config["MARKETPLACE_OZON_PUBLICATION_ENABLED"] = False
        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch:
            response = self.client.post(
                "/marketplaces/ozon/uploads/from-drafts",
                json={
                    "account_id": self.account.id,
                    "draft_ids": [self.draft.id],
                    "expected_versions": {str(self.draft.id): self.draft.version},
                    "request_key": "route-disabled-0123456789abcdef",
                    "confirm_write": True,
                },
            )
        self.assertEqual(response.status_code, 409)
        self.assertEqual(
            response.get_json()["code"],
            "ozon_publication_disabled",
        )
        self.assertEqual(
            BackgroundJob.query.filter_by(
                job_type="ozon_bulk_upload",
            ).count(),
            0,
        )

    def test_json_preparation_requires_strict_explicit_confirmation(self):
        user_patch, login_patch = self._auth_patches(
            self.seller,
            self.user.id,
        )
        with user_patch, login_patch, patch.object(
            OzonBulkUploadService,
            "accept_source_prepare",
        ) as create_run:
            missing = self.client.post(
                "/marketplaces/ozon/uploads/",
                json={
                    "account_id": self.account.id,
                    "imported_product_ids": [self.source.id],
                    "request_key": "route-confirm-0123456789abcdef",
                },
            )
            string_value = self.client.post(
                "/marketplaces/ozon/uploads/",
                json={
                    "account_id": self.account.id,
                    "imported_product_ids": [self.source.id],
                    "request_key": "route-confirm-0123456789abcdef",
                    "confirm_prepare": "true",
                },
            )

        self.assertEqual(missing.status_code, 400)
        self.assertEqual(string_value.status_code, 400)
        self.assertEqual(
            missing.get_json()["code"],
            "ozon_bulk_upload_invalid",
        )
        create_run.assert_not_called()

    def _create_for_route_test(self):
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            return OzonBulkUploadService.create_run(
                seller_id=self.seller.id,
                account_id=self.account.id,
                imported_product_ids=[self.source.id],
                created_by_user_id=self.user.id,
            )


class MarketplacePublicationServiceTest(OzonPublicationFixture, unittest.TestCase):
    def test_product_write_freezes_media_before_any_ozon_call(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.app.config.update(
            SECRET_KEY="publication-media-secret",
            PUBLIC_BASE_URL="https://seller.test",
            MARKETPLACE_IMAGE_ASSET_DIR=temporary.name,
            MARKETPLACE_IMAGE_ASSETS_TEST_ENFORCE=True,
            OZON_MEDIA_ASSET_URLS_PER_ATTEMPT=10,
        )
        image = BytesIO()
        from PIL import Image

        Image.new("RGB", (800, 1000), (20, 40, 60)).save(
            image,
            format="JPEG",
        )
        adapter = SyntheticPublicationAdapter()
        with patch(
            "services.marketplace_image_assets.download_public_image",
            return_value=image.getvalue(),
        ):
            operation = self.start(
                adapter,
                key="publication-key-media-assets",
            )

        self.assertEqual(operation.status, "submitted")
        self.assertEqual(operation.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)
        submitted_url = adapter.submitted_payloads[0]["items"][0]["images"][0]
        self.assertIn(
            "https://seller.test/marketplace-assets/images/",
            submitted_url,
        )
        self.assertNotEqual(submitted_url, "https://img.test/product.jpg")
        summary = json.loads(operation.request_summary_json)
        self.assertEqual(summary["media_asset_state"], "ready")
        self.assertEqual(summary["media_asset_prepared"], 1)
        self.assertEqual(summary["media_asset_source_total"], 1)
        self.assertEqual(
            operation.request_fingerprint,
            operation.snapshot.submitted_fingerprint,
        )
        self.assertEqual(
            OzonProductImportContract.fingerprint(
                json.loads(operation.snapshot.submitted_state_json)
            ),
            operation.request_fingerprint,
        )
        self.assertEqual(
            len(list(Path(temporary.name).rglob("*.jpg"))),
            1,
        )

    def test_invalid_media_fails_with_zero_provider_attempts(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.app.config.update(
            SECRET_KEY="publication-media-secret",
            PUBLIC_BASE_URL="https://seller.test",
            MARKETPLACE_IMAGE_ASSET_DIR=temporary.name,
            MARKETPLACE_IMAGE_ASSETS_TEST_ENFORCE=True,
            OZON_MEDIA_ASSET_URLS_PER_ATTEMPT=10,
        )
        adapter = SyntheticPublicationAdapter()
        with patch(
            "services.marketplace_image_assets.download_public_image",
            side_effect=ImageLabError(
                "Источник фото вернул HTML без безопасного redirect"
            ),
        ):
            operation = self.start(
                adapter,
                key="publication-key-invalid-media",
            )

        self.assertEqual(operation.status, "failed")
        self.assertEqual(operation.attempt_count, 0)
        self.assertEqual(operation.error_code, "media_source_invalid_image")
        self.assertEqual(adapter.list_calls, [])
        self.assertEqual(adapter.submitted_payloads, [])

    def test_media_preparation_resumes_from_committed_partial_snapshot(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.app.config.update(
            SECRET_KEY="publication-media-secret",
            PUBLIC_BASE_URL="https://seller.test",
            MARKETPLACE_IMAGE_ASSET_DIR=temporary.name,
            MARKETPLACE_IMAGE_ASSETS_TEST_ENFORCE=True,
            OZON_MEDIA_ASSET_URLS_PER_ATTEMPT=2,
        )
        self.draft.media_json = json.dumps({
            "images": [
                f"https://img.test/product-{index}.jpg"
                for index in range(4)
            ],
        })
        db.session.commit()
        self.expected_version = self.draft.version

        images = []
        from PIL import Image

        for index in range(4):
            output = BytesIO()
            Image.new(
                "RGB",
                (800, 1000),
                (20 + index * 30, 40, 60),
            ).save(output, format="JPEG")
            images.append(output.getvalue())
        adapter = SyntheticPublicationAdapter()
        with patch(
            "services.marketplace_image_assets.download_public_image",
            side_effect=images,
        ) as download:
            first = self.start(
                adapter,
                key="publication-key-partial-media",
            )
            self.assertEqual(first.status, "queued")
            self.assertEqual(first.attempt_count, 0)
            self.assertEqual(first.error_code, "media_preparation_pending")
            self.assertEqual(adapter.list_calls, [])
            partial = json.loads(first.snapshot.submitted_state_json)
            self.assertTrue(
                all(
                    "/marketplace-assets/images/" in value
                    for value in partial["items"][0]["images"][:2]
                )
            )
            self.assertTrue(
                all(
                    value.startswith("https://img.test/")
                    for value in partial["items"][0]["images"][2:]
                )
            )

            completed = MarketplacePublicationService.poll_operation(
                seller_id=self.seller.id,
                operation_id=first.id,
                adapter=adapter,
                credentials=SYNTHETIC_CREDENTIALS,
            )

        self.assertEqual(download.call_count, 4)
        self.assertEqual(completed.status, "submitted")
        self.assertEqual(completed.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)
        self.assertTrue(
            all(
                "/marketplace-assets/images/" in value
                for value in adapter.submitted_payloads[0]["items"][0][
                    "images"
                ]
            )
        )

    def test_update_keeps_ozon_baseline_media_and_can_remain_noop(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.app.config.update(
            SECRET_KEY="publication-media-secret",
            PUBLIC_BASE_URL="https://seller.test",
            MARKETPLACE_IMAGE_ASSET_DIR=temporary.name,
            MARKETPLACE_IMAGE_ASSETS_TEST_ENFORCE=True,
            OZON_MEDIA_ASSET_URLS_PER_ATTEMPT=10,
        )
        desired = self.desired_payload()
        self.attach_listing(desired)
        self.draft.media_json = json.dumps({"images": []})
        db.session.commit()
        self.expected_version = self.draft.version
        adapter = SyntheticFullStateAdapter(desired)

        with patch(
            "services.marketplace_image_assets.download_public_image",
        ) as download:
            operation = self.start_update(
                adapter,
                key="product-update-key-noop-media-baseline",
            )

        self.assertFalse(download.called)
        self.assertEqual(operation.status, "succeeded")
        self.assertEqual(operation.attempt_count, 0)
        self.assertEqual(
            json.loads(operation.item_results_json)[0]["status"],
            "already_current",
        )
        self.assertEqual(adapter.submitted_payloads, [])
        summary = json.loads(operation.request_summary_json)
        self.assertEqual(summary["media_asset_slots"], [])
        self.assertEqual(summary["media_asset_state"], "ready")
        submitted = json.loads(operation.snapshot.submitted_state_json)
        self.assertEqual(
            (
                [submitted["items"][0]["primary_image"]]
                + submitted["items"][0]["images"]
            ),
            desired["items"][0]["images"],
        )

    def test_due_poll_prioritizes_attempted_reconciliation_over_older_queue(self):
        now = datetime.utcnow()
        queued = MarketplaceOperation(
            seller_id=self.seller.id,
            marketplace_id=self.marketplace.id,
            account_id=self.account.id,
            operation_kind="product_import",
            status="queued",
            idempotency_key="priority-queued-operation",
            request_fingerprint="b" * 64,
            contract_version="ozon-product-import-v3-2026-07-10",
            request_summary_json="{}",
            quota_snapshot_json="{}",
            provider_request_ids_json="[]",
            item_results_json="[]",
            next_poll_at=now - timedelta(minutes=2),
        )
        submitted = MarketplaceOperation(
            seller_id=self.seller.id,
            marketplace_id=self.marketplace.id,
            account_id=self.account.id,
            operation_kind="product_import",
            status="submitted",
            idempotency_key="priority-submitted-operation",
            request_fingerprint="c" * 64,
            contract_version="ozon-product-import-v3-2026-07-10",
            request_summary_json="{}",
            quota_snapshot_json="{}",
            provider_request_ids_json="[]",
            item_results_json="[]",
            attempt_count=1,
            next_poll_at=now - timedelta(minutes=1),
        )
        db.session.add_all([queued, submitted])
        db.session.commit()

        with patch.object(
            MarketplacePublicationService,
            "poll_operation",
        ) as poll:
            result = MarketplacePublicationService.poll_due_operations(
                limit=1,
                now=now,
                allow_submission=True,
            )

        self.assertEqual(result["selected"], 1)
        self.assertEqual(result["processed"], 1)
        self.assertEqual(
            poll.call_args.kwargs["operation_id"],
            submitted.id,
        )

    def test_due_poll_serves_queue_behind_twenty_pending_tasks(self):
        now = datetime.utcnow()
        operations = []
        for index in range(20):
            operations.append(MarketplaceOperation(
                seller_id=self.seller.id,
                marketplace_id=self.marketplace.id,
                account_id=self.account.id,
                operation_kind="product_import",
                status="polling",
                idempotency_key=f"fair-pending-operation-{index:02d}",
                request_fingerprint="a" * 64,
                contract_version="ozon-product-import-v3-2026-07-10",
                request_summary_json="{}",
                quota_snapshot_json="{}",
                provider_request_ids_json="[]",
                item_results_json="[]",
                attempt_count=1,
                created_at=now - timedelta(minutes=2),
                next_poll_at=now - timedelta(minutes=1),
            ))
        queued = MarketplaceOperation(
            seller_id=self.seller.id,
            marketplace_id=self.marketplace.id,
            account_id=self.account.id,
            operation_kind="product_import",
            status="queued",
            idempotency_key="fair-queued-operation-21",
            request_fingerprint="b" * 64,
            contract_version="ozon-product-import-v3-2026-07-10",
            request_summary_json="{}",
            quota_snapshot_json="{}",
            provider_request_ids_json="[]",
            item_results_json="[]",
            created_at=now,
            next_poll_at=now - timedelta(minutes=1),
        )
        db.session.add_all([*operations, queued])
        db.session.commit()
        pending_ids = {operation.id for operation in operations}

        with patch.object(MarketplacePublicationService, "poll_operation") as poll:
            selected = MarketplacePublicationService.poll_due_operations(
                limit=20, now=now,
            )
            selected_ids = [call.kwargs["operation_id"] for call in poll.call_args_list]
        self.assertEqual(selected["selected"], 20)
        self.assertEqual(len(pending_ids.intersection(selected_ids)), 19)
        self.assertIn(queued.id, selected_ids)

        # The one-slot caller also makes progress after each pending task has
        # been served once; no transient in-memory lane counter is needed.
        seen = []
        for tick in range(21):
            current = now + timedelta(seconds=tick + 1)

            def record_poll(*, operation_id, **_kwargs):
                seen.append(operation_id)
                if operation_id in pending_ids:
                    db.session.get(MarketplaceOperation, operation_id).last_polled_at = current
                    db.session.commit()

            with patch.object(
                MarketplacePublicationService, "poll_operation",
                side_effect=record_poll,
            ):
                MarketplacePublicationService.poll_due_operations(
                    limit=1, now=current,
                )
            if queued.id in seen:
                break
        self.assertIn(queued.id, seen)

        # Available slots are not discarded when the pending lane is short.
        for operation in operations[1:]:
            operation.status = 'queued'
        db.session.commit()
        with patch.object(MarketplacePublicationService, 'poll_operation') as poll:
            selected = MarketplacePublicationService.poll_due_operations(limit=20, now=now)
        self.assertEqual(selected['selected'], 20)
        self.assertEqual(poll.call_count, 20)

    def test_quota_rate_limit_defers_without_provider_write(self):
        adapter = SyntheticPublicationAdapter()
        before = datetime.utcnow()
        with patch.object(
            adapter, "get_operation_limits",
            side_effect=OzonAPIError(
                "rate limit", code="ozon_rate_limited",
                status_code=429, retry_after=3600.25,
                retriable=True,
            ),
        ):
            deferred = self.start(adapter, key="quota-rate-limit-001")
        self.assertEqual(deferred.status, "queued")
        self.assertEqual(deferred.attempt_count, 0)
        self.assertEqual(adapter.submitted_payloads, [])
        self.assertGreaterEqual(
            deferred.next_poll_at, before + timedelta(seconds=3601),
        )
        self.assertEqual(
            json.loads(deferred.request_summary_json)["provider_read_not_before"],
            deferred.next_poll_at.isoformat(),
        )
        still_deferred = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id, operation_id=deferred.id,
            adapter=adapter, credentials=SYNTHETIC_CREDENTIALS,
            now=before + timedelta(minutes=10),
        )
        self.assertEqual(still_deferred.status, "queued")
        self.assertEqual(adapter.submitted_payloads, [])
        submitted = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id, operation_id=deferred.id,
            adapter=adapter, credentials=SYNTHETIC_CREDENTIALS,
            now=deferred.next_poll_at + timedelta(seconds=1),
        )
        self.assertEqual(submitted.status, "submitted")
        self.assertEqual(submitted.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_quota_permission_failure_requires_action_before_write(self):
        adapter = SyntheticPublicationAdapter()
        with patch.object(
            adapter, "get_operation_limits",
            side_effect=OzonAPIError(
                "forbidden", code="ozon_auth_error", status_code=403,
            ),
        ):
            failed = self.start(adapter, key="quota-forbidden-001")
        self.assertEqual(failed.status, "failed")
        self.assertEqual(failed.attempt_count, 0)
        self.assertEqual(adapter.submitted_payloads, [])

    def test_quota_server_error_defers_before_write(self):
        adapter = SyntheticPublicationAdapter()
        with patch.object(
            adapter, "get_operation_limits",
            side_effect=OzonAPIError(
                "server error", code="ozon_server_error", status_code=503,
            ),
        ):
            deferred = self.start(adapter, key="quota-server-error-001")
        self.assertEqual(deferred.status, "queued")
        self.assertEqual(deferred.attempt_count, 0)
        self.assertEqual(adapter.submitted_payloads, [])

    def test_account_quota_capacity_subtracts_local_active_reservations(self):
        adapter = SyntheticPublicationAdapter()
        operation = self.start(adapter, key="publication-key-quota-capacity")
        self.assertEqual(operation.status, "submitted")
        self.assertEqual(operation.quota_reserved, 1)

        capacity = MarketplacePublicationService.get_account_quota_capacity(
            seller_id=self.seller.id,
            account_id=self.account.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
            now=datetime.utcnow(),
        )

        self.assertEqual(capacity["provider_remaining"], 96)
        self.assertEqual(capacity["local_reserved"], 1)
        self.assertEqual(capacity["available"], 95)

    def test_submit_poll_finalize_and_idempotent_retry(self):
        adapter = SyntheticPublicationAdapter()
        operation = self.start(adapter)
        self.assertEqual(operation.status, "submitted")
        self.assertEqual(operation.attempt_count, 1)
        self.assertEqual(operation.external_task_id, "456")
        self.assertEqual(len(adapter.submitted_payloads), 1)
        before = json.loads(operation.snapshot.before_state_json)
        self.assertIs(before["exists"], False)

        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(completed.quota_reserved, 0)
        listing = MarketplaceListing.query.filter_by(
            account_id=self.account.id,
            offer_id="safe-offer",
        ).one()
        self.assertEqual(listing.external_product_id, "987654")
        self.assertEqual(listing.imported_product_id, self.source.id)
        db.session.refresh(self.draft)
        self.assertEqual(self.draft.status, "published")
        self.assertEqual(self.draft.published_listing_id, listing.id)
        self.assertEqual(completed.snapshot.rollback_status, "available")

        repeated = self.start(adapter)
        self.assertEqual(repeated.id, completed.id)
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_imported_task_waits_for_complete_live_observation(self):
        adapter = SyntheticPublicationAdapter()
        operation = self.start(adapter, key="create-full-read-pending-001")
        now = datetime.utcnow()
        with patch.object(
            adapter, "get_product_attributes",
            side_effect=OzonAPIError(
                "rate limit", code="ozon_rate_limited", status_code=429,
                retry_after=90.5, retriable=True,
            ),
        ):
            pending = MarketplacePublicationService.poll_operation(
                seller_id=self.seller.id, operation_id=operation.id,
                adapter=adapter, credentials=SYNTHETIC_CREDENTIALS,
                now=now,
            )
        self.assertEqual(pending.status, "polling")
        self.assertEqual(pending.attempt_count, 1)
        self.assertEqual(pending.error_code, "create_live_read_unavailable")
        self.assertGreaterEqual(
            pending.next_poll_at, now + timedelta(seconds=91),
        )
        self.assertEqual(pending.snapshot.confirmed_state_json, "{}")
        self.assertEqual(MarketplaceListing.query.count(), 0)
        self.assertEqual(len(adapter.submitted_payloads), 1)

        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id, operation_id=operation.id,
            adapter=adapter, credentials=SYNTHETIC_CREDENTIALS,
            now=pending.next_poll_at + timedelta(seconds=1),
        )
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(len(adapter.submitted_payloads), 1)
        self.assertEqual(
            {name for name, _ in adapter.full_read_calls},
            {"info", "attributes", "prices", "pictures"},
        )
        listing = MarketplaceListing.query.one()
        self.assertIsNotNone(listing.info_synced_at)
        self.assertEqual(
            json.loads(completed.snapshot.confirmed_state_json)["source"],
            "task_status_and_live_state",
        )

    def test_imported_task_with_live_price_drift_remains_unconfirmed(self):
        adapter = SyntheticPublicationAdapter()
        operation = self.start(adapter, key="create-live-price-drift-001")
        actual_read_prices = adapter.read_prices

        def changed_price(credentials, payload):
            response = actual_read_prices(credentials, payload)
            response["items"][0]["price"]["price"] = "999"
            return response

        now = datetime.utcnow()
        with patch.object(adapter, "read_prices", side_effect=changed_price):
            pending = MarketplacePublicationService.poll_operation(
                seller_id=self.seller.id, operation_id=operation.id,
                adapter=adapter, credentials=SYNTHETIC_CREDENTIALS,
                now=now,
            )
        self.assertEqual(pending.status, "polling")
        self.assertEqual(pending.error_code, "create_live_state_pending")
        self.assertEqual(MarketplaceListing.query.count(), 0)
        self.assertEqual(pending.snapshot.confirmed_state_json, "{}")

        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id, operation_id=operation.id,
            adapter=adapter, credentials=SYNTHETIC_CREDENTIALS,
            now=pending.next_poll_at + timedelta(seconds=1),
        )
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_imported_task_readback_cooldown_past_deadline_is_uncertain(self):
        adapter = SyntheticPublicationAdapter()
        operation = self.start(adapter, key="create-readback-deadline-001")
        now = operation.deadline_at - timedelta(seconds=70)
        with patch.object(
            adapter, "get_products",
            side_effect=OzonAPIError(
                "rate limit", code="ozon_rate_limited", status_code=429,
                retry_after=120, retriable=True,
            ),
        ):
            unresolved = MarketplacePublicationService.poll_operation(
                seller_id=self.seller.id, operation_id=operation.id,
                adapter=adapter, credentials=SYNTHETIC_CREDENTIALS,
                now=now,
            )
        self.assertEqual(unresolved.status, "uncertain")
        self.assertEqual(unresolved.error_code, "create_live_read_deadline_exceeded")
        self.assertIsNone(unresolved.next_poll_at)
        self.assertEqual(unresolved.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)
        self.assertEqual(MarketplaceListing.query.count(), 0)

    def test_existing_offer_fails_before_quota_and_write(self):
        adapter = SyntheticPublicationAdapter(offer_exists=True)
        operation = self.start(adapter, key="publication-key-existing")
        self.assertEqual(operation.status, "failed")
        self.assertEqual(operation.error_code, "offer_exists_upstream")
        self.assertEqual(operation.attempt_count, 0)
        self.assertEqual(adapter.submitted_payloads, [])

    def test_exact_offer_preflight_accepts_current_ozon_opaque_cursor(self):
        adapter = SyntheticPublicationAdapter(
            offer_exists=True,
            cursor_on_exact=True,
        )
        found = MarketplacePublicationService._offer_lookup(
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
            offer_id=self.draft.offer_id,
        )
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0]["offer_id"], self.draft.offer_id)

    def test_ambiguous_write_reconciles_live_without_retry(self):
        adapter = SyntheticPublicationAdapter(ambiguous=True)
        operation = self.start(adapter, key="publication-key-ambiguous")
        self.assertEqual(operation.status, "uncertain")
        self.assertEqual(operation.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_manual_uncertain_resolution_releases_only_local_quota(self):
        adapter = SyntheticPublicationAdapter(ambiguous=True)
        operation = self.start(adapter, key="publication-key-manual-resolution")
        self.assertEqual(operation.status, "uncertain")
        self.assertEqual(operation.quota_reserved, 1)

        resolved = MarketplacePublicationService.resolve_uncertain(
            seller_id=self.seller.id,
            operation_id=operation.id,
            expected_version=operation.version,
            reason="Проверено вручную в кабинете; результат пока неясен",
            resolved_by_user_id=self.user.id,
        )
        self.assertEqual(resolved.status, "uncertain")
        self.assertEqual(resolved.quota_reserved, 0)
        self.assertIsNone(resolved.next_poll_at)
        self.assertEqual(resolved.error_code, "manual_uncertain_resolution")
        summary = json.loads(resolved.request_summary_json)
        self.assertEqual(
            summary["manual_resolution"]["upstream_outcome"],
            "still_uncertain",
        )
        self.assertEqual(len(adapter.submitted_payloads), 1)

        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(completed.reconcile_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_prewrite_outage_stays_queued_and_rollout_flag_prevents_submit(self):
        adapter = SyntheticPublicationAdapter(preflight_error=True)
        operation = self.start(adapter, key="publication-key-prewrite")
        self.assertEqual(operation.status, "queued")
        self.assertEqual(operation.attempt_count, 0)
        adapter.preflight_error = False

        still_queued = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
            allow_submission=False,
        )
        self.assertEqual(still_queued.status, "queued")
        self.assertEqual(adapter.submitted_payloads, [])

        submitted = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
            allow_submission=True,
        )
        self.assertEqual(submitted.status, "submitted")
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_task_poll_outage_stops_automatic_retries_after_deadline(self):
        adapter = SyntheticPublicationAdapter()
        operation = self.start(adapter, key="publication-key-poll-deadline")
        adapter.get_submission = MagicMock(side_effect=OzonAPIError(
            "synthetic status outage",
            code="synthetic_status_outage",
            request_id="synthetic-status-request",
        ))

        stopped = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
            now=operation.deadline_at + timedelta(seconds=1),
        )

        self.assertEqual(stopped.status, "uncertain")
        self.assertEqual(
            stopped.error_code,
            "ozon_task_poll_deadline_exceeded",
        )
        self.assertIsNone(stopped.next_poll_at)
        self.assertEqual(stopped.poll_count, 1)

    def test_task_cooldown_survives_reload_and_manual_poll_without_rewrite(self):
        adapter = SyntheticPublicationAdapter()
        operation = self.start(adapter)
        operation_id = operation.id
        now = operation.submitted_at + timedelta(seconds=20)
        original = adapter.get_submission
        adapter.get_submission = MagicMock(side_effect=OzonAPIError(
            "synthetic throttling", status_code=429, retry_after=7200.25,
        ))
        def poll(at):
            return MarketplacePublicationService.poll_operation(
                seller_id=self.seller.id, operation_id=operation_id,
                adapter=adapter, credentials=SYNTHETIC_CREDENTIALS, now=at,
                allow_submission=False,
            )
        pending = poll(now)
        self.assertEqual(pending.next_poll_at, now + timedelta(seconds=7201))
        db.session.expire_all()
        poll(now + timedelta(minutes=30))
        adapter.get_submission.assert_called_once()
        adapter.get_submission = original
        completed = poll(now + timedelta(seconds=7201))
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(completed.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_task_cooldown_past_deadline_stops_but_keeps_manual_read_gate(self):
        adapter = SyntheticPublicationAdapter()
        operation = self.start(adapter)
        now = operation.submitted_at + timedelta(seconds=20)
        adapter.get_submission = MagicMock(side_effect=OzonAPIError(
            "synthetic throttling", status_code=429, retry_after=172800,
        ))
        for at in (now, now + timedelta(days=1)):
            operation = MarketplacePublicationService.poll_operation(
                seller_id=self.seller.id, operation_id=operation.id,
                adapter=adapter, credentials=SYNTHETIC_CREDENTIALS, now=at,
            )
            self.assertEqual(operation.status, "uncertain")
            self.assertIsNone(operation.next_poll_at)
        adapter.get_submission.assert_called_once()
        self.assertEqual(operation.attempt_count, 1)

    def test_update_accepted_task_read_outage_respects_deadline(self):
        prior = self.prior_payload()
        self.attach_listing(prior)
        adapter = SyntheticFullStateAdapter(prior)
        operation = self.start_update(adapter)
        with patch.object(OzonProductStateContract, "read_full_payload",
                          side_effect=OzonAPIError("synthetic outage")):
            operation = MarketplacePublicationService.poll_operation(
                seller_id=self.seller.id, operation_id=operation.id,
                adapter=adapter, credentials=SYNTHETIC_CREDENTIALS,
                now=operation.deadline_at + timedelta(seconds=1),
            )
        self.assertEqual(operation.status, "uncertain")
        self.assertEqual(operation.error_code, "update_live_read_deadline_exceeded")
        self.assertIsNone(operation.next_poll_at)
        self.assertEqual(operation.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_update_live_read_cooldown_then_exact_success(self):
        prior = self.prior_payload()
        self.attach_listing(prior)
        adapter = SyntheticFullStateAdapter(prior)
        operation = self.start_update(adapter)
        now = operation.submitted_at + timedelta(seconds=20)
        def poll(at):
            return MarketplacePublicationService.poll_operation(
                seller_id=self.seller.id, operation_id=operation.id,
                adapter=adapter, credentials=SYNTHETIC_CREDENTIALS, now=at,
            )
        with patch.object(OzonProductStateContract, "read_full_payload",
                          side_effect=OzonAPIError("throttled", retry_after=7200)) as read:
            pending = poll(now)
            self.assertEqual(pending.next_poll_at, now + timedelta(hours=2))
            poll(now + timedelta(minutes=30))
            read.assert_called_once()
        self.assertEqual(poll(now + timedelta(hours=2)).status, "succeeded")
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_prewrite_cooldown_does_not_allow_early_manual_submission(self):
        adapter = SyntheticPublicationAdapter()
        with patch.object(adapter, "list_products", side_effect=OzonAPIError(
            "synthetic throttling", retry_after=7200,
        )) as read:
            operation = self.start(adapter)
            self.assertEqual(operation.status, "queued")
            self.assertEqual(operation.attempt_count, 0)
            MarketplacePublicationService.poll_operation(
                seller_id=self.seller.id, operation_id=operation.id,
                adapter=adapter, credentials=SYNTHETIC_CREDENTIALS,
                now=operation.next_poll_at - timedelta(seconds=1),
            )
            read.assert_called_once()
        self.assertEqual(adapter.submitted_payloads, [])

    def test_active_operation_blocks_draft_mutation_and_foreign_read(self):
        operation = self.start(
            SyntheticPublicationAdapter(status="pending"),
            key="publication-key-active",
        )
        with self.assertRaises(MarketplaceDraftConflict):
            MarketplaceDraftService.update_draft(
                seller_id=self.seller.id,
                draft_id=self.draft.id,
                expected_version=self.draft.version,
                patch={"offer_id": "changed-offer"},
            )
        with self.assertRaises(MarketplacePublicationNotFound):
            MarketplacePublicationService.get_operation(
                seller_id=self.foreign_seller.id,
                operation_id=operation.id,
            )

    def test_public_serializer_never_exposes_payload_idempotency_or_secret(self):
        operation = self.start(
            SyntheticPublicationAdapter(),
            key="publication-key-public",
        )
        document = operation.to_public_dict(detail=True)
        encoded = json.dumps(document, ensure_ascii=False)
        self.assertNotIn("synthetic-key", encoded)
        self.assertNotIn("idempotency_key", document)
        self.assertNotIn("submitted_state", encoded)
        self.assertNotIn("api_key", encoded.lower())
        self.assertEqual(
            MarketplaceOperation.query.filter_by(seller_id=self.seller.id).count(),
            1,
        )

    def test_full_state_exact_price_read_accepts_current_ozon_cursor(self):
        prior = self.prior_payload()
        adapter = SyntheticFullStateAdapter(
            prior,
            cursor_on_exact=True,
        )
        state = OzonProductStateContract.read_full_payload(
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
            product_id=987654,
            offer_id=prior["items"][0]["offer_id"],
        )
        self.assertEqual(
            state["fingerprint"],
            OzonProductStateContract.fingerprint(prior),
        )

    def test_update_uses_listing_fallbacks_without_deleting_full_state(self):
        prior = self.prior_payload()
        self.attach_listing(prior)
        self.draft.media_json = json.dumps({"images": []})
        self.draft.dimensions_json = "{}"
        self.draft.barcodes_json = "[]"
        self.draft.commercial_json = "{}"
        db.session.commit()
        self.expected_version = self.draft.version
        adapter = SyntheticFullStateAdapter(prior)

        operation = self.start_update(
            adapter,
            key="product-update-key-preserve-listing",
        )

        self.assertEqual(operation.status, "submitted")
        self.assertEqual(len(adapter.submitted_payloads), 1)
        submitted = adapter.submitted_payloads[0]["items"][0]
        before = prior["items"][0]
        for field in (
            "width",
            "height",
            "depth",
            "dimension_unit",
            "weight",
            "weight_unit",
            "price",
            "old_price",
            "vat",
            "currency_code",
            "barcode",
        ):
            self.assertEqual(submitted[field], before[field])
        submitted_gallery = (
            [submitted["primary_image"]]
            + submitted.get("images", [])
        )
        before_gallery = (
            ([before["primary_image"]] if before.get("primary_image") else [])
            + before.get("images", [])
        )
        self.assertEqual(submitted_gallery, before_gallery)
        self.assertEqual(submitted["name"], "Безопасный товар")
        self.assertEqual(
            next(
                item for item in submitted["attributes"]
                if item["id"] == 4191
            )["values"],
            [{"value": "Наблюдаемое описание"}],
        )

    def test_explicit_attribute_removal_survives_full_state_preflight(self):
        prior = self.prior_payload()
        prior["items"][0]["attributes"].append({
            "id": 99999,
            "complex_id": 0,
            "values": [{"value": "Legacy"}],
        })
        self.attach_listing(prior)
        self.draft = MarketplaceDraftService.update_draft(
            seller_id=self.seller.id,
            draft_id=self.draft.id,
            expected_version=self.draft.version,
            patch={"attribute_removals": [{
                "attribute_id": "99999",
                "complex_id": "0",
            }]},
            corrected_by_user_id=self.user.id,
        )
        adapter = SyntheticFullStateAdapter(prior)
        with patch.object(
            MarketplaceDraftService,
            "_build_validation_result",
            return_value=self.validation_result(),
        ):
            self.draft = MarketplaceDraftService.validate_draft(
                seller_id=self.seller.id,
                draft_id=self.draft.id,
                expected_version=self.draft.version,
            )
            self.expected_version = self.draft.version
            operation = self.start_update(
                adapter,
                key="product-update-explicit-attribute-cleanup",
            )

        self.assertEqual(operation.status, "submitted")
        self.assertEqual(len(adapter.submitted_payloads), 1)
        self.assertEqual(
            json.loads(operation.request_summary_json)[
                "explicit_attribute_removal_count"
            ],
            1,
        )
        submitted_ids = {
            item["id"]
            for item in adapter.submitted_payloads[0]["items"][0][
                "attributes"
            ]
        }
        self.assertNotIn(99999, submitted_ids)
        before = json.loads(operation.snapshot.before_state_json)
        self.assertIn(
            99999,
            {
                item["id"]
                for item in before["payload"]["items"][0]["attributes"]
            },
        )

    def test_update_rejects_stale_listing_projection_before_operation(self):
        prior = self.prior_payload()
        listing = self.attach_listing(prior)
        listing.attributes_synced_at = datetime.utcnow() - timedelta(hours=49)
        db.session.commit()

        with self.assertRaises(MarketplaceDraftConflict) as caught:
            MarketplaceDraftService.publication_documents(self.draft)

        self.assertEqual(
            caught.exception.code,
            "ozon_listing_snapshot_stale",
        )
        self.assertEqual(MarketplaceOperation.query.count(), 0)

    def test_update_catalog_baseline_drift_fails_before_provider_write(self):
        prior = self.prior_payload()
        self.attach_listing(prior)
        drifted = deepcopy(prior)
        drifted["items"][0]["name"] = "Внешнее изменение после sync"
        adapter = SyntheticFullStateAdapter(drifted)

        operation = self.start_update(
            adapter,
            key="product-update-key-before-drift",
        )

        self.assertEqual(operation.status, "failed")
        self.assertEqual(operation.error_code, "update_before_state_drift")
        self.assertEqual(operation.attempt_count, 0)
        self.assertEqual(adapter.submitted_payloads, [])

    def test_import_only_type_omission_is_already_current_without_write(self):
        self.enable_import_only_type_attribute()
        desired = self.desired_payload()
        live = deepcopy(desired)
        live["items"][0]["attributes"] = [
            attribute
            for attribute in live["items"][0]["attributes"]
            if attribute["id"] != 8229
        ]
        self.attach_listing(live)
        adapter = SyntheticTypeOmittingAdapter(live)

        operation = self.start_update(
            adapter,
            key="product-update-key-import-only-noop",
        )

        self.assertEqual(operation.status, "succeeded")
        self.assertEqual(operation.attempt_count, 0)
        self.assertEqual(adapter.submitted_payloads, [])
        self.assertEqual(
            operation.snapshot.confirmed_fingerprint,
            OzonProductStateContract.fingerprint(live),
        )
        self.assertNotEqual(
            operation.snapshot.confirmed_fingerprint,
            operation.request_fingerprint,
        )
        self.assertEqual(
            json.loads(operation.request_summary_json)[
                "provider_roundtrip_omitted_attribute_ids"
            ],
            ["8229"],
        )

    def test_task_confirmed_type_omission_preserves_actual_live_and_rolls_back(self):
        self.enable_import_only_type_attribute()
        prior = self.prior_payload()
        prior["items"][0]["attributes"] = [
            attribute
            for attribute in prior["items"][0]["attributes"]
            if attribute["id"] != 8229
        ]
        listing = self.attach_listing(prior)
        adapter = SyntheticTypeOmittingAdapter(prior)

        operation = self.start_update(
            adapter,
            key="product-update-key-import-only-write",
        )
        self.assertEqual(operation.status, "submitted")
        self.assertEqual(operation.attempt_count, 1)

        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )

        expected_visible = deepcopy(adapter.live_payload)
        expected_visible["items"][0]["attributes"] = [
            attribute
            for attribute in expected_visible["items"][0]["attributes"]
            if attribute["id"] != 8229
        ]
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(completed.attempt_count, 1)
        self.assertEqual(completed.snapshot.rollback_status, "available")
        self.assertEqual(
            completed.snapshot.confirmed_fingerprint,
            OzonProductStateContract.fingerprint(expected_visible),
        )
        db.session.refresh(listing)
        self.assertNotIn(
            "8229",
            {
                str(attribute["id"])
                for attribute in json.loads(listing.attributes_json)
            },
        )
        rollback_payload = json.loads(
            completed.snapshot.rollback_state_json
        )["payload"]
        self.assertIn(
            8229,
            {
                attribute["id"]
                for attribute in rollback_payload["items"][0]["attributes"]
            },
        )

        rollback = MarketplacePublicationService.start_update_rollback(
            seller_id=self.seller.id,
            operation_id=completed.id,
            expected_version=completed.version,
            idempotency_key="product-update-rollback-import-only",
            created_by_user_id=self.user.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(rollback.status, "submitted")
        restored = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=rollback.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        restored_visible = deepcopy(adapter.live_payload)
        restored_visible["items"][0]["attributes"] = [
            attribute
            for attribute in restored_visible["items"][0]["attributes"]
            if attribute["id"] != 8229
        ]
        self.assertEqual(restored.status, "succeeded")
        self.assertEqual(
            OzonProductStateContract.fingerprint(restored_visible),
            OzonProductStateContract.fingerprint(prior),
        )
        db.session.refresh(completed.snapshot)
        self.assertEqual(completed.snapshot.rollback_status, "succeeded")
        self.assertEqual(len(adapter.submitted_payloads), 2)
        repeated = MarketplacePublicationService.start_update_rollback(
            seller_id=self.seller.id,
            operation_id=completed.id,
            expected_version=completed.version,
            idempotency_key="product-update-rollback-import-only",
            created_by_user_id=self.user.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(repeated.id, restored.id)
        self.assertEqual(len(adapter.submitted_payloads), 2)

    def test_new_required_compliance_makes_old_state_rollback_unavailable(self):
        self.enable_import_only_type_attribute()
        self.add_required_dictionary_attribute(
            external_attribute_id="22232",
            name="ТН ВЭД",
            external_value_id="971397758",
            value="3307900008 - Косметические средства, прочие",
        )
        prior = self.prior_payload()
        prior["items"][0]["attributes"] = [
            attribute
            for attribute in prior["items"][0]["attributes"]
            if attribute["id"] not in {8229, 22232}
        ]
        self.attach_listing(prior)
        adapter = SyntheticTypeOmittingAdapter(prior)

        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=self.start_update(
                adapter,
                key="product-update-key-new-required-compliance",
            ).id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )

        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(completed.snapshot.rollback_status, "unavailable")
        self.assertEqual(
            completed.snapshot.rollback_error_code,
            "update_rollback_prior_state_not_publishable",
        )
        self.assertEqual(completed.snapshot.rollback_state_json, "{}")

    def test_type_omission_plus_visible_drift_remains_uncertain(self):
        self.enable_import_only_type_attribute()
        prior = self.prior_payload()
        prior["items"][0]["attributes"] = [
            attribute
            for attribute in prior["items"][0]["attributes"]
            if attribute["id"] != 8229
        ]
        self.attach_listing(prior)
        adapter = SyntheticTypeOmittingAdapter(prior)
        operation = self.start_update(
            adapter,
            key="product-update-key-import-only-drift",
        )
        adapter.pending_payload["items"][0]["name"] = "Внешняя правка"

        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )

        self.assertEqual(completed.status, "uncertain")
        self.assertEqual(completed.error_code, "update_postwrite_drift")
        self.assertEqual(completed.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_type_omission_without_task_id_does_not_resolve_ambiguous_write(self):
        self.enable_import_only_type_attribute()
        prior = self.prior_payload()
        prior["items"][0]["attributes"] = [
            attribute
            for attribute in prior["items"][0]["attributes"]
            if attribute["id"] != 8229
        ]
        self.attach_listing(prior)
        adapter = SyntheticTypeOmittingAdapter(prior)
        operation = self.start_update(
            adapter,
            key="product-update-key-import-only-no-task",
        )
        adapter.live_payload = deepcopy(adapter.pending_payload)
        adapter.pending_payload = None
        operation.external_task_id = None
        operation.status = "uncertain"
        operation.next_poll_at = None
        db.session.commit()

        reconciled = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )

        self.assertEqual(reconciled.status, "uncertain")
        self.assertEqual(reconciled.error_code, "update_postwrite_drift")
        self.assertEqual(reconciled.attempt_count, 1)
        self.assertEqual(reconciled.reconcile_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)

    def test_full_state_update_and_exact_prior_state_rollback(self):
        prior = self.prior_payload()
        listing = self.attach_listing(prior)
        adapter = SyntheticFullStateAdapter(prior)

        operation = self.start_update(adapter)
        self.assertEqual(operation.operation_kind, "product_update")
        self.assertEqual(operation.status, "submitted")
        self.assertEqual(operation.snapshot.snapshot_kind, "product_update")
        self.assertEqual(
            operation.snapshot.before_fingerprint,
            OzonProductStateContract.fingerprint(prior),
        )
        self.assertEqual(operation.attempt_count, 1)
        self.assertEqual(len(adapter.submitted_payloads), 1)

        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=operation.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(completed.snapshot.rollback_status, "available")
        submitted = adapter.submitted_payloads[0]
        self.assertEqual(
            completed.snapshot.confirmed_fingerprint,
            OzonProductStateContract.fingerprint(submitted),
        )
        self.assertEqual(
            submitted["items"][0]["primary_image"],
            "https://img.test/prior.jpg",
        )
        self.assertEqual(
            submitted["items"][0]["images"],
            ["https://img.test/product.jpg"],
        )
        db.session.refresh(listing)
        self.assertEqual(json.loads(listing.stock_summary_json), {"preserve": True})
        self.assertEqual(listing.title, "Безопасный товар")

        rollback = MarketplacePublicationService.start_update_rollback(
            seller_id=self.seller.id,
            operation_id=completed.id,
            expected_version=completed.version,
            idempotency_key="product-update-rollback-key-0001",
            created_by_user_id=self.user.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(rollback.operation_kind, "product_update_rollback")
        self.assertEqual(rollback.status, "submitted")
        db.session.refresh(completed.snapshot)
        self.assertEqual(completed.snapshot.rollback_status, "pending")

        restored = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=rollback.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(restored.status, "succeeded")
        self.assertEqual(
            OzonProductStateContract.fingerprint(adapter.live_payload),
            OzonProductStateContract.fingerprint(prior),
        )
        db.session.refresh(completed.snapshot)
        self.assertEqual(completed.snapshot.rollback_status, "succeeded")
        self.assertEqual(len(adapter.submitted_payloads), 2)

    def test_update_rollback_drift_fails_before_second_write(self):
        prior = self.prior_payload()
        self.attach_listing(prior)
        adapter = SyntheticFullStateAdapter(prior)
        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=self.start_update(
                adapter,
                key="product-update-key-drift",
            ).id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        drifted = deepcopy(adapter.live_payload)
        drifted["items"][0]["name"] = "Внешнее изменение"
        adapter.live_payload = drifted

        rollback = MarketplacePublicationService.start_update_rollback(
            seller_id=self.seller.id,
            operation_id=completed.id,
            expected_version=completed.version,
            idempotency_key="product-update-rollback-key-drift",
            created_by_user_id=self.user.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(rollback.status, "failed")
        self.assertEqual(rollback.error_code, "update_before_state_drift")
        self.assertEqual(len(adapter.submitted_payloads), 1)
        db.session.refresh(completed.snapshot)
        self.assertEqual(completed.snapshot.rollback_status, "conflict")

    def test_create_compensation_archives_only_unchanged_created_listing(self):
        desired = self.desired_payload()
        adapter = SyntheticFullStateAdapter(create_mode=True)
        created = self.start(adapter, key="publication-key-archive")
        completed = MarketplacePublicationService.poll_operation(
            seller_id=self.seller.id,
            operation_id=created.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(completed.status, "succeeded")
        self.assertEqual(completed.snapshot.rollback_status, "available")
        self.assertEqual(
            OzonProductStateContract.fingerprint(adapter.live_payload),
            OzonProductStateContract.fingerprint(desired),
        )

        archived = MarketplacePublicationService.start_create_rollback(
            seller_id=self.seller.id,
            operation_id=completed.id,
            expected_version=completed.version,
            idempotency_key="product-create-archive-key-0001",
            created_by_user_id=self.user.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(archived.status, "succeeded")
        self.assertEqual(archived.operation_kind, "product_import_rollback")
        self.assertEqual(adapter.archive_calls, [{"product_id": [987654]}])
        db.session.refresh(completed.snapshot)
        self.assertEqual(completed.snapshot.rollback_status, "succeeded")
        listing = db.session.get(MarketplaceListing, completed.listing_id)
        self.assertTrue(listing.is_archived)
        db.session.refresh(self.draft)
        self.assertEqual(self.draft.status, "archived")
        repeated = MarketplacePublicationService.start_create_rollback(
            seller_id=self.seller.id,
            operation_id=completed.id,
            expected_version=completed.version,
            idempotency_key="product-create-archive-key-0001",
            created_by_user_id=self.user.id,
            adapter=adapter,
            credentials=SYNTHETIC_CREDENTIALS,
        )
        self.assertEqual(repeated.id, archived.id)
        self.assertEqual(adapter.archive_calls, [{"product_id": [987654]}])


if __name__ == "__main__":
    unittest.main()
