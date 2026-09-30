"""Bounded, read-only XLSX of one exact completed finance selection."""
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from time import monotonic

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, select
from sqlalchemy.exc import OperationalError

from models import (db, MarketplaceFinanceFact as Fact, MarketplaceFinanceFactItem as Item,
                    MarketplaceFinanceComponent as Component, MarketplaceFinanceSync as Sync,
                    MarketplaceListing as Listing)
from services.marketplace_finance import (MarketplaceFinanceService as Finance,
    MarketplaceFinanceError, MarketplaceFinanceNotFound, MarketplaceFinanceValidationError)

MAX_FACTS = 10_000
MAX_CHILDREN = 50_000  # products and explanatory components together
MAX_TEXT_BYTES = 16 * 1024 * 1024
MAX_FILE_BYTES = 16 * 1024 * 1024
MAX_SECONDS = 20
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
CATEGORIES = {'POSTING':'По отправлению', 'ITEM':'По товару',
              'NON_ITEM':'По продавцу', 'UNSPECIFIED':'Категория не указана'}
MATCHES = {'matched':'Карточка связана', 'unmatched':'Карточка не найдена',
           'ambiguous':'Несколько карточек', 'unavailable':'Связанная карточка недоступна'}


class FinanceExportLimit(MarketplaceFinanceError):
    status_code = 422
    code = 'finance_export_limit'


def _check_budget(deadline):
    if monotonic() > deadline:
        raise FinanceExportLimit('Подготовка файла заняла слишком много времени. Сократите период или уточните фильтры.')


def _limit(rows, maximum):
    if len(rows) > maximum:
        raise FinanceExportLimit('Выборка слишком большая для одного файла: до 10 000 начислений и 50 000 строк состава. Сократите период или уточните фильтры.')


def _amount(value):
    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        raise MarketplaceFinanceValidationError('В сохранённых данных есть неподдерживаемая сумма. Обновите начисления.') from None
    if not number.is_finite() or number.as_tuple().exponent < -4:
        raise MarketplaceFinanceValidationError('В сохранённых данных есть неподдерживаемая сумма. Обновите начисления.')
    return number


def _money_cells(value):
    number = _amount(value)
    exact = format(number, 'f')
    # Excel retains at most 15 significant decimal digits. Large values stay text.
    visible = exact if len(number.normalize().as_tuple().digits) > 15 else number
    return visible, exact


def _collect(*, seller_id, account_id, snapshot_id, as_of, period_code, category, amount_sign, type_id, search):
    if snapshot_id is None or not as_of:
        raise MarketplaceFinanceValidationError('Для выгрузки сначала откройте готовые начисления за выбранные даты.')
    deadline = monotonic() + MAX_SECONDS
    meta = Finance.list_facts(seller_id=seller_id, account_id=account_id, snapshot_id=snapshot_id,
        as_of=as_of, period_code=period_code, category=category, amount_sign=amount_sign,
        type_id=type_id, search=search, page=1, per_page=1, compact=True)
    if meta['pagination']['total'] > MAX_FACTS:
        _limit([None] * (MAX_FACTS + 1), MAX_FACTS)
    account = Finance._owned_account(seller_id=seller_id, account_id=account_id)
    label = account.label or 'Ozon'
    query = Finance._filtered_facts(account=account, snapshot_id=snapshot_id,
        period_start=date.fromisoformat(meta['coverage']['requested_start']),
        period_end=date.fromisoformat(meta['coverage']['requested_end']), category=category,
        amount_sign=amount_sign, type_id=type_id, search=search)
    ids = query.with_entities(Fact.id).subquery()
    # Explicit SQLite read transaction: retention cannot prune half the workbook.
    # No ORM mutation, provider I/O or XLSX serialization occurs in this transaction.
    with db.engine.connect() as connection:
        raw = connection.connection.driver_connection
        sqlite = connection.dialect.name == 'sqlite'
        try:
            if sqlite:
                connection.exec_driver_sql('BEGIN')
                raw.set_progress_handler(lambda: int(monotonic() > deadline), 10000)
            present = connection.execute(select(Sync.id).where(
                Sync.id == snapshot_id, Sync.seller_id == account.seller_id,
                Sync.account_id == account.id, Sync.marketplace_id == account.marketplace_id,
                Sync.status == 'completed', Sync.contract_version == Finance.CONTRACT_VERSION)).first()
            if not present:
                raise MarketplaceFinanceNotFound('Этот снимок начислений недоступен. Откройте актуальные данные.')
            fact_query = query.with_entities(Fact.id, Fact.accrual_id, Fact.fact_date, Fact.unit_number,
                Fact.accrued_category, Fact.total_amount, Fact.currency, Fact.amount_sign,
                Fact.source_endpoint, Fact.observed_at).order_by(Fact.fact_date.desc(), Fact.id.desc())
            facts = [dict(r._mapping) for r in connection.execute(fact_query.limit(MAX_FACTS + 1).statement)]
            _limit(facts, MAX_FACTS)
            item_query = select(Item.id, Item.fact_id, Item.external_sku, Item.match_status,
                Item.listing_id, Listing.id.label('owned_listing_id'), Listing.offer_id, Listing.title).join(
                    ids, Item.fact_id == ids.c.id).outerjoin(Listing, and_(
                        Listing.id == Item.listing_id, Listing.seller_id == account.seller_id,
                        Listing.account_id == account.id, Listing.marketplace_id == account.marketplace_id)).where(
                            Item.seller_id == account.seller_id, Item.account_id == account.id).order_by(Item.fact_id, Item.id)
            items = [dict(r._mapping) for r in connection.execute(item_query.limit(MAX_CHILDREN + 1))]
            _limit(items, MAX_CHILDREN)
            component_query = select(Component.id, Component.fact_id, Component.component_kind,
                Component.external_type_id, Component.type_name, Component.external_sku,
                Component.amount, Component.currency).join(ids, Component.fact_id == ids.c.id).where(
                    Component.seller_id == account.seller_id, Component.account_id == account.id).order_by(Component.fact_id, Component.id)
            components = [dict(r._mapping) for r in connection.execute(component_query.limit(MAX_CHILDREN - len(items) + 1))]
            _limit(components, MAX_CHILDREN - len(items))
        except OperationalError:
            if monotonic() > deadline:
                _check_budget(deadline)
            raise
        finally:
            if sqlite:
                raw.set_progress_handler(None, 0)
            connection.rollback()
    _check_budget(deadline)
    totals = defaultdict(lambda: {'count':0, 'positive':Decimal(0), 'negative':Decimal(0), 'net':Decimal(0)})
    for fact in facts:
        amount = _amount(fact['total_amount']); part = totals[fact['currency']]
        part['count'] += 1; part['net'] += amount
        if amount > 0:part['positive'] += amount
        if amount < 0:part['negative'] += amount
    expected = {r['currency']:r for r in meta['totals']}
    if (len(facts) != meta['pagination']['total'] or set(totals) != set(expected)
        or any(v['count'] != expected[c]['fact_count'] or any(v[k] != Decimal(expected[c][k])
            for k in ['positive','negative','net']) for c,v in totals.items())):
        raise MarketplaceFinanceValidationError('Суммы выгрузки не совпали с журналом. Обновите страницу перед повторной выгрузкой.')
    return meta, label, facts, items, components, totals, deadline


def _text(value):
    if len(value) > 32767 or any(not (c in '\t\r\n' or 0x20 <= ord(c) <= 0xD7FF
            or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF) for c in value):
        raise MarketplaceFinanceValidationError('Один из текстов не поддерживается форматом Excel. Обновите начисления или уточните выборку.')
    return value


def build_workbook(*, seller_id, account_id, snapshot_id, as_of, period_code='30d',
                   category=None, amount_sign=None, type_id=None, search=''):
    meta, label, facts, items, components, totals, deadline = _collect(
        seller_id=seller_id, account_id=account_id, snapshot_id=snapshot_id, as_of=as_of,
        period_code=period_code, category=category, amount_sign=amount_sign, type_id=type_id, search=search)
    parents = {f['id']:f['accrual_id'] for f in facts}
    coverage = meta['coverage']; snapshot = meta['snapshot_sync']
    notes = [
        ['Магазин', label], ['Кабинет Seller Hub', str(account_id)], ['Версия данных', str(snapshot_id)],
        ['Период с', coverage['requested_start']], ['Период по', coverage['requested_end']],
        ['Полученные даты с', coverage['snapshot_start']], ['Полученные даты по', coverage['snapshot_end']],
        ['Покрытие периода', 'Полное' if coverage['complete'] else 'НЕПОЛНОЕ — показаны только имеющиеся начисления'],
        ['Данные обновлены, UTC', snapshot.get('completed_at') or 'Неизвестно'],
        ['Файл подготовлен, UTC', datetime.utcnow().isoformat(timespec='seconds')],
        ['Поиск', Finance._search(search) or 'Не задан'], ['Категория', CATEGORIES.get(category, 'Все категории')],
        ['Направление', {'positive':'Положительные','negative':'Отрицательные','zero':'Нулевые'}.get(amount_sign,'Все суммы')],
        ['Тип в расшифровке', str(type_id) if type_id else 'Любой тип'],
        ['Начислений', len(facts)], ['Товарных строк', len(items)], ['Компонентов', len(components)],
        ['Основа итогов', 'Только сумма каждого начисления: accruals[].total_amount. Валюты считаются отдельно. Это не прибыль и не банковская выплата.'],
        ['Расшифровка', 'Компоненты поясняют начисления, не прибавляются повторно и могут не совпадать с основной суммой. Фильтр типа выбирает целое начисление.'],
        ['Точность', 'Номера и артикулы сохранены текстом. Для каждой суммы есть точная текстовая колонка. Суммы с более чем 15 значащими цифрами сохранены текстом и в основной колонке.'],
        ['Карточки товаров', 'Названия и артикулы — текущий справочник точно связанных карточек; это не историческое название на дату начисления.'],
        ['Пустая выборка', 'Отсутствие строк не означает нулевой баланс магазина.'],
    ]
    summaries = [[c,v['count'],*_money_cells(v['positive']),*_money_cells(v['negative']),*_money_cells(v['net'])]
                 for c,v in sorted(totals.items())]
    fact_rows = [[str(f['id']), f['accrual_id'], f['fact_date'], f['unit_number'] or '',
        CATEGORIES.get(f['accrued_category'], f['accrued_category']), *_money_cells(f['total_amount']),
        f['currency'], f['source_endpoint'], f['observed_at'].isoformat() if f['observed_at'] else 'Неизвестно'] for f in facts]
    item_rows = [[parents[i['fact_id']], i['external_sku'], i['offer_id'] or '', i['title'] or '',
        MATCHES.get('unavailable' if not i['owned_listing_id'] and i['match_status']=='matched' else i['match_status'], 'Связь не подтверждена')] for i in items]
    component_rows = [[parents[c['fact_id']], {'item_fee':'По товару','non_item_fee':'По продавцу',
        'delivery_service':'Услуга доставки'}.get(c['component_kind'], c['component_kind']),
        str(c['external_type_id']), c['type_name'] or '', c['external_sku'] or '',
        *_money_cells(c['amount']), c['currency']] for c in components]
    sheets = [
        ('О выгрузке', ['Параметр','Значение'], notes, [30,100]),
        ('Итоги', ['Валюта','Начислений','Положительные суммы','Положительные — точно, текст','Отрицательные суммы','Отрицательные — точно, текст','Итог начислений','Итог — точно, текст'], summaries, [12,15,24,32,24,32,24,32]),
        ('Начисления', ['ID Seller Hub','Номер начисления Ozon','Дата начисления','Номер отправления / документа','Категория','Сумма','Сумма — точно, текст','Валюта','Источник','Наблюдение, UTC'], fact_rows, [18,27,20,32,24,24,32,12,38,29]),
        ('Товары', ['Номер начисления Ozon','SKU','Артикул продавца','Название товара','Связь с карточкой'], item_rows, [27,23,28,60,34]),
        ('Расшифровка', ['Номер начисления Ozon','Вид компонента','ID типа Ozon','Название типа Ozon','SKU','Сумма компонента','Сумма — точно, текст','Валюта'], component_rows, [27,24,20,52,23,24,32,12]),
    ]
    text_bytes = 0
    for _, headers, rows, _ in sheets:
        for row in [headers, *rows]:
            for value in row:
                if isinstance(value, str):
                    text_bytes += len(_text(value).encode('utf-8'))
                    if text_bytes > MAX_TEXT_BYTES:
                        raise FinanceExportLimit('Для одного файла слишком много текста. Сократите период или уточните фильтры.')
            _check_budget(deadline)
    workbook = Workbook(write_only=True)
    font = Font(name='Calibri', size=11, color='1A1A1A')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1A1A1A')
    alignment = Alignment(vertical='top', wrap_text=True)
    for name, number_format in [('text', '@'), ('money', '#,##0.00##;[Red]-#,##0.00##;0.00'),
                                 ('date', 'yyyy-mm-dd'), ('number', 'General')]:
        workbook.add_named_style(NamedStyle(name=name, font=font, alignment=alignment, number_format=number_format))
    workbook.add_named_style(NamedStyle(name='heading', font=header_font, fill=header_fill, alignment=alignment))
    try:
        for title, headers, rows, widths in sheets:
            sheet = workbook.create_sheet(title);sheet.freeze_panes = 'A2';sheet.sheet_view.showGridLines = False
            for n,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(n)].width = width
            if title != 'О выгрузке':sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows)+1}'
            sheet.row_dimensions[1].height = 32
            for index, row in enumerate([headers, *rows]):
                cells = []
                for value in row:
                    cell = WriteOnlyCell(sheet, value=value)
                    kind = 'text' if isinstance(value, str) else 'money' if isinstance(value, Decimal) else 'date' if isinstance(value, date) else 'number'
                    cell.style = 'heading' if index == 0 else kind
                    if isinstance(value, str):cell.data_type = 's'
                    cells.append(cell)
                sheet.append(cells)
                _check_budget(deadline)
        output = BytesIO();workbook.save(output)
        if output.tell() > MAX_FILE_BYTES:
            raise FinanceExportLimit('Файл больше допустимых 16 МБ. Сократите период или уточните фильтры.')
        _check_budget(deadline)
        suffix = '' if coverage['complete'] else '-partial'
        name = f'ozon-finance-{account_id}-{coverage["requested_start"]}-{coverage["requested_end"]}-snapshot-{snapshot_id}{suffix}.xlsx'
        return output.getvalue(), name
    finally:
        # Write-only mode uses owned temporary XML files, including on a budget/error exit.
        for sheet in workbook.worksheets:
            if not sheet.closed:sheet.close()
            writer = getattr(sheet, '_writer', None)
            if writer and isinstance(writer.out, str):
                from pathlib import Path
                if Path(writer.out).is_file():writer.cleanup()
        workbook.close()
