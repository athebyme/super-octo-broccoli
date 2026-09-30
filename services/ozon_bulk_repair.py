"""Bounded XLSX round-trip for local repair of an Ozon upload run.

The workbook changes seller-owned drafts only.  Importing it never creates a
marketplace operation and never calls Ozon; a separate explicitly confirmed
retry remains the only path to a provider write.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import re
import zipfile
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import or_
from sqlalchemy.orm import joinedload

from models import (
    ImportedProduct,
    MarketplaceAttributeDefinition,
    MarketplaceAttributeValue,
    MarketplaceListing,
    MarketplaceProductDraft,
    MarketplaceProductType,
    db,
)
from services.marketplace_drafts import (
    MarketplaceDraftError,
    MarketplaceDraftService,
    MarketplaceDraftValidationError,
)
from services.source_photo_display import imported_photo_previews
from services.marketplace_operation_locks import (
    release_account_operation_lock,
    try_account_operation_lock,
)
from services.ozon_compliance_suggestions import (
    OzonComplianceSuggestionService,
)
from services.ozon_bulk_upload import (
    OzonBulkUploadConflict,
    OzonBulkUploadService,
    OzonBulkUploadValidationError,
)
from services.ozon_reference_service import OzonReferenceService
from services.ozon_product_state import (
    OzonProductStateContract,
    OzonProductStateError,
)


class OzonBulkRepairService:
    """Export and import one exact, seller-scoped repair workbook."""

    CONTRACT_VERSION = "seller-hub-ozon-repair-v1"
    CARD_SHEET = "Карточки"
    LIST_SHEET = "_lists"
    META_SHEET = "_meta"
    MAX_FILE_BYTES = 2 * 1024 * 1024
    MAX_EDITOR_FORM_BYTES = 2 * 1024 * 1024
    MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
    MAX_ZIP_MEMBERS = 100
    MAX_ROWS = OzonBulkUploadService.MAX_ITEMS
    MAX_COLUMNS = 100
    MAX_EDITOR_ATTRIBUTE_CELLS = 5_000
    MAX_ATTRIBUTE_VALUES = 100
    MAX_EDITOR_VALIDATION_ERRORS = 20
    EXPORTABLE_STATUSES = (
        OzonBulkUploadService.RETRYABLE_ITEM_STATUSES
        | {"excluded"}
    )
    ATTRIBUTE_KEY = re.compile(r"^attribute:([1-9][0-9]{0,19})$")
    FLAT_ATTRIBUTE_PATH = re.compile(r"^attributes\[([0-9]+)\]$")
    COMPLEX_ATTRIBUTE_PATH = re.compile(
        r"^complex_attributes\[([0-9]+)\]\.attributes\[([0-9]+)\]$"
    )
    SCHEMA_CLEANUP_ERROR_CODES = frozenset({
        "unknown_attribute",
        "attribute_disabled",
        "attribute_values_empty",
        "attribute_max_value_count",
        "attribute_not_collection",
        "dictionary_value_id_missing",
        "dictionary_value_out_of_scope",
        "dictionary_value_restricted",
        "dictionary_display_mismatch",
        "unexpected_dictionary_value_id",
        "attribute_type_mismatch",
        "unsupported_attribute_type",
        "attribute_value_not_string",
        "malformed_attribute_value",
        "complex_id_mismatch",
        "complex_attribute_outside_group",
        "simple_attribute_inside_group",
        "complex_attribute_repeated",
        "duplicate_attribute",
    })
    FIXED_COLUMNS: Sequence[Tuple[str, str]] = (
        ("action", "Действие"),
        ("draft_id", "ID черновика — не менять"),
        ("draft_version", "Версия — не менять"),
        ("imported_product_id", "ID товара — не менять"),
        ("offer_id", "Артикул Ozon — справочно"),
        ("title", "Название — справочно"),
        ("current_status", "Статус — справочно"),
        ("blocking_reason", "Что мешает загрузке — справочно"),
        ("price_rub", "Цена, ₽"),
        ("package_width_mm", "Ширина упаковки, мм"),
        ("package_height_mm", "Высота упаковки, мм"),
        ("package_depth_mm", "Длина упаковки, мм"),
        ("package_weight_g", "Вес с упаковкой, г"),
        ("description", "Описание"),
    )
    EDITABLE_FIXED = {
        "action",
        "price_rub",
        "package_width_mm",
        "package_height_mm",
        "package_depth_mm",
        "package_weight_g",
        "description",
    }
    ACTION_REPAIR = "ИСПРАВИТЬ"
    ACTION_EXCLUDE = "ИСКЛЮЧИТЬ"

    @staticmethod
    def _safe_text(value: Any, maximum: int = 500) -> str:
        if value is None:
            return ""
        rendered = str(value).strip()
        if len(rendered) > maximum:
            rendered = rendered[:maximum]
        return rendered

    @classmethod
    def _conflict(cls, code: str, message: str) -> OzonBulkUploadConflict:
        error = OzonBulkUploadConflict(message)
        error.code = code
        return error

    @classmethod
    def _owned_completed_run(cls, *, seller_id: int, job_uid: str):
        job = OzonBulkUploadService.get_run(
            seller_id=seller_id,
            job_uid=job_uid,
            reconcile=True,
        )
        if job.status == "running":
            raise cls._conflict(
                "ozon_repair_run_still_active",
                "Дождитесь завершения текущей загрузки перед массовым исправлением",
            )
        document = OzonBulkUploadService._load_progress(job)
        if not document:
            raise cls._conflict(
                "ozon_repair_run_corrupt",
                "Сводка запуска повреждена; черновики сохранены отдельно",
            )
        return job, document

    @classmethod
    def _repairable_items(cls, document: dict) -> List[dict]:
        return [
            item
            for item in document["items"]
            if (
                item.get("status") in cls.EXPORTABLE_STATUSES
                and isinstance(item.get("draft_id"), int)
                and not isinstance(item.get("draft_id"), bool)
                and item["draft_id"] > 0
            )
        ]

    @classmethod
    def _drafts(
        cls,
        *,
        seller_id: int,
        account_id: int,
        items: Sequence[dict],
    ) -> Dict[int, MarketplaceProductDraft]:
        draft_ids = [item["draft_id"] for item in items]
        rows = MarketplaceProductDraft.query.options(
            joinedload(MarketplaceProductDraft.account),
            joinedload(MarketplaceProductDraft.imported_product).joinedload(
                ImportedProduct.supplier_product
            ),
            joinedload(MarketplaceProductDraft.product_type).joinedload(
                MarketplaceProductType.category
            ),
            joinedload(MarketplaceProductDraft.published_listing),
        ).filter(
            MarketplaceProductDraft.seller_id == seller_id,
            MarketplaceProductDraft.account_id == account_id,
            MarketplaceProductDraft.id.in_(draft_ids),
        ).all()
        result = {row.id: row for row in rows}
        if len(result) != len(set(draft_ids)):
            raise cls._conflict(
                "ozon_repair_draft_scope_changed",
                "Один из черновиков удалён или больше не относится к этому кабинету",
            )
        for item in items:
            draft = result[item["draft_id"]]
            if (
                draft.imported_product_id != item.get("imported_product_id")
                or draft.account_id != account_id
                or draft.imported_product is None
                or draft.imported_product.seller_id != seller_id
            ):
                raise cls._conflict(
                    "ozon_repair_draft_identity_changed",
                    "Связь товара с черновиком изменилась; сформируйте новую таблицу",
                )
        return result

    @staticmethod
    def _stored_json(raw: Any, fallback: Any) -> Any:
        try:
            value = json.loads(raw or "")
        except (TypeError, ValueError):
            return fallback
        return value if isinstance(value, type(fallback)) else fallback

    @classmethod
    def _validation_errors(
        cls,
        draft: MarketplaceProductDraft,
        *,
        validation_result: Optional[dict] = None,
    ) -> List[dict]:
        value = (
            validation_result
            if isinstance(validation_result, dict)
            else cls._stored_json(draft.validation_result_json, {})
        )
        errors = value.get("errors") if isinstance(value, dict) else []
        if not isinstance(errors, list):
            return []
        return [
            item
            for item in errors
            if isinstance(item, dict)
        ][: MarketplaceDraftService.MAX_VALIDATION_ITEMS]

    @classmethod
    def _baseline_documents(
        cls,
        draft: MarketplaceProductDraft,
    ) -> Optional[dict]:
        """Reconstruct the already-scoped local listing snapshot, without I/O."""
        listing = draft.published_listing
        if (
            not isinstance(listing, MarketplaceListing)
            or listing.id != draft.published_listing_id
            or listing.seller_id != draft.seller_id
            or listing.account_id != draft.account_id
            or listing.imported_product_id != draft.imported_product_id
            or listing.offer_id != draft.offer_id
        ):
            return None
        try:
            state = OzonProductStateContract.from_listing_projection({
                "product_id": listing.external_product_id,
                "offer_id": listing.offer_id,
                "category_id": listing.external_category_id,
                "type_id": listing.external_type_id,
                "title": listing.title,
                "attributes": cls._stored_json(
                    listing.attributes_json,
                    [],
                ),
                "complex_attributes": cls._stored_json(
                    listing.complex_attributes_json,
                    [],
                ),
                "media": cls._stored_json(listing.media_json, {}),
                "dimensions": cls._stored_json(
                    listing.dimensions_json,
                    {},
                ),
                "barcodes": cls._stored_json(
                    listing.barcodes_json,
                    [],
                ),
                "price": cls._stored_json(
                    listing.price_summary_json,
                    {},
                ),
            })
            return OzonProductStateContract.draft_documents(
                state["payload"]
            )
        except OzonProductStateError:
            return None

    @classmethod
    def _schema_cleanup_candidates(
        cls,
        *,
        draft: MarketplaceProductDraft,
        definitions: Dict[
            Tuple[int, str],
            MarketplaceAttributeDefinition,
        ],
        validation_result: Optional[dict] = None,
    ) -> List[dict]:
        """Map current validation paths to exact, seller-reviewable removals."""
        if draft.published_listing_id is None:
            return []
        baseline = cls._baseline_documents(draft)
        if baseline is None:
            return []
        try:
            effective, _snapshot = (
                MarketplaceDraftService.publication_documents(draft)
            )
        except MarketplaceDraftError:
            effective = MarketplaceDraftService._stored_draft_documents(
                draft
            )

        path_items: Dict[str, dict] = {}
        for index, item in enumerate(effective.get("attributes", [])):
            if isinstance(item, dict):
                path_items[f"attributes[{index}]"] = item
        for group_index, group in enumerate(
            effective.get("complex_attributes", [])
        ):
            if not isinstance(group, dict):
                continue
            for item_index, item in enumerate(
                group.get("attributes", [])
            ):
                if isinstance(item, dict):
                    path_items[
                        "complex_attributes"
                        f"[{group_index}].attributes[{item_index}]"
                    ] = item

        def identities(documents: dict) -> set:
            result = {
                MarketplaceDraftService._attribute_identity(item)
                for item in documents.get("attributes", [])
            }
            for group in documents.get("complex_attributes", []):
                if not isinstance(group, dict):
                    continue
                result.update(
                    MarketplaceDraftService._attribute_identity(item)
                    for item in group.get("attributes", [])
                )
            result.discard(None)
            return result

        baseline_identities = identities(baseline)
        draft_identities = identities(
            MarketplaceDraftService._stored_draft_documents(draft)
        )
        stored_removals = {
            (
                item.get("attribute_id"),
                item.get("complex_id", "0"),
            )
            for item in cls._stored_json(
                draft.attribute_removals_json,
                [],
            )
            if (
                isinstance(item, dict)
                and isinstance(item.get("attribute_id"), str)
                and isinstance(item.get("complex_id", "0"), str)
            )
        }
        candidates: Dict[Tuple[str, str], dict] = {}

        def ensure(identity: Tuple[str, str]) -> dict:
            external_id, complex_id = identity
            definition = definitions.get(
                (draft.product_type_id, external_id)
            )
            candidate = candidates.get(identity)
            if candidate is None:
                candidate = {
                    "attribute_id": external_id,
                    "complex_id": complex_id,
                    "name": (
                        definition.name
                        if definition is not None and definition.name
                        else f"Атрибут Ozon {external_id}"
                    ),
                    "issues": [],
                    "in_live_baseline": identity in baseline_identities,
                    "in_draft_overlay": identity in draft_identities,
                    "already_confirmed": identity in stored_removals,
                    "replacement_required": bool(
                        definition is not None
                        and definition.is_required
                        and not definition.attribute_complex_id
                    ),
                }
                candidates[identity] = candidate
            return candidate

        for error in cls._validation_errors(
            draft,
            validation_result=validation_result,
        ):
            if error.get("code") not in cls.SCHEMA_CLEANUP_ERROR_CODES:
                continue
            field = error.get("field")
            if not isinstance(field, str):
                continue
            item = path_items.get(field)
            identity = MarketplaceDraftService._attribute_identity(item)
            if (
                identity is None
                or identity
                not in baseline_identities | draft_identities
            ):
                continue
            candidate = ensure(identity)
            message = cls._safe_text(error.get("message"), 240)
            if message and message not in candidate["issues"]:
                candidate["issues"].append(message)

        for identity in stored_removals:
            if identity in baseline_identities:
                candidate = ensure(identity)
                if not candidate["issues"]:
                    candidate["issues"].append(
                        "Удаление уже подтверждено продавцом"
                    )

        return sorted(
            candidates.values(),
            key=lambda item: (
                int(item["attribute_id"]),
                int(item["complex_id"]),
            ),
        )

    @classmethod
    def _missing_attribute_ids(
        cls,
        draft: MarketplaceProductDraft,
        *,
        validation_result: Optional[dict] = None,
    ) -> List[str]:
        result = []
        for error in cls._validation_errors(
            draft,
            validation_result=validation_result,
        ):
            field = error.get("field")
            if (
                error.get("code") == "required_attribute_missing"
                and isinstance(field, str)
                and field.startswith("attributes.")
            ):
                external_id = field.split(".", 1)[1]
                if external_id.isascii() and external_id.isdigit():
                    result.append(external_id)
        return list(dict.fromkeys(result))

    @classmethod
    def _definition_map(
        cls,
        drafts: Iterable[MarketplaceProductDraft],
    ) -> Dict[Tuple[int, str], MarketplaceAttributeDefinition]:
        product_type_ids = {
            draft.product_type_id
            for draft in drafts
            if draft.product_type_id is not None
        }
        if not product_type_ids:
            return {}
        rows = MarketplaceAttributeDefinition.query.filter(
            MarketplaceAttributeDefinition.product_type_id.in_(
                product_type_ids
            ),
            MarketplaceAttributeDefinition.is_available.is_(True),
            MarketplaceAttributeDefinition.is_enabled.is_(True),
        ).all()
        result: Dict[Tuple[int, str], MarketplaceAttributeDefinition] = {}
        duplicates = set()
        for row in rows:
            key = (row.product_type_id, row.external_attribute_id)
            if key in result:
                duplicates.add(key)
            result[key] = row
        for key in duplicates:
            result.pop(key, None)
        return result

    @classmethod
    def _current_attribute_values(
        cls,
        draft: MarketplaceProductDraft,
    ) -> Dict[str, List[str]]:
        result: Dict[str, List[str]] = {}
        for item in cls._stored_json(draft.attributes_json, []):
            if (
                not isinstance(item, dict)
                or item.get("complex_id", "0") != "0"
                or not isinstance(item.get("attribute_id"), str)
            ):
                continue
            values = item.get("values")
            if not isinstance(values, list):
                continue
            rendered = [
                value["value"]
                for value in values
                if (
                    isinstance(value, dict)
                    and isinstance(value.get("value"), str)
                    and value["value"].strip()
                )
            ]
            if rendered:
                result[item["attribute_id"]] = rendered
        return result

    @staticmethod
    def _set_excel_value(cell, value: Any) -> None:
        """Write text as text without changing the seller value.

        ``openpyxl`` otherwise interprets a leading ``=`` as a formula.  A
        prefixed apostrophe is not suitable here because it becomes part of
        the value on the XLSX round-trip and could leak into the draft.
        """
        cell.value = value
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            cell.data_type = "s"

    @staticmethod
    def _decimal_text(value: Any) -> str:
        if value in (None, ""):
            return ""
        try:
            number = Decimal(str(value))
        except (InvalidOperation, TypeError, ValueError):
            return ""
        if not number.is_finite():
            return ""
        return format(number.normalize(), "f")

    @classmethod
    def _physical_export_values(cls, draft: MarketplaceProductDraft) -> dict:
        dimensions = cls._stored_json(draft.dimensions_json, {})
        dimension_factors = {
            "MILLIMETERS": Decimal("1"),
            "CENTIMETERS": Decimal("10"),
            "INCHES": Decimal("25.4"),
        }
        weight_factors = {
            "GRAMS": Decimal("1"),
            "KILOGRAMS": Decimal("1000"),
            "POUNDS": Decimal("453.59237"),
        }
        dimension_factor = dimension_factors.get(
            dimensions.get("dimension_unit")
        )
        weight_factor = weight_factors.get(dimensions.get("weight_unit"))
        result = {}

        def converted(raw_value: Any, factor: Optional[Decimal]) -> str:
            if factor is None or raw_value in (None, ""):
                return ""
            try:
                return cls._decimal_text(Decimal(str(raw_value)) * factor)
            except (InvalidOperation, TypeError, ValueError):
                return ""

        for source, target in (
            ("width", "package_width_mm"),
            ("height", "package_height_mm"),
            ("depth", "package_depth_mm"),
        ):
            result[target] = converted(
                dimensions.get(source),
                dimension_factor,
            )
        result["package_weight_g"] = converted(
            dimensions.get("weight"),
            weight_factor,
        )
        return result

    @classmethod
    def _dictionary_values(
        cls,
        definition: MarketplaceAttributeDefinition,
    ) -> List[str]:
        if (
            not definition.dictionary_id
            or not OzonReferenceService.dictionary_is_fresh(definition)
        ):
            return []
        query = MarketplaceAttributeValue.query.filter_by(
            attribute_id=definition.id,
            is_available=True,
        )
        restriction = set(definition.restriction_value_ids)
        if restriction:
            query = query.filter(
                MarketplaceAttributeValue.external_value_id.in_(
                    restriction
                )
            )
        return [
            row.value
            for row in query.order_by(
                MarketplaceAttributeValue.value_normalized.asc(),
                MarketplaceAttributeValue.id.asc(),
            ).limit(5_001).all()
            if isinstance(row.value, str) and row.value.strip()
        ][:5_000]

    @classmethod
    def editor_document(
        cls,
        *,
        seller_id: int,
        job_uid: str,
    ) -> dict:
        """Return a bounded platform-native editor model for one upload run."""
        job, document = cls._owned_completed_run(
            seller_id=seller_id,
            job_uid=job_uid,
        )
        items = cls._repairable_items(document)
        if not items:
            raise OzonBulkUploadValidationError(
                "В запуске нет карточек для массового исправления"
            )
        drafts = cls._drafts(
            seller_id=seller_id,
            account_id=document["account_id"],
            items=items,
        )
        definitions = cls._definition_map(drafts.values())
        groups: Dict[Tuple[Any, ...], dict] = {}
        bulk_attribute_labels: Dict[str, set] = {}
        editor_attribute_cells = 0
        has_schema_cleanup = False

        for item in items:
            draft = drafts[item["draft_id"]]
            # Validation paths for attributes are positional.  Always derive
            # them from the same current local listing projection used below
            # to resolve exact cleanup identities.  A catalog refresh can
            # change that projection without changing the draft version, so a
            # stored older result must never authorize removal by a stale
            # ``attributes[N]`` path.
            validation_result = (
                MarketplaceDraftService._build_validation_result(draft)
            )
            errors = cls._validation_errors(
                draft,
                validation_result=validation_result,
            )
            cleanup_candidates = cls._schema_cleanup_candidates(
                draft=draft,
                definitions=definitions,
                validation_result=validation_result,
            )
            has_schema_cleanup = bool(
                has_schema_cleanup or cleanup_candidates
            )
            cleanup_replacement_ids = {
                candidate["attribute_id"]
                for candidate in cleanup_candidates
                if candidate["replacement_required"]
            }
            error_codes = {
                error.get("code")
                for error in errors
                if isinstance(error.get("code"), str)
            }
            action = (
                cls.ACTION_EXCLUDE
                if (
                    item.get("status") == "excluded"
                    or "ozon_brand_forbidden" in error_codes
                )
                else cls.ACTION_REPAIR
            )
            content = cls._stored_json(draft.content_json, {})
            commercial = cls._stored_json(draft.commercial_json, {})
            current_attributes = cls._current_attribute_values(draft)
            raw_attributes = cls._stored_json(draft.attributes_json, [])
            dictionary_ids = {}
            for attribute in raw_attributes:
                if (
                    not isinstance(attribute, dict)
                    or attribute.get("complex_id", "0") != "0"
                    or not isinstance(attribute.get("attribute_id"), str)
                    or not isinstance(attribute.get("values"), list)
                    or len(attribute["values"]) != 1
                    or not isinstance(attribute["values"][0], dict)
                    or not isinstance(
                        attribute["values"][0].get(
                            "dictionary_value_id"
                        ),
                        str,
                    )
                ):
                    continue
                dictionary_ids[attribute["attribute_id"]] = (
                    attribute["values"][0]["dictionary_value_id"]
                )

            attribute_fields = []
            editable_ids = list(dict.fromkeys([
                *cls._missing_attribute_ids(
                    draft,
                    validation_result=validation_result,
                ),
                *sorted(
                    cleanup_replacement_ids,
                    key=int,
                ),
            ]))
            for external_id in editable_ids:
                editor_attribute_cells += 1
                if (
                    editor_attribute_cells
                    > cls.MAX_EDITOR_ATTRIBUTE_CELLS
                ):
                    raise cls._conflict(
                        "ozon_repair_editor_too_many_fields",
                        "В запуске слишком много разных полей для одного "
                        "массового экрана; разделите карточки на меньшие партии",
                    )
                definition = definitions.get(
                    (draft.product_type_id, external_id)
                )
                replaces_invalid = (
                    external_id in cleanup_replacement_ids
                )
                name = (
                    definition.name
                    if definition is not None and definition.name
                    else f"Атрибут Ozon {external_id}"
                )
                if (
                    definition is not None
                    and not definition.attribute_complex_id
                ):
                    bulk_attribute_labels.setdefault(
                        external_id,
                        set(),
                    ).add(name)
                attribute_fields.append({
                    "external_id": external_id,
                    "name": name,
                    "value": (
                        ""
                        if replaces_invalid
                        else "\n".join(
                            current_attributes.get(external_id, [])
                        )
                    ),
                    "dictionary_value_id": (
                        ""
                        if replaces_invalid
                        else dictionary_ids.get(external_id, "")
                    ),
                    "editable": bool(
                        definition is not None
                        and not definition.attribute_complex_id
                    ),
                    "dictionary": bool(
                        definition is not None
                        and definition.dictionary_id
                    ),
                    "dictionary_fresh": bool(
                        definition is not None
                        and (
                            not definition.dictionary_id
                            or OzonReferenceService.dictionary_is_fresh(
                                definition
                            )
                        )
                    ),
                    "data_type": (
                        definition.data_type
                        if definition is not None else None
                    ),
                    "is_collection": bool(
                        definition is not None
                        and definition.is_collection
                    ),
                    "max_values": min(
                        (definition.max_value_count or cls.MAX_ATTRIBUTE_VALUES)
                        if definition is not None and definition.is_collection else 1,
                        cls.MAX_ATTRIBUTE_VALUES,
                    ),
                    "replaces_invalid": replaces_invalid,
                    "suggestions": (
                        OzonComplianceSuggestionService.tnved_suggestions(
                            draft=draft,
                            definition=definition,
                        )
                        if (
                            definition is not None
                            and external_id
                            == OzonComplianceSuggestionService.TNVED_ATTRIBUTE_ID
                        )
                        else []
                    ),
                })

            try:
                identities = (
                    MarketplaceDraftService._mapping_identities(
                        draft.imported_product
                    )
                    if draft.imported_product is not None
                    else []
                )
            except MarketplaceDraftError:
                identities = []
            if identities:
                identity = identities[0]
                group_identity = (
                    identity["scope_key"],
                    identity["source_category_normalized"],
                )
                group_label = (
                    identity.get("source_category")
                    or "Исходная категория"
                )
            else:
                group_identity = ("draft", draft.id)
                group_label = "Без исходной категории"
            group = groups.get(group_identity)
            if group is None:
                group = {
                    "id": f"group-{len(groups) + 1}",
                    "source_category": group_label,
                    "rows": [],
                    "_product_types": {},
                    "_type_ids": [],
                }
                groups[group_identity] = group
            group["_type_ids"].append(draft.product_type_id)
            if draft.product_type is not None:
                group["_product_types"][draft.product_type.id] = {
                    "id": draft.product_type.id,
                    "name": draft.product_type.name,
                    "category_path": (
                        draft.product_type.category.full_path
                        if draft.product_type.category is not None
                        else ""
                    ),
                    "schema_fresh": (
                        OzonReferenceService.reference_is_fresh(
                            draft.product_type
                        )
                    ),
                }
            group["rows"].append({
                "draft_id": draft.id,
                "draft_version": draft.version,
                "imported_product_id": draft.imported_product_id,
                "offer_id": draft.offer_id,
                "title": (
                    draft.imported_product.title
                    if draft.imported_product is not None
                    else item.get("title") or ""
                ),
                "status": item.get("status"),
                "message": item.get("message") or "",
                "repair_error": item.get("repair_error") or "",
                "action": action,
                "product_type_id": draft.product_type_id,
                "product_type_name": draft.product_type.name if draft.product_type else "",
                "photo_url": next(iter(imported_photo_previews(draft.imported_product).values()), None),
                "type_change_impact": {
                    "attributes": len(raw_attributes) if isinstance(raw_attributes, list) else 0,
                    "complex_groups": len(cls._stored_json(draft.complex_attributes_json, [])),
                    "removals": len(cls._stored_json(draft.attribute_removals_json, [])),
                },
                "price_rub": commercial.get("price", ""),
                "description": content.get("description", ""),
                **cls._physical_export_values(draft),
                "attributes": attribute_fields,
                "cleanup_candidates": cleanup_candidates,
                "schema_cleanup": bool(
                    cleanup_candidates
                    and all(
                        candidate["already_confirmed"]
                        for candidate in cleanup_candidates
                    )
                ),
                "validation_error_count": len(errors),
                "validation_errors_truncated": (
                    len(errors) > cls.MAX_EDITOR_VALIDATION_ERRORS
                ),
                "validation_errors": [
                    {
                        "code": cls._safe_text(
                            error.get("code"),
                            80,
                        ),
                        "field": cls._safe_text(
                            error.get("field"),
                            120,
                        ),
                        "message": cls._safe_text(
                            error.get("message"),
                            300,
                        ),
                    }
                    for error in errors[
                        : cls.MAX_EDITOR_VALIDATION_ERRORS
                    ]
                ],
            })

        public_groups = []
        for group in groups.values():
            product_types = list(group.pop("_product_types").values())
            type_ids = set(group.pop("_type_ids"))
            group["product_type"] = (
                product_types[0]
                if len(type_ids) == 1 and None not in type_ids
                else None
            )
            group["mixed_product_types"] = len(type_ids) > 1
            public_groups.append(group)

        bulk_fields = [
            {"key": "price_rub", "label": "Цена продавца, ₽"},
            {
                "key": "package_width_mm",
                "label": "Ширина упаковки, мм",
            },
            {
                "key": "package_height_mm",
                "label": "Высота упаковки, мм",
            },
            {
                "key": "package_depth_mm",
                "label": "Длина упаковки, мм",
            },
            {
                "key": "package_weight_g",
                "label": "Вес с упаковкой, г",
            },
            {"key": "description", "label": "Описание"},
        ]
        if has_schema_cleanup:
            bulk_fields.append({
                "key": "schema_cleanup",
                "label": (
                    "Очистить несовместимые атрибуты "
                    "(да / нет)"
                ),
            })
        for external_id in sorted(
            bulk_attribute_labels,
            key=lambda value: int(value),
        ):
            labels = sorted(bulk_attribute_labels[external_id])
            label = labels[0] if len(labels) == 1 else " / ".join(
                labels[:2]
            )
            bulk_fields.append({
                "key": f"attribute:{external_id}",
                "label": label,
            })
        if len(bulk_fields) > cls.MAX_COLUMNS:
            raise cls._conflict(
                "ozon_repair_editor_too_many_columns",
                "Для выбранных типов слишком много разных обязательных "
                "полей; разделите карточки на более однородные партии",
            )
        return {
            "job_uid": job.job_uid,
            "account_id": document["account_id"],
            "account_label": document.get("account_label") or "",
            "groups": public_groups,
            "bulk_fields": bulk_fields,
            "total": sum(len(group["rows"]) for group in public_groups),
        }

    @classmethod
    def search_dictionary_values(
        cls,
        *,
        seller_id: int,
        job_uid: str,
        draft_id: int,
        external_attribute_id: str,
        query: Any = "",
        limit: int = 30,
    ) -> dict:
        """Search one fresh exact type-scoped dictionary locally."""
        _job, document = cls._owned_completed_run(
            seller_id=seller_id,
            job_uid=job_uid,
        )
        draft_id = cls._positive_integer_cell(draft_id, "draft_id")
        if (
            not isinstance(external_attribute_id, str)
            or cls.ATTRIBUTE_KEY.fullmatch(
                f"attribute:{external_attribute_id}"
            ) is None
        ):
            raise OzonBulkUploadValidationError(
                "Некорректный ID атрибута Ozon"
            )
        item = next(
            (
                value
                for value in cls._repairable_items(document)
                if value.get("draft_id") == draft_id
            ),
            None,
        )
        if item is None:
            raise cls._conflict(
                "ozon_repair_row_not_in_run",
                "Черновик отсутствует среди исправляемых строк запуска",
            )
        draft = cls._drafts(
            seller_id=seller_id,
            account_id=document["account_id"],
            items=[item],
        )[draft_id]
        # The editor derives its fields from current local facts. Search must
        # use that same projection, including replacements after schema drift;
        # an older (or absent) saved validation is not the current form.
        validation_result = MarketplaceDraftService._build_validation_result(draft)
        editable_ids = set(cls._missing_attribute_ids(
            draft, validation_result=validation_result,
        ))
        editable_ids.update(
            candidate["attribute_id"]
            for candidate in cls._schema_cleanup_candidates(
                draft=draft, definitions=cls._definition_map([draft]),
                validation_result=validation_result,
            )
            if candidate["replacement_required"]
        )
        if external_attribute_id not in editable_ids:
            raise OzonBulkUploadValidationError(
                "Атрибут не относится к недостающим полям этой строки"
            )
        definitions = MarketplaceAttributeDefinition.query.filter_by(
            product_type_id=draft.product_type_id,
            external_attribute_id=external_attribute_id,
            is_available=True,
            is_enabled=True,
        ).all()
        if (
            len(definitions) != 1
            or definitions[0].attribute_complex_id
            or not definitions[0].dictionary_id
        ):
            raise OzonBulkUploadValidationError(
                "Для этого поля нет простого официального справочника"
            )
        definition = definitions[0]
        if not OzonReferenceService.dictionary_is_fresh(definition):
            raise cls._conflict(
                "ozon_repair_dictionary_stale",
                "Справочник Ozon ещё загружается; повторите позже",
            )
        if (
            isinstance(query, bool)
            or query is None
            or not isinstance(query, str)
            or len(query.strip()) > 200
        ):
            raise OzonBulkUploadValidationError(
                "Поиск по справочнику должен быть строкой до 200 символов"
            )
        if (
            isinstance(limit, bool)
            or not isinstance(limit, int)
            or limit < 1
            or limit > 50
        ):
            raise OzonBulkUploadValidationError(
                "limit должен быть целым числом от 1 до 50"
            )
        values = MarketplaceAttributeValue.query.filter_by(
            attribute_id=definition.id,
            is_available=True,
        )
        restriction = set(definition.restriction_value_ids)
        if restriction:
            values = values.filter(
                MarketplaceAttributeValue.external_value_id.in_(
                    restriction
                )
            )
        search = query.strip()
        if search:
            escaped = (
                search.replace("\\", "\\\\")
                .replace("%", "\\%")
                .replace("_", "\\_")
            )
            pattern = f"%{escaped}%"
            values = values.filter(or_(
                MarketplaceAttributeValue.value.ilike(
                    pattern,
                    escape="\\",
                ),
                MarketplaceAttributeValue.external_value_id.ilike(
                    pattern,
                    escape="\\",
                ),
            ))
        rows = values.order_by(
            MarketplaceAttributeValue.value_normalized.asc(),
            MarketplaceAttributeValue.id.asc(),
        ).limit(limit + 1).all()
        return {
            "items": [
                {
                    "id": row.external_value_id,
                    "value": row.value,
                }
                for row in rows[:limit]
            ],
            "truncated": len(rows) > limit,
        }

    @classmethod
    def search_product_types(
        cls,
        *,
        seller_id: int,
        job_uid: str,
        query: Any,
        limit: int = 30,
    ) -> list:
        cls._owned_completed_run(
            seller_id=seller_id,
            job_uid=job_uid,
        )
        if (
            isinstance(query, bool)
            or not isinstance(query, str)
            or not query.strip()
            or len(query.strip()) > 200
        ):
            raise OzonBulkUploadValidationError(
                "Введите название типа Ozon для поиска"
            )
        return MarketplaceDraftService.search_product_types(
            seller_id=seller_id,
            query=query.strip(),
            limit=limit,
        )

    @classmethod
    def export_workbook(
        cls,
        *,
        seller_id: int,
        job_uid: str,
    ) -> Tuple[str, bytes]:
        job, document = cls._owned_completed_run(
            seller_id=seller_id,
            job_uid=job_uid,
        )
        items = cls._repairable_items(document)
        if not items:
            raise OzonBulkUploadValidationError(
                "В запуске нет карточек для массового исправления"
            )
        drafts = cls._drafts(
            seller_id=seller_id,
            account_id=document["account_id"],
            items=items,
        )
        definitions = cls._definition_map(drafts.values())

        missing_by_draft = {
            draft.id: cls._missing_attribute_ids(draft)
            for draft in drafts.values()
        }
        attribute_ids = sorted(
            {
                external_id
                for values in missing_by_draft.values()
                for external_id in values
            },
            key=lambda value: (
                {
                    "22232": 0,
                    "23536": 1,
                    "9048": 2,
                }.get(value, 3),
                int(value),
            ),
        )
        labels = {}
        for external_id in attribute_ids:
            names = sorted({
                definition.name
                for (product_type_id, attribute_id), definition
                in definitions.items()
                if attribute_id == external_id and definition.name
            })
            label = names[0] if len(names) == 1 else (
                " / ".join(names[:2]) if names else "Обязательный атрибут"
            )
            labels[external_id] = (
                f"{label} · Ozon {external_id}"
            )[:240]

        columns = list(cls.FIXED_COLUMNS) + [
            (f"attribute:{external_id}", labels[external_id])
            for external_id in attribute_ids
        ]
        if len(columns) > cls.MAX_COLUMNS:
            raise cls._conflict(
                "ozon_repair_too_many_columns",
                "Для выбранных типов слишком много разных обязательных полей; "
                "разделите запуск на более однородные группы",
            )

        workbook = Workbook()
        cards = workbook.active
        cards.title = cls.CARD_SHEET
        lists = workbook.create_sheet(cls.LIST_SHEET)
        meta = workbook.create_sheet(cls.META_SHEET)

        header_fill = PatternFill("solid", fgColor="17324D")
        editable_fill = PatternFill("solid", fgColor="FFF4CC")
        locked_fill = PatternFill("solid", fgColor="E8EDF2")
        blocked_fill = PatternFill("solid", fgColor="FDE8E8")
        ready_fill = PatternFill("solid", fgColor="E7F7ED")
        white_font = Font(color="FFFFFF", bold=True)

        for index, (key, label) in enumerate(columns, start=1):
            key_cell = cards.cell(row=1, column=index, value=key)
            key_cell.font = Font(color="6B7280", size=8)
            label_cell = cards.cell(row=2, column=index, value=label)
            label_cell.fill = header_fill
            label_cell.font = white_font
            label_cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            if key == "action":
                label_cell.comment = Comment(
                    "ИСПРАВИТЬ — применить заполненные данные локально. "
                    "ИСКЛЮЧИТЬ — убрать строку только из повторов этого запуска. "
                    "Импорт таблицы ничего не отправляет в Ozon.",
                    "Seller Hub",
                )
            elif key.startswith("attribute:"):
                label_cell.comment = Comment(
                    "Для справочного поля выбирайте точное официальное "
                    "значение из выпадающего списка. Если значений несколько, "
                    "укажите каждое с новой строки.",
                    "Seller Hub",
                )
        cards.row_dimensions[1].hidden = True
        cards.row_dimensions[2].height = 44
        cards.freeze_panes = "A3"
        cards.auto_filter.ref = (
            f"A2:{get_column_letter(len(columns))}{len(items) + 2}"
        )

        item_by_draft = {item["draft_id"]: item for item in items}
        validation_targets: Dict[
            Tuple[int, str],
            List[str],
        ] = {}
        for row_index, item in enumerate(items, start=3):
            draft = drafts[item["draft_id"]]
            errors = cls._validation_errors(draft)
            error_codes = {
                error.get("code")
                for error in errors
                if isinstance(error.get("code"), str)
            }
            action = (
                cls.ACTION_EXCLUDE
                if (
                    item.get("status") == "excluded"
                    or "ozon_brand_forbidden" in error_codes
                )
                else cls.ACTION_REPAIR
            )
            content = cls._stored_json(draft.content_json, {})
            commercial = cls._stored_json(draft.commercial_json, {})
            physical = cls._physical_export_values(draft)
            current_attributes = cls._current_attribute_values(draft)
            blocking_reason = " | ".join(
                cls._safe_text(error.get("message"), 240)
                for error in errors[:8]
                if cls._safe_text(error.get("message"), 240)
            )
            values: Dict[str, Any] = {
                "action": action,
                "draft_id": draft.id,
                "draft_version": draft.version,
                "imported_product_id": draft.imported_product_id,
                "offer_id": draft.offer_id,
                "title": (
                    draft.imported_product.title
                    if draft.imported_product else ""
                ),
                "current_status": item.get("status"),
                "blocking_reason": blocking_reason,
                "price_rub": commercial.get("price", ""),
                "description": content.get("description", ""),
                **physical,
            }
            for external_id in attribute_ids:
                values[f"attribute:{external_id}"] = "\n".join(
                    current_attributes.get(external_id, [])
                )

            for column_index, (key, _label) in enumerate(columns, start=1):
                cell = cards.cell(
                    row=row_index,
                    column=column_index,
                )
                cls._set_excel_value(cell, values.get(key, ""))
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=key in {
                        "title",
                        "blocking_reason",
                        "description",
                    } or key.startswith("attribute:"),
                )
                editable = key in cls.EDITABLE_FIXED
                if key.startswith("attribute:"):
                    external_id = key.split(":", 1)[1]
                    definition = definitions.get(
                        (draft.product_type_id, external_id)
                    )
                    editable = bool(
                        definition is not None
                        and not definition.attribute_complex_id
                    )
                    if editable and definition.dictionary_id:
                        validation_targets.setdefault(
                            (definition.id, external_id),
                            [],
                        ).append(cell.coordinate)
                cell.fill = editable_fill if editable else locked_fill
                cell.protection = Protection(locked=not editable)
                if key in {
                    "price_rub",
                    "package_width_mm",
                    "package_height_mm",
                    "package_depth_mm",
                    "package_weight_g",
                }:
                    cell.number_format = "0.########"
                elif key.startswith("attribute:"):
                    cell.number_format = "@"
            if action == cls.ACTION_EXCLUDE:
                cards.cell(row=row_index, column=1).fill = blocked_fill
            elif not errors:
                cards.cell(row=row_index, column=1).fill = ready_fill

        action_validation = DataValidation(
            type="list",
            formula1=f'"{cls.ACTION_REPAIR},{cls.ACTION_EXCLUDE}"',
            allow_blank=False,
            error="Выберите ИСПРАВИТЬ или ИСКЛЮЧИТЬ",
            errorTitle="Неизвестное действие",
        )
        cards.add_data_validation(action_validation)
        action_validation.add(f"A3:A{len(items) + 2}")

        list_column = 1
        definition_by_id = {
            definition.id: definition
            for definition in definitions.values()
        }
        for (definition_id, external_id), coordinates in sorted(
            validation_targets.items()
        ):
            definition = definition_by_id[definition_id]
            official_values = cls._dictionary_values(definition)
            if not official_values:
                continue
            list_name = f"L_{definition.id}_{external_id}"
            lists.cell(row=1, column=list_column, value=list_name)
            for value_index, value in enumerate(official_values, start=2):
                cell = lists.cell(
                    row=value_index,
                    column=list_column,
                )
                cls._set_excel_value(cell, value)
                cell.number_format = "@"
            reference = (
                f"{quote_sheetname(cls.LIST_SHEET)}!"
                f"${get_column_letter(list_column)}$2:"
                f"${get_column_letter(list_column)}${len(official_values) + 1}"
            )
            workbook.defined_names.add(
                DefinedName(list_name, attr_text=reference)
            )
            validation = DataValidation(
                type="list",
                formula1=f"={list_name}",
                allow_blank=True,
                error="Выберите точное официальное значение из списка",
                errorTitle="Значение вне справочника Ozon",
            )
            cards.add_data_validation(validation)
            for coordinate in coordinates:
                validation.add(coordinate)
            list_column += 1

        cards.conditional_formatting.add(
            f"A3:A{len(items) + 2}",
            FormulaRule(
                formula=[f'$A3="{cls.ACTION_EXCLUDE}"'],
                fill=blocked_fill,
            ),
        )
        widths = {
            "action": 15,
            "draft_id": 13,
            "draft_version": 12,
            "imported_product_id": 13,
            "offer_id": 24,
            "title": 38,
            "current_status": 16,
            "blocking_reason": 52,
            "price_rub": 14,
            "package_width_mm": 17,
            "package_height_mm": 17,
            "package_depth_mm": 17,
            "package_weight_g": 18,
            "description": 56,
        }
        for index, (key, _label) in enumerate(columns, start=1):
            cards.column_dimensions[get_column_letter(index)].width = (
                widths.get(key, 28)
            )

        meta_values = {
            "contract_version": cls.CONTRACT_VERSION,
            "job_uid": job.job_uid,
            "seller_id": seller_id,
            "account_id": document["account_id"],
            "exported_at": datetime.utcnow().isoformat(),
            "columns_json": json.dumps(
                [key for key, _label in columns],
                ensure_ascii=False,
                separators=(",", ":"),
            ),
        }
        for row_index, (key, value) in enumerate(
            meta_values.items(),
            start=1,
        ):
            meta.cell(row=row_index, column=1, value=key)
            meta.cell(row=row_index, column=2, value=value)
        lists.sheet_state = "hidden"
        meta.sheet_state = "hidden"

        output = BytesIO()
        workbook.save(output)
        payload = output.getvalue()
        if len(payload) > cls.MAX_FILE_BYTES:
            raise cls._conflict(
                "ozon_repair_workbook_too_large",
                "Таблица исправлений превысила safety limit; разделите запуск",
            )
        return f"ozon-repair-{job.job_uid}.xlsx", payload

    @classmethod
    def _load_workbook(cls, payload: bytes):
        if not isinstance(payload, bytes) or not payload:
            raise OzonBulkUploadValidationError(
                "Загрузите непустой файл XLSX"
            )
        if len(payload) > cls.MAX_FILE_BYTES:
            raise OzonBulkUploadValidationError(
                "Файл XLSX больше допустимых 2 МБ"
            )
        stream = BytesIO(payload)
        if not zipfile.is_zipfile(stream):
            raise OzonBulkUploadValidationError(
                "Ожидался настоящий файл XLSX"
            )
        stream.seek(0)
        try:
            with zipfile.ZipFile(stream) as archive:
                members = archive.infolist()
                if (
                    len(members) > cls.MAX_ZIP_MEMBERS
                    or sum(member.file_size for member in members)
                    > cls.MAX_UNCOMPRESSED_BYTES
                    or any(
                        member.flag_bits & 0x1
                        or member.filename.startswith(("/", "\\"))
                        or ".." in member.filename.replace("\\", "/").split("/")
                        for member in members
                    )
                ):
                    raise OzonBulkUploadValidationError(
                        "Небезопасная или слишком большая структура XLSX"
                    )
        except zipfile.BadZipFile:
            raise OzonBulkUploadValidationError(
                "Повреждённый файл XLSX"
            ) from None
        stream.seek(0)
        try:
            workbook = load_workbook(
                stream,
                read_only=False,
                data_only=False,
                keep_links=False,
            )
        except Exception:
            raise OzonBulkUploadValidationError(
                "Не удалось прочитать XLSX; скачайте новую таблицу и повторите"
            ) from None
        if (
            cls.CARD_SHEET not in workbook.sheetnames
            or cls.META_SHEET not in workbook.sheetnames
        ):
            raise OzonBulkUploadValidationError(
                "XLSX не является таблицей исправлений Seller Hub"
            )
        return workbook

    @staticmethod
    def _metadata(workbook) -> dict:
        sheet = workbook[OzonBulkRepairService.META_SHEET]
        result = {}
        for row in sheet.iter_rows(min_col=1, max_col=2, values_only=True):
            if isinstance(row[0], str):
                result[row[0]] = row[1]
        return result

    @staticmethod
    def _positive_integer_cell(value: Any, field_name: str) -> int:
        if isinstance(value, bool):
            raise MarketplaceDraftValidationError(
                f"{field_name} должен быть положительным целым числом"
            )
        if isinstance(value, int):
            parsed = value
        elif isinstance(value, float) and value.is_integer():
            parsed = int(value)
        elif (
            isinstance(value, str)
            and value.strip().isascii()
            and value.strip().isdigit()
        ):
            parsed = int(value.strip())
        else:
            raise MarketplaceDraftValidationError(
                f"{field_name} должен быть положительным целым числом"
            )
        if parsed <= 0:
            raise MarketplaceDraftValidationError(
                f"{field_name} должен быть положительным целым числом"
            )
        return parsed

    @staticmethod
    def _cell_text(value: Any, field_name: str, maximum: int) -> str:
        if value in (None, ""):
            return ""
        if isinstance(value, bool):
            raise MarketplaceDraftValidationError(
                f"{field_name} содержит недопустимое значение"
            )
        rendered = str(value).strip()
        if not rendered:
            return ""
        if len(rendered) > maximum:
            raise MarketplaceDraftValidationError(
                f"{field_name} длиннее {maximum} символов"
            )
        return rendered

    @classmethod
    def _positive_decimal_cell(
        cls,
        value: Any,
        field_name: str,
        *,
        integer: bool = False,
    ) -> str:
        rendered = cls._cell_text(value, field_name, 80).replace(",", ".")
        try:
            parsed = Decimal(rendered)
        except (InvalidOperation, TypeError, ValueError):
            raise MarketplaceDraftValidationError(
                f"{field_name} должен быть положительным числом"
            ) from None
        if (
            not parsed.is_finite()
            or parsed <= 0
            or (integer and parsed != parsed.to_integral_value())
        ):
            suffix = " целым числом" if integer else " числом"
            raise MarketplaceDraftValidationError(
                f"{field_name} должен быть положительным{suffix}"
            )
        return format(parsed.normalize(), "f")

    @classmethod
    def _resolve_attribute_values(
        cls,
        *,
        definition: MarketplaceAttributeDefinition,
        raw_value: Any,
        expected_external_value_id: Any = None,
    ) -> list:
        rendered = cls._cell_text(
            raw_value,
            f"атрибут {definition.external_attribute_id}",
            10_000,
        )
        if not rendered:
            return []
        values = [
            value.strip()
            for value in rendered.replace("\r\n", "\n").split("\n")
            if value.strip()
        ]
        values = list(dict.fromkeys(values))
        if (
            not values
            or len(values) > cls.MAX_ATTRIBUTE_VALUES
            or (
                not definition.is_collection
                and len(values) > 1
            )
            or (
                definition.max_value_count
                and len(values) > definition.max_value_count
            )
        ):
            raise MarketplaceDraftValidationError(
                f"Некорректное число значений для «{definition.name}»"
            )

        if not definition.dictionary_id:
            if expected_external_value_id not in (None, ""):
                raise MarketplaceDraftValidationError(
                    f"У «{definition.name}» нет dictionary value ID"
                )
            normalized = MarketplaceDraftService._value_strings(values)
            if len(normalized) != len(values):
                raise MarketplaceDraftValidationError(
                    f"Некорректное значение для «{definition.name}»"
                )
            return [{"value": value} for value in normalized]

        if not OzonReferenceService.dictionary_is_fresh(definition):
            raise MarketplaceDraftValidationError(
                f"Справочник «{definition.name}» ещё не синхронизирован"
            )
        expected_id = ""
        if expected_external_value_id not in (None, ""):
            expected_id = cls._cell_text(
                expected_external_value_id,
                f"ID значения «{definition.name}»",
                200,
            )
            if len(values) != 1:
                raise MarketplaceDraftValidationError(
                    f"ID значения «{definition.name}» допустим только "
                    "для одного выбранного значения"
                )
        normalized_values = {
            OzonReferenceService.normalize_value(value)
            for value in values
        }
        query = MarketplaceAttributeValue.query.filter(
            MarketplaceAttributeValue.attribute_id == definition.id,
            MarketplaceAttributeValue.value_normalized.in_(
                normalized_values
            ),
            MarketplaceAttributeValue.is_available.is_(True),
        )
        if expected_id:
            query = query.filter(
                MarketplaceAttributeValue.external_value_id == expected_id
            )
        rows = query.all()
        by_normalized: Dict[str, List[MarketplaceAttributeValue]] = {}
        for row in rows:
            by_normalized.setdefault(row.value_normalized, []).append(row)
        restriction = set(definition.restriction_value_ids)
        result = []
        for value in values:
            normalized = OzonReferenceService.normalize_value(value)
            matches = by_normalized.get(normalized, [])
            if len(matches) != 1:
                raise MarketplaceDraftValidationError(
                    f"«{value}» не является одним точным официальным "
                    f"значением «{definition.name}»"
                )
            row = matches[0]
            if restriction and row.external_value_id not in restriction:
                raise MarketplaceDraftValidationError(
                    f"«{value}» запрещено для «{definition.name}» этого типа"
                )
            result.append({
                "dictionary_value_id": row.external_value_id,
                "value": row.value,
            })
        return result

    @classmethod
    def _row_patch(
        cls,
        *,
        draft: MarketplaceProductDraft,
        values: Dict[str, Any],
        attribute_keys: Sequence[str],
        schema_cleanup: Optional[bool] = None,
        cleanup_candidates: Sequence[dict] = (),
    ) -> dict:
        patch = {}
        if schema_cleanup is not None and not isinstance(
            schema_cleanup,
            bool,
        ):
            raise MarketplaceDraftValidationError(
                "schema_cleanup должен быть boolean или отсутствовать"
            )
        cleanup_identities = {
            (
                candidate["attribute_id"],
                candidate.get("complex_id", "0"),
            )
            for candidate in cleanup_candidates
            if (
                isinstance(candidate, dict)
                and isinstance(candidate.get("attribute_id"), str)
                and isinstance(candidate.get("complex_id", "0"), str)
            )
        }
        if schema_cleanup and not cleanup_identities:
            raise MarketplaceDraftValidationError(
                "Для карточки нет подтверждённых кандидатов на очистку"
            )
        current_removals = cls._stored_json(
            draft.attribute_removals_json,
            [],
        )
        desired_removals = (
            [
                {
                    "attribute_id": candidate["attribute_id"],
                    "complex_id": candidate.get("complex_id", "0"),
                }
                for candidate in cleanup_candidates
                if candidate.get("in_live_baseline") is True
            ]
            if schema_cleanup
            else (
                current_removals
                if schema_cleanup is None
                else []
            )
        )
        if desired_removals != current_removals:
            patch["attribute_removals"] = desired_removals

        content = cls._stored_json(draft.content_json, {})
        description = cls._cell_text(
            values.get("description"),
            "Описание",
            100_000,
        )
        if description and description != content.get("description"):
            next_content = dict(content)
            next_content["description"] = description
            patch["content"] = next_content

        commercial = cls._stored_json(draft.commercial_json, {})
        if values.get("price_rub") not in (None, ""):
            price = cls._positive_decimal_cell(
                values["price_rub"],
                "Цена",
            )
            if price != commercial.get("price"):
                next_commercial = dict(commercial)
                next_commercial["price"] = price
                patch["commercial"] = next_commercial

        physical_keys = (
            "package_width_mm",
            "package_height_mm",
            "package_depth_mm",
            "package_weight_g",
        )
        if any(values.get(key) not in (None, "") for key in physical_keys):
            if any(values.get(key) in (None, "") for key in physical_keys):
                raise MarketplaceDraftValidationError(
                    "Габариты упаковки заполняются только полным набором: "
                    "ширина, высота, длина и вес"
                )
            dimensions = {
                "width": cls._positive_decimal_cell(
                    values["package_width_mm"],
                    "Ширина упаковки",
                    integer=True,
                ),
                "height": cls._positive_decimal_cell(
                    values["package_height_mm"],
                    "Высота упаковки",
                    integer=True,
                ),
                "depth": cls._positive_decimal_cell(
                    values["package_depth_mm"],
                    "Длина упаковки",
                    integer=True,
                ),
                "weight": cls._positive_decimal_cell(
                    values["package_weight_g"],
                    "Вес с упаковкой",
                    integer=True,
                ),
                "dimension_unit": "MILLIMETERS",
                "weight_unit": "GRAMS",
            }
            if dimensions != cls._stored_json(draft.dimensions_json, {}):
                patch["dimensions"] = dimensions

        current_attributes = cls._stored_json(draft.attributes_json, [])
        next_attributes = [
            item
            for item in current_attributes
            if (
                not schema_cleanup
                or MarketplaceDraftService._attribute_identity(item)
                not in cleanup_identities
            )
        ]
        positions = {
            (
                item.get("attribute_id"),
                item.get("complex_id", "0"),
            ): index
            for index, item in enumerate(next_attributes)
            if isinstance(item, dict)
        }
        attributes_changed = next_attributes != current_attributes
        for key in attribute_keys:
            raw_value = values.get(key)
            if raw_value in (None, ""):
                continue
            external_id = key.split(":", 1)[1]
            definitions = MarketplaceAttributeDefinition.query.filter_by(
                product_type_id=draft.product_type_id,
                external_attribute_id=external_id,
                is_available=True,
                is_enabled=True,
            ).all()
            if (
                len(definitions) != 1
                or definitions[0].attribute_complex_id
            ):
                raise MarketplaceDraftValidationError(
                    f"Атрибут Ozon {external_id} недоступен как простое поле "
                    "этого типа"
                )
            definition = definitions[0]
            resolved = cls._resolve_attribute_values(
                definition=definition,
                raw_value=raw_value,
                expected_external_value_id=values.get(
                    f"attribute_value_id:{external_id}"
                ),
            )
            if not resolved:
                continue
            identity = (external_id, "0")
            item = {
                "attribute_id": external_id,
                "complex_id": "0",
                "values": resolved,
            }
            position = positions.get(identity)
            if position is None:
                positions[identity] = len(next_attributes)
                next_attributes.append(item)
                attributes_changed = True
            elif next_attributes[position] != item:
                next_attributes[position] = item
                attributes_changed = True
        if attributes_changed:
            patch["attributes"] = next_attributes

        current_complex = cls._stored_json(
            draft.complex_attributes_json,
            [],
        )
        next_complex = current_complex
        if schema_cleanup:
            next_complex = []
            for group in current_complex:
                if not isinstance(group, dict):
                    next_complex.append(group)
                    continue
                remaining = [
                    item
                    for item in group.get("attributes", [])
                    if MarketplaceDraftService._attribute_identity(item)
                    not in cleanup_identities
                ]
                if remaining:
                    next_complex.append({"attributes": remaining})
        if next_complex != current_complex:
            patch["complex_attributes"] = next_complex
        return patch

    @classmethod
    def _item_validation_update(
        cls,
        *,
        item: dict,
        draft: MarketplaceProductDraft,
    ) -> dict:
        errors = cls._validation_errors(draft)
        item["completeness"] = MarketplaceDraftService.completeness_summary(
            draft
        )
        item["updated_at"] = datetime.utcnow().isoformat()
        if draft.status == "ready" and not errors:
            item.update({
                "status": "ready_to_retry",
                "code": "local_repair_ready",
                "message": (
                    "Данные заполнены и локально проверены. "
                    "Нажмите «Повторить готовые», чтобы отдельно подтвердить "
                    "отправку в Ozon"
                ),
            })
            item.pop("validation_errors", None)
            return {"status": "ready_to_retry", "errors": []}
        bounded = [
            {
                "code": OzonBulkUploadService._safe_text(
                    error.get("code"),
                    80,
                ) or "draft_not_ready",
                "field": OzonBulkUploadService._safe_text(
                    error.get("field"),
                    120,
                ) or "draft",
                "message": OzonBulkUploadService._safe_text(
                    error.get("message"),
                    300,
                ) or "Черновик требует проверки",
            }
            for error in errors[
                : OzonBulkUploadService.MAX_VALIDATION_ERRORS
            ]
        ]
        item.update({
            "status": "needs_input",
            "code": (
                bounded[0]["code"] if bounded else "draft_not_ready"
            ),
            "message": (
                bounded[0]["message"]
                if bounded else "Черновик требует проверки"
            ),
            "validation_errors": bounded,
        })
        return {"status": "needs_input", "errors": bounded}

    @classmethod
    def _apply_parsed_rows(
        cls,
        *,
        seller_id: int,
        job_uid: str,
        account_id: int,
        parsed_rows: Sequence[dict],
        default_attribute_keys: Sequence[str] = (),
        corrected_by_user_id: Optional[int] = None,
    ) -> dict:
        """Apply normalized local edits under one exact account claim."""
        job, document = cls._owned_completed_run(
            seller_id=seller_id,
            job_uid=job_uid,
        )
        if document["account_id"] != account_id:
            raise cls._conflict(
                "ozon_repair_account_mismatch",
                "Кабинет запуска изменился; обновите страницу",
            )
        item_by_draft = {
            item["draft_id"]: item
            for item in cls._repairable_items(document)
        }
        if any(row["draft_id"] not in item_by_draft for row in parsed_rows):
            raise cls._conflict(
                "ozon_repair_row_not_in_run",
                "Исправление содержит черновик вне текущего запуска",
            )

        claim = try_account_operation_lock(account_id)
        if claim is None:
            raise cls._conflict(
                "ozon_repair_account_busy",
                "Кабинет сейчас занят другой операцией; повторите через минуту",
            )
        report = {
            "total": len(parsed_rows),
            "updated": 0,
            "ready_to_retry": 0,
            "needs_input": 0,
            "excluded": 0,
            "failed": 0,
            "errors": [],
            "rows": [],
        }
        try:
            # Re-read after the account claim. A stale editor/workbook must not
            # race a publication queued between initial render and submit.
            db.session.expire_all()
            job, document = cls._owned_completed_run(
                seller_id=seller_id,
                job_uid=job_uid,
            )
            if document["account_id"] != account_id:
                raise cls._conflict(
                    "ozon_repair_account_mismatch",
                    "Кабинет запуска изменился; обновите страницу",
                )
            item_by_draft = {
                item["draft_id"]: item
                for item in cls._repairable_items(document)
            }
            for row in parsed_rows:
                item = item_by_draft.get(row["draft_id"])
                if item is None:
                    report["failed"] += 1
                    report["rows"].append({"draft_id": row["draft_id"], "status": "failed",
                        "message": "Статус строки изменился; обновите страницу"})
                    report["errors"].append({
                        "row": row.get("row_number"),
                        "draft_id": row["draft_id"],
                        "message": (
                            "Статус строки изменился; обновите страницу"
                        ),
                    })
                    continue
                try:
                    draft = MarketplaceDraftService.get_draft(
                        seller_id=seller_id,
                        draft_id=row["draft_id"],
                    )
                    if (
                        draft.account_id != account_id
                        or draft.imported_product_id
                        != row["imported_product_id"]
                    ):
                        raise MarketplaceDraftValidationError(
                            "Связь черновика с товаром или кабинетом изменилась"
                        )
                    if draft.version != row["expected_version"]:
                        raise MarketplaceDraftValidationError(
                            "Черновик изменился после открытия; обновите страницу"
                        )
                    action = cls._cell_text(
                        row["values"].get("action"),
                        "Действие",
                        30,
                    ).upper()
                    if action not in {
                        cls.ACTION_REPAIR,
                        cls.ACTION_EXCLUDE,
                    }:
                        raise MarketplaceDraftValidationError(
                            "Действие должно быть ИСПРАВИТЬ или ИСКЛЮЧИТЬ"
                        )
                    if action == cls.ACTION_EXCLUDE:
                        item.update({
                            "status": "excluded",
                            "code": "seller_excluded_from_retry",
                            "message": (
                                "Исключено продавцом только из повторов этого "
                                "запуска; товар и черновик не удалены"
                            ),
                            "updated_at": datetime.utcnow().isoformat(),
                        })
                        item.pop("validation_errors", None)
                        item.pop("repair_error", None)
                        report["excluded"] += 1
                        OzonBulkUploadService._persist(job, document)
                        report["rows"].append({"draft_id": draft.id, "status": "excluded",
                                               "version": draft.version})
                        continue

                    attribute_keys = row.get(
                        "attribute_keys",
                        default_attribute_keys,
                    )
                    patch = cls._row_patch(
                        draft=draft,
                        values=row["values"],
                        attribute_keys=attribute_keys,
                        schema_cleanup=row.get(
                            "schema_cleanup",
                        ),
                        cleanup_candidates=row.get(
                            "cleanup_candidates",
                            (),
                        ),
                    )
                    selected_type_id = row.get("product_type_id")
                    save_mapping = row.get("save_mapping", False)
                    if selected_type_id is not None:
                        if (
                            selected_type_id != draft.product_type_id
                            and any(
                                row["values"].get(key) not in (None, "")
                                for key in attribute_keys
                            )
                        ):
                            raise MarketplaceDraftValidationError(
                                "Сначала сохраните новый тип Ozon, затем "
                                "заполните его обязательные атрибуты"
                            )
                        if (
                            selected_type_id != draft.product_type_id
                            and row.get("schema_cleanup") is True
                        ):
                            raise MarketplaceDraftValidationError(
                                "Сначала сохраните новый тип Ozon, затем "
                                "повторно проверьте очистку его атрибутов"
                            )
                        if selected_type_id != draft.product_type_id:
                            patch["product_type_id"] = selected_type_id
                    if save_mapping:
                        if selected_type_id is None:
                            raise MarketplaceDraftValidationError(
                                "Для сохранения категории выберите тип Ozon"
                            )
                        patch["product_type_id"] = selected_type_id
                        patch["save_mapping"] = True
                    if patch:
                        draft = MarketplaceDraftService.update_draft(
                            seller_id=seller_id,
                            draft_id=draft.id,
                            expected_version=draft.version,
                            patch=patch,
                            corrected_by_user_id=corrected_by_user_id,
                            # This batch already validated its own reviewed
                            # category/reset contract above; no single-draft token.
                            category_review_required=False,
                        )
                        report["updated"] += 1
                    draft = MarketplaceDraftService.validate_draft(
                        seller_id=seller_id,
                        draft_id=draft.id,
                        expected_version=draft.version,
                    )
                    outcome = cls._item_validation_update(
                        item=item,
                        draft=draft,
                    )
                    item.pop("repair_error", None)
                    report[outcome["status"]] += 1
                    OzonBulkUploadService._persist(job, document)
                    report["rows"].append({"draft_id": draft.id, "status": outcome["status"],
                                           "version": draft.version})
                except MarketplaceDraftError as exc:
                    db.session.rollback()
                    report["failed"] += 1
                    safe_message = cls._safe_text(exc, 500)
                    report["rows"].append({"draft_id": row["draft_id"], "status": "failed",
                                           "message": safe_message})
                    if len(report["errors"]) < 30:
                        report["errors"].append({
                            "row": row.get("row_number"),
                            "draft_id": row["draft_id"],
                            "message": safe_message,
                        })
                    # Already committed rows remain durable. Reload both the
                    # run and exact item after rollback, then persist this
                    # row-local editor error so the UI explains the failure.
                    job = OzonBulkUploadService.get_run(
                        seller_id=seller_id,
                        job_uid=job_uid,
                        reconcile=False,
                    )
                    document = OzonBulkUploadService._load_progress(job)
                    item_by_draft = {
                        current["draft_id"]: current
                        for current in cls._repairable_items(document)
                    }
                    failed_item = item_by_draft.get(row["draft_id"])
                    if failed_item is not None:
                        failed_item["repair_error"] = safe_message
                        failed_item["updated_at"] = (
                            datetime.utcnow().isoformat()
                        )
                        OzonBulkUploadService._persist(job, document)
            return report
        finally:
            release_account_operation_lock(claim)

    @classmethod
    def apply_editor_rows(
        cls,
        *,
        seller_id: int,
        job_uid: str,
        rows: Sequence[dict],
        corrected_by_user_id: Optional[int] = None,
    ) -> dict:
        """Apply edits submitted by the platform-native mass editor."""
        if (
            not isinstance(rows, (list, tuple))
            or not rows
            or len(rows) > cls.MAX_ROWS
        ):
            raise OzonBulkUploadValidationError(
                f"Выберите от 1 до {cls.MAX_ROWS} карточек"
            )
        editor = cls.editor_document(
            seller_id=seller_id,
            job_uid=job_uid,
        )
        editable_attributes = {
            row["draft_id"]: {
                field["external_id"]
                for field in row["attributes"]
                if field["editable"]
            }
            for group in editor["groups"]
            for row in group["rows"]
        }
        identity_by_draft = {
            row["draft_id"]: row
            for group in editor["groups"]
            for row in group["rows"]
        }
        parsed_rows = []
        seen = set()
        fixed_value_keys = set(cls.EDITABLE_FIXED)
        for index, row in enumerate(rows, start=1):
            if not isinstance(row, dict) or set(row) - {
                "draft_id",
                "expected_version",
                "imported_product_id",
                "action",
                "product_type_id",
                "save_mapping",
                "schema_cleanup",
                "values",
            }:
                raise OzonBulkUploadValidationError(
                    f"Строка {index}: неизвестная структура"
                )
            for field_name in (
                "draft_id",
                "expected_version",
                "imported_product_id",
            ):
                value = row.get(field_name)
                if (
                    isinstance(value, bool)
                    or not isinstance(value, int)
                    or value <= 0
                ):
                    raise OzonBulkUploadValidationError(
                        f"Строка {index}: {field_name} должен быть "
                        "положительным целым числом"
                    )
            draft_id = row["draft_id"]
            if draft_id in seen:
                raise OzonBulkUploadValidationError(
                    f"Черновик {draft_id} продублирован"
                )
            seen.add(draft_id)
            expected = identity_by_draft.get(draft_id)
            if (
                expected is None
                or expected["imported_product_id"]
                != row["imported_product_id"]
            ):
                raise cls._conflict(
                    "ozon_repair_row_not_in_run",
                    "Строка не относится к этому запуску",
                )
            product_type_id = row.get("product_type_id")
            if product_type_id is not None and (
                isinstance(product_type_id, bool)
                or not isinstance(product_type_id, int)
                or product_type_id <= 0
            ):
                raise OzonBulkUploadValidationError(
                    f"Строка {index}: product_type_id некорректен"
                )
            save_mapping = row.get("save_mapping", False)
            if not isinstance(save_mapping, bool):
                raise OzonBulkUploadValidationError(
                    f"Строка {index}: save_mapping должен быть boolean"
                )
            schema_cleanup = row.get("schema_cleanup", False)
            if not isinstance(schema_cleanup, bool):
                raise OzonBulkUploadValidationError(
                    f"Строка {index}: schema_cleanup должен быть boolean"
                )
            if schema_cleanup and not expected["cleanup_candidates"]:
                raise OzonBulkUploadValidationError(
                    f"Строка {index}: для карточки нет атрибутов на очистку"
                )
            values = row.get("values")
            if not isinstance(values, dict) or len(values) > cls.MAX_COLUMNS:
                raise OzonBulkUploadValidationError(
                    f"Строка {index}: слишком много полей"
                )
            attribute_keys = []
            allowed_value_keys = set(fixed_value_keys)
            allowed_value_keys.add("action")
            for external_id in editable_attributes.get(draft_id, set()):
                key = f"attribute:{external_id}"
                attribute_keys.append(key)
                allowed_value_keys.add(key)
                allowed_value_keys.add(
                    f"attribute_value_id:{external_id}"
                )
            if set(values) - allowed_value_keys:
                raise OzonBulkUploadValidationError(
                    f"Строка {index}: передано поле вне текущей схемы"
                )
            values = dict(values)
            values["action"] = row.get("action")
            parsed_rows.append({
                "row_number": index,
                "draft_id": draft_id,
                "expected_version": row["expected_version"],
                "imported_product_id": row["imported_product_id"],
                "product_type_id": product_type_id,
                "save_mapping": save_mapping,
                "schema_cleanup": schema_cleanup,
                "cleanup_candidates": expected[
                    "cleanup_candidates"
                ],
                "values": values,
                "attribute_keys": sorted(
                    attribute_keys,
                    key=lambda key: int(key.split(":", 1)[1]),
                ),
            })
        return cls._apply_parsed_rows(
            seller_id=seller_id,
            job_uid=job_uid,
            account_id=editor["account_id"],
            parsed_rows=parsed_rows,
            corrected_by_user_id=corrected_by_user_id,
        )

    @classmethod
    def import_workbook(
        cls,
        *,
        seller_id: int,
        job_uid: str,
        payload: bytes,
        corrected_by_user_id: Optional[int] = None,
    ) -> dict:
        workbook = cls._load_workbook(payload)
        metadata = cls._metadata(workbook)
        if (
            metadata.get("contract_version") != cls.CONTRACT_VERSION
            or metadata.get("job_uid") != job_uid
        ):
            raise OzonBulkUploadValidationError(
                "Таблица относится к другому запуску или версии формата"
            )
        try:
            metadata_seller_id = cls._positive_integer_cell(
                metadata.get("seller_id"),
                "seller_id",
            )
            metadata_account_id = cls._positive_integer_cell(
                metadata.get("account_id"),
                "account_id",
            )
            expected_columns = json.loads(
                metadata.get("columns_json") or "[]"
            )
        except (MarketplaceDraftValidationError, TypeError, ValueError):
            raise OzonBulkUploadValidationError(
                "Метаданные XLSX повреждены; скачайте новую таблицу"
            ) from None
        if metadata_seller_id != seller_id:
            raise cls._conflict(
                "ozon_repair_tenant_mismatch",
                "Таблица принадлежит другому продавцу",
            )

        cards = workbook[cls.CARD_SHEET]
        if cards.max_row > cls.MAX_ROWS + 2 or cards.max_column > cls.MAX_COLUMNS:
            raise OzonBulkUploadValidationError(
                "В XLSX слишком много строк или столбцов"
            )
        columns = [
            cell.value
            for cell in cards[1]
            if cell.column <= cards.max_column
        ]
        if (
            not isinstance(expected_columns, list)
            or columns != expected_columns
            or len(columns) != len(set(columns))
        ):
            raise OzonBulkUploadValidationError(
                "Структура столбцов изменена; не удаляйте и не добавляйте столбцы"
            )
        fixed_keys = {key for key, _label in cls.FIXED_COLUMNS}
        if not fixed_keys.issubset(columns) or any(
            (
                not isinstance(key, str)
                or (
                    key not in fixed_keys
                    and cls.ATTRIBUTE_KEY.fullmatch(key) is None
                )
            )
            for key in columns
        ):
            raise OzonBulkUploadValidationError(
                "XLSX содержит неизвестные столбцы"
            )
        column_index = {
            key: index
            for index, key in enumerate(columns, start=1)
        }
        attribute_keys = [
            key for key in columns if key.startswith("attribute:")
        ]
        parsed_rows = []
        seen_draft_ids = set()
        for row_number in range(3, cards.max_row + 1):
            if all(
                cards.cell(row=row_number, column=index).value in (None, "")
                for index in range(1, len(columns) + 1)
            ):
                continue
            for index in range(1, len(columns) + 1):
                if cards.cell(row=row_number, column=index).data_type == "f":
                    raise OzonBulkUploadValidationError(
                        f"Строка {row_number}: формулы запрещены; "
                        "введите итоговое значение"
                    )
            try:
                draft_id = cls._positive_integer_cell(
                    cards.cell(
                        row=row_number,
                        column=column_index["draft_id"],
                    ).value,
                    "draft_id",
                )
                expected_version = cls._positive_integer_cell(
                    cards.cell(
                        row=row_number,
                        column=column_index["draft_version"],
                    ).value,
                    "draft_version",
                )
                imported_product_id = cls._positive_integer_cell(
                    cards.cell(
                        row=row_number,
                        column=column_index["imported_product_id"],
                    ).value,
                    "imported_product_id",
                )
            except MarketplaceDraftValidationError as exc:
                raise OzonBulkUploadValidationError(
                    f"Строка {row_number}: {exc}"
                ) from None
            if draft_id in seen_draft_ids:
                raise OzonBulkUploadValidationError(
                    f"Черновик {draft_id} продублирован в XLSX"
                )
            seen_draft_ids.add(draft_id)
            parsed_rows.append({
                "row_number": row_number,
                "draft_id": draft_id,
                "expected_version": expected_version,
                "imported_product_id": imported_product_id,
                "values": {
                    key: cards.cell(
                        row=row_number,
                        column=index,
                    ).value
                    for key, index in column_index.items()
                },
            })
        if not parsed_rows:
            raise OzonBulkUploadValidationError(
                "В XLSX нет строк для обработки"
            )

        return cls._apply_parsed_rows(
            seller_id=seller_id,
            job_uid=job_uid,
            account_id=metadata_account_id,
            parsed_rows=parsed_rows,
            default_attribute_keys=attribute_keys,
            corrected_by_user_id=corrected_by_user_id,
        )
